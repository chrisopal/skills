#!/usr/bin/env python3
"""Build a conservative editable working draft from a document specification.

Not a universal template-preserving engine. Fixed forms/complex attachments require host tools.
"""
from __future__ import annotations
import argparse, json, sys
from pathlib import Path

def build(spec_path: Path,out_path: Path,asset_root: Path):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    spec=json.loads(spec_path.read_text(encoding='utf-8'))
    if out_path.exists():raise ValueError('拒绝覆盖已有文档，请使用新版本文件名')
    document=Document();section=document.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(2.3);section.left_margin=section.right_margin=Cm(2.5)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=document.styles[name];style.font.name='Arial';style.font.size=Pt(11 if name=='Normal' else (22 if name=='Title' else 13))
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'SimSun' if name=='Normal' else 'SimHei')
        style.paragraph_format.space_after=Pt(7);style.paragraph_format.line_spacing=1.25
    document.core_properties.title=spec['title'];document.core_properties.author='';document.core_properties.last_modified_by=''
    section.header.paragraphs[0].text='内部工作稿｜需人工复核与签署'
    footer=section.footer.paragraphs[0];footer.alignment=2
    footer.add_run('第 ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    footer.add_run(' 页')
    document.add_heading(spec['title'],0)
    if spec.get('subtitle'):document.add_paragraph(spec['subtitle'],'Subtitle')
    document.add_paragraph('本文件由基础排版脚本生成，不代表真实投标通过审核。客户指定版式、目录域与签署要求需在正式交付前另行落实。')
    for item in spec['sections']:
        document.add_heading(item['title'],level=max(1,min(int(item.get('level',1)),3)))
        for text in item.get('paragraphs',[]):document.add_paragraph(str(text))
        for table_spec in item.get('tables',[]):
            headers=table_spec['headers'];rows=table_spec['rows']
            if not headers or any(len(row)!=len(headers) for row in rows):raise ValueError('表格列数不一致')
            table=document.add_table(rows=1,cols=len(headers));table.style='Table Grid'
            trPr=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
            for i,v in enumerate(headers):table.rows[0].cells[i].text=str(v)
            for values in rows:
                cells=table.add_row().cells
                for i,v in enumerate(values):cells[i].text=str(v)
            document.add_paragraph('')
        for image in item.get('images',[]):
            rel=Path(image['path']);path=(asset_root/rel).resolve()
            if rel.is_absolute() or '..' in rel.parts or not path.is_relative_to(asset_root.resolve()):raise ValueError('图片路径越界')
            if not path.is_file():raise ValueError('图片缺失：'+str(rel))
            if path.suffix.lower() not in {'.png','.jpg','.jpeg'}:raise ValueError('此基础生成器仅直接插入PNG/JPEG，SVG须先受控渲染')
            document.add_picture(str(path),width=Cm(min(float(image.get('width_cm',15)),16)))
            document.add_paragraph(image.get('caption',''))
    out_path.parent.mkdir(parents=True,exist_ok=True);document.save(out_path)
    return {'output':str(out_path),'status':'draft_only','visual_qa':'NOT_RUN','notes':['字体文件未随包分发，请在运行环境配置合法字体。','不自动更新目录，不代替复杂原始模板填报和证明材料插页。']}

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--spec',type=Path,required=True);p.add_argument('--out',type=Path,required=True);p.add_argument('--asset-root',type=Path);a=p.parse_args()
    try:print(json.dumps(build(a.spec,a.out,a.asset_root or a.spec.parent),ensure_ascii=False,indent=2));return 0
    except (ImportError,OSError,ValueError,KeyError) as e:print('DOCX生成失败：'+str(e),file=sys.stderr);return 2
if __name__=='__main__':raise SystemExit(main())
