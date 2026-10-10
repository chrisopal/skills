#!/usr/bin/env python3
"""Deterministic completeness checks for the DOCX/PDF built from document-spec.

This module checks the bytes which were actually written.  It does not certify
semantic correctness, business approval, signatures, or visual acceptance.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import platform
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Any

try:
    from .inspect_artifact import inspect as inspect_artifact
except ImportError:  # pragma: no cover - permits running this file directly
    try:
        from inspect_artifact import inspect as inspect_artifact
    except ImportError:
        inspect_artifact = None


SCHEMA_VERSION = "export-acceptance-v1"
_TOC_PLACEHOLDER = "目录字段将在支持的Word交付引擎中更新"
_BAD_TEXT = ("\ufffd", "\ue200", "\u2eda")


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _normal(value: Any) -> str:
    """Ignore layout whitespace, preserving compatibility glyph anomalies."""
    text = unicodedata.normalize("NFC", str(value))
    return re.sub(r"\s+", "", text)


def _pixels(content: bytes) -> dict:
    import fitz
    pixmap = fitz.Pixmap(content)
    if pixmap.alpha:
        pixmap = fitz.Pixmap(pixmap, 0)
    pixmap = fitz.Pixmap(fitz.csRGB, pixmap)
    return {"width": pixmap.width, "height": pixmap.height,
            "sha256": hashlib.sha256(pixmap.samples).hexdigest()}


def _relative(path: Path, root: Path) -> str:
    return str(path.resolve().relative_to(root.resolve()))


def _read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("document-spec顶层必须是对象")
    return value


def _project_id(project_root: Path) -> str | None:
    metadata = project_root / "work" / "project.json"
    if not metadata.is_file():
        return None
    value = _read_json(metadata).get("project_id")
    return value if isinstance(value, str) and value.strip() else None


def _facts(spec: dict[str, Any]) -> list[dict[str, Any]]:
    raw = spec.get("expected_facts", [])
    if not isinstance(raw, list):
        raise ValueError("expected_facts必须是数组")
    result = []
    for index, item in enumerate(raw):
        if isinstance(item, str):
            result.append({"id": f"FACT-{index + 1}", "value": item, "required": True})
            continue
        if not isinstance(item, dict) or not isinstance(item.get("value"), (str, int, float)):
            raise ValueError(f"expected_facts[{index}]必须是文本或包含value的对象")
        result.append(
            {
                "id": str(item.get("id", f"FACT-{index + 1}")),
                "value": str(item["value"]),
                "required": bool(item.get("required", True)),
            }
        )
    return result


def _expected(spec: dict[str, Any]) -> dict[str, Any]:
    sections = spec.get("sections")
    if not isinstance(spec.get("title"), str) or not spec["title"].strip():
        raise ValueError("document-spec.title必须是非空文本")
    if not isinstance(sections, list) or not sections:
        raise ValueError("document-spec.sections必须是非空数组")
    from build_docx import header_footer_settings
    settings = spec.get("export_settings", {})
    if not isinstance(settings, dict):
        raise ValueError("export_settings必须是对象")
    cover = settings.get("cover", {})
    toc = settings.get("toc", {})
    if not isinstance(cover, dict) or not isinstance(toc, dict):
        raise ValueError("export_settings.cover/toc必须是对象")
    cover_enabled = cover.get("enabled", True)
    toc_enabled = toc.get("enabled", False)
    if not isinstance(cover_enabled, bool) or not isinstance(toc_enabled, bool):
        raise ValueError("cover.enabled/toc.enabled必须是布尔值")

    cover_paragraphs = []
    tables: list[dict[str, Any]] = []
    image_count = 0
    image_sources = []
    if cover_enabled:
        cover_paragraphs.append(str(cover.get("title", spec["title"])))
        subtitle = cover.get("subtitle", spec.get("subtitle", ""))
        if subtitle:
            cover_paragraphs.append(str(subtitle))
        metadata = cover.get("metadata_rows", [])
        if not isinstance(metadata, list):
            raise ValueError("cover.metadata_rows必须是数组")
        if metadata:
            if any(not isinstance(row, dict) for row in metadata):
                raise ValueError("cover.metadata_rows中的行必须是对象")
            tables.append(
                {
                    "kind": "cover",
                    "rows": [
                        [str(row.get("label", "")), str(row.get("value", ""))] for row in metadata
                    ],
                }
            )
        if cover.get("image") is not None:
            image_count += 1
            image_sources.append(cover["image"]["path"])

    headings = []
    paragraphs = []
    ordered = list(cover_paragraphs)
    for table in tables:
        ordered.extend(value for row in table["rows"] for value in row)
    for index, section in enumerate(sections):
        if not isinstance(section, dict) or not isinstance(section.get("title"), str):
            raise ValueError(f"sections[{index}]必须是包含title的对象")
        title = section["title"]
        level = max(1, min(int(section.get("level", 1)), 3))
        headings.append({"text": title, "level": level})
        ordered.append(title)
        section_paragraphs = section.get("paragraphs", [])
        if not isinstance(section_paragraphs, list):
            raise ValueError(f"sections[{index}].paragraphs必须是数组")
        for value in section_paragraphs:
            if str(value).strip():
                text = str(value)
                paragraphs.append(text)
                ordered.append(text)
        for table in section.get("tables", []):
            if (
                not isinstance(table, dict)
                or not isinstance(table.get("headers"), list)
                or not isinstance(table.get("rows", []), list)
            ):
                raise ValueError(f"sections[{index}].tables中的表格无效")
            rows = [[str(value) for value in table["headers"]]] + [
                [str(value) for value in row] for row in table.get("rows", [])
            ]
            if any(len(row) != len(rows[0]) for row in rows):
                raise ValueError(f"sections[{index}]表格列数不一致")
            tables.append({"kind": "section", "rows": rows})
            ordered.extend(value for row in rows for value in row)
        for image in section.get("images", []):
            if not isinstance(image, dict):
                raise ValueError(f"sections[{index}].images中的图片无效")
            image_count += 1
            image_sources.append(image["path"])
            if image.get("caption", ""):
                ordered.append(str(image["caption"]))

    if toc_enabled:
        cover_table_values = sum(
            len(row) for table in tables if table["kind"] == "cover" for row in table["rows"]
        )
        ordered.insert(len(cover_paragraphs) + cover_table_values, str(toc.get("title", "目录")))
    return {
        "cover_paragraphs": cover_paragraphs,
        "headings": headings,
        "paragraphs": paragraphs,
        "tables": tables,
        "images": image_count,
        "image_sources": image_sources,
        "toc_levels": int(toc.get("levels", 3)),
        "facts": _facts(spec),
        "toc_enabled": toc_enabled,
        "header_footer": [header_footer_settings(settings.get(name, {}), name) for name in ("header", "footer")],
        "toc_title": str(toc.get("title", "目录")) if toc_enabled else None,
        "ordered_text": ordered,
    }


def _bad_text_findings(text: str) -> list[str]:
    findings = []
    for token in _BAD_TEXT:
        if token in text:
            findings.append(f"文字层含异常编码字符：{token!r}")
    if re.search(r"[\u2f00-\u2fdf]", text):
        findings.append("文字层含康熙部首编码，需修复导出字体映射")
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", text):
        findings.append("文字层含不可解释控制字符")
    if not _normal(text):
        findings.append("没有可检查文字层，不能自动通过")
    return findings


def _subsequence(expected: list[str], actual: list[str]) -> tuple[list[str], list[str]]:
    present, absent, cursor = [], [], 0
    for value in expected:
        target = _normal(value)
        found = next((i for i in range(cursor, len(actual)) if _normal(actual[i]) == target), None)
        if found is None:
            absent.append(value)
        else:
            present.append(value)
            cursor = found + 1
    return present, absent


def _occurrence_order(expected: list[str], text: str) -> tuple[list[str], list[str]]:
    present, absent, cursor = [], [], 0
    compact = _normal(text)
    for value in expected:
        target = _normal(value)
        position = compact.find(target, cursor) if target else -1
        if position < 0:
            absent.append(value)
        else:
            present.append(value)
            cursor = position + len(target)
    return present, absent


def _artifact_meta(path: Path, project_root: Path, pages: int | None = None) -> dict[str, Any]:
    try:
        relative = path.resolve().relative_to(project_root.resolve()).as_posix()
    except ValueError:
        relative = str(path)
    return {"path": relative, "sha256": _sha256(path), "bytes": path.stat().st_size, "pages": pages}


def _docx_snapshot(path: Path, expected: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    findings: list[str] = []
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    heading_records = []
    for paragraph in document.paragraphs:
        match = re.search(r"(?:Heading|标题)\s*([1-9])", paragraph.style.name, flags=re.I)
        if match and paragraph.text.strip():
            heading_records.append({"text": paragraph.text, "level": int(match.group(1))})
    expected_heading_text = [item["text"] for item in expected["headings"]]
    actual_heading_text = [item["text"] for item in heading_records]
    heading_present, heading_absent = _subsequence(expected_heading_text, actual_heading_text)
    if heading_absent:
        findings.append("DOCX缺少章节标题：" + "、".join(heading_absent))
    for want, got in zip(expected["headings"], heading_records):
        if _normal(want["text"]) == _normal(got["text"]) and want["level"] != got["level"]:
            findings.append(f"DOCX章节层级不一致：{want['text']}")

    cover_present, cover_absent = _subsequence(expected["cover_paragraphs"], paragraphs)
    if cover_absent:
        findings.append("DOCX缺少封面标题或副标题：" + "、".join(cover_absent))
    paragraph_present, paragraph_absent = _subsequence(expected["paragraphs"], paragraphs)
    if paragraph_absent:
        findings.append("DOCX缺少正文段落：" + "、".join(paragraph_absent))
    actual_tables = [
        [[cell.text for cell in row.cells] for row in table.rows] for table in document.tables
    ]
    table_present, table_absent = [], []
    for index, want in enumerate(expected["tables"]):
        rows = [[_normal(value) for value in row] for row in want["rows"]]
        if index >= len(actual_tables):
            table_absent.append(f"table-{index + 1}")
        elif [[_normal(value) for value in row] for row in actual_tables[index]] != rows:
            table_absent.append(f"table-{index + 1}")
            findings.append(f"DOCX表格内容或顺序不一致：table-{index + 1}")
        else:
            table_present.append(f"table-{index + 1}")
    if len(actual_tables) > len(expected["tables"]):
        findings.append("DOCX包含document-spec未声明的表格")
    if table_absent and not any("表格内容" in item for item in findings):
        findings.append("DOCX缺少表格：" + "、".join(table_absent))

    media_rels = [rel for rel in document.part.rels.values() if rel.reltype == RT.IMAGE]
    inline_images = len(document.inline_shapes)
    with zipfile.ZipFile(path) as archive:
        media_parts = [name for name in archive.namelist() if name.startswith("word/media/")]
        document_xml = archive.read("word/document.xml").decode("utf-8", errors="strict")
    if (inline_images != expected["images"]
            or len(media_rels) != len(media_parts)
            or (expected["images"] and not media_parts)):
        findings.append(
            f"DOCX插图数量不一致：期望{expected['images']}，inline_shapes={inline_images}，关系={len(media_rels)}，media={len(media_parts)}"
        )

    from docx.oxml.ns import qn
    actual_images = []
    for blip in document.element.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid:
            actual_images.append(_pixels(document.part.related_parts[rid].blob))
    if actual_images != [binding["pixels"] for binding in expected["image_bindings"]]:
        findings.append("DOCX插图内容或顺序与源图不一致")
    for section in document.sections:
        for name, settings in zip(('header', 'footer'), expected.get('header_footer', [])):
            if not settings:
                continue
            container = getattr(section, name)
            text = ''.join(node.text or '' for node in container._element.iter(qn('w:t')))
            if settings.get('text') and _normal(settings['text']) not in _normal(text):
                findings.append('DOCX缺少声明的' + name + '文字')
            if settings.get('page_field'):
                codes = [node.text or '' for node in container._element.iter(qn('w:instrText'))]
                codes += [node.get(qn('w:instr'), '') for node in container._element.iter(qn('w:fldSimple'))]
                if not any(re.search(r'\bPAGE\b', code) for code in codes):
                    findings.append('DOCX缺少声明的PAGE页码域')
    full_text = "\n".join(document.element.itertext())
    # Include refreshed TOC inside content controls, which Document.paragraphs omits.
    from docx.oxml.ns import qn
    full_text = "\n".join(node.text or "" for node in document.element.iter(qn("w:t")))
    findings.extend(_bad_text_findings(full_text))
    fact_values = [item["value"] for item in expected["facts"] if item["required"]]
    fact_present = [value for value in fact_values if _normal(value) in _normal(full_text)]
    fact_absent = [value for value in fact_values if value not in fact_present]
    if fact_absent:
        findings.append("DOCX缺少关键事实：" + "、".join(fact_absent))
    if expected["toc_enabled"]:
        if _TOC_PLACEHOLDER in full_text:
            findings.append("DOCX目录仍是未刷新的占位文本")
        if not re.search(r"\bTOC\b[^<]*\\o\b", document_xml):
            findings.append("DOCX缺少document-spec要求的目录字段")
        for title in [h["text"] for h in expected["headings"] if h["level"] <= expected["toc_levels"]]:
            if _normal(full_text).count(_normal(title)) < 2:
                findings.append(f"DOCX目录未发现已刷新的章节条目：{title}")
    elif re.search(r"\bTOC\b[^<]*\\o\b", document_xml):
        findings.append("DOCX包含document-spec未启用的目录字段")

    return {
        "counts": {
            "paragraphs": len([p for p in paragraphs if p.strip()]),
            "headings": len(heading_records),
            "tables": len(actual_tables),
            "images": inline_images,
        },
        "coverage": {
            "headings": {
                "expected": expected_heading_text,
                "present": heading_present,
                "absent": heading_absent,
            },
            "cover": {
                "expected": expected["cover_paragraphs"],
                "present": cover_present,
                "absent": cover_absent,
            },
            "paragraphs": {
                "expected": expected["paragraphs"],
                "present": paragraph_present,
                "absent": paragraph_absent,
            },
            "tables": {
                "expected": [item["rows"] for item in expected["tables"]],
                "present": table_present,
                "absent": table_absent,
            },
            "facts": {
                "expected": [item["value"] for item in expected["facts"]],
                "present": fact_present,
                "absent": fact_absent,
            },
        },
    }, findings


def _pdf_snapshot(path: Path, expected: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    import fitz

    findings: list[str] = []
    with fitz.open(path) as document:
        pages = len(document)
        text = "\n".join(page.get_text("text") for page in document)
        image_count = sum(len(page.get_image_info()) for page in document)
        for index, page in enumerate(document):
            page_text = _normal(page.get_text('text'))
            for name, settings in zip(('header', 'footer'), expected.get('header_footer', [])):
                if settings.get('text') and _normal(settings['text']) not in page_text:
                    findings.append(f'PDF第{index + 1}页缺少声明的{name}文字')
                if settings.get('page_field'):
                    pattern = re.escape(_normal(settings.get('text', '') + settings.get('page_prefix', ''))) + r'\d+' + re.escape(_normal(settings.get('page_suffix', '')))
                    if not re.search(pattern, page_text):
                        findings.append(f'PDF第{index + 1}页缺少声明的页码')
    findings.extend(_bad_text_findings(text))
    # Page headers/footers can interrupt a paragraph spanning two pages.
    running = expected.get('header_footer', [])
    lines = []
    for line in text.splitlines():
        normalized = _normal(line)
        ignore = False
        for item in running:
            if item.get('text') and normalized == _normal(item['text']):
                ignore = True
            if item.get('page_field'):
                page_pattern = re.escape(_normal(item.get('text', '') + item.get('page_prefix', ''))) + r'\d+' + re.escape(_normal(item.get('page_suffix', '')))
                if re.fullmatch(page_pattern, normalized):
                    ignore = True
        if not ignore:
            lines.append(line)
    text = '\n'.join(lines)
    if not _normal(text):
        findings.append("PDF文字层未知或为空，不能自动通过")
    ordered_present, ordered_absent = _occurrence_order(expected["ordered_text"], text)
    if ordered_absent:
        findings.append("PDF缺少正文/标题/表格/图注或顺序不一致：" + "、".join(ordered_absent))
    fact_values = [item['value'] for item in expected['facts'] if item['required']]
    fact_present = [value for value in fact_values if _normal(value) in _normal(text)]
    fact_absent = [value for value in fact_values if value not in fact_present]
    if fact_absent:
        findings.append("PDF缺少关键事实：" + "、".join(fact_absent))
    if image_count != expected["images"]:
        findings.append(f"PDF插图数量不一致：期望{expected['images']}，实际{image_count}")
    if expected["toc_enabled"]:
        if _TOC_PLACEHOLDER in _normal(text):
            findings.append("PDF目录仍是未刷新的占位文本")
        for title in [item["text"] for item in expected["headings"] if item["level"] <= expected["toc_levels"]]:
            if _normal(text).count(_normal(title)) < 2:
                findings.append(f"PDF目录未发现已刷新的章节条目：{title}")
    return {
        "counts": {"pages": pages, "extracted_chars": len(text), "images": image_count},
        "coverage": {
            "ordered_text": {
                "expected": expected["ordered_text"],
                "present": ordered_present,
                "absent": ordered_absent,
            },
            "facts": {
                "expected": [item["value"] for item in expected["facts"]],
                "present": fact_present,
                "absent": fact_absent,
            },
        },
    }, findings


def _tool_versions() -> dict[str, Any]:
    versions = {"python": platform.python_version(), "checker": SCHEMA_VERSION}
    for package, key in (("python-docx", "python_docx"), ("PyMuPDF", "pymupdf")):
        try:
            versions[key] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            versions[key] = "unavailable"
    return versions


def _check_current_chapters(spec, project_root):
    outline = project_root / 'artifacts/08-outline.json'
    writing = project_root / 'artifacts/11-technical-content.json'
    findings = []
    commercial = project_root / 'artifacts/12-commercial-content.json'
    current_paths = [p for p in (outline, writing, commercial) if p.is_file()]
    for path in current_paths:
        if _read_json(path).get('project_id') != _project_id(project_root):
            findings.append('导出源产物项目身份不一致：' + str(path.name))
    bindings = {item['relative_path']: item['sha256'] for item in spec.get('inputs', [])}
    for path in current_paths:
        relative = _relative(path, project_root)
        if bindings.get(relative) != _sha256(path):
            findings.append('导出输入未绑定当前版本：' + relative)
    sections = spec['sections']
    if outline.is_file():
        from outline_view import ordered_sections
        expected_ids = [item['id'] for item in ordered_sections(_read_json(outline)['data']['sections'])]
        if [item.get('section_id') for item in sections] != expected_ids:
            findings.append('导出章节与当前目录不一致或顺序错误')
    commercial = project_root / 'artifacts/12-commercial-content.json'
    if commercial.is_file():
        value = _read_json(commercial)
        if value.get('project_id') != _project_id(project_root):
            findings.append('导出源商务产物项目身份不一致')
        if bindings.get(_relative(commercial, project_root)) != _sha256(commercial):
            findings.append('导出输入未绑定当前商务版本')
        exported = ''.join(str(p) for section in sections for p in section.get('paragraphs', []))
        exported += ''.join(str(v) for section in sections for table in section.get('tables', []) for row in [table['headers']] + table.get('rows', []) for v in row)
        for form in value.get('data', {}).get('forms', []):
            for field in form.get('fields', []):
                expected_value = field.get('value')
                if expected_value is not None and str(expected_value).strip() and _normal(expected_value) not in _normal(exported):
                    findings.append('导出规格漏商务字段：' + str(form['id']) + ':' + str(field['key']))
    by_id = {item.get('section_id'): item for item in sections}
    for chapter in (_read_json(writing)['data']['chapters'] if writing.is_file() else []):
        section = by_id.get(chapter['section_id'], {})
        content = ''.join(str(value) for value in section.get('paragraphs', []))
        for table in section.get('tables', []):
            content += ''.join(str(v) for row in [table['headers']] + table.get('rows', []) for v in row)
        for line in chapter.get('body_markdown', chapter.get('body', '')).splitlines():
            line = re.sub(r'!\[[^]]*\]\([^)]*\)', '', line)
            line = re.sub(r'^\s{0,3}#{1,6}\s*', '', line).replace('**', '')
            if re.fullmatch(r'[| :\-]+', line or ' '):
                continue
            line = line.replace('|', '').strip()
            if line and _normal(line) not in _normal(content + section.get('title', '')):
                findings.append('导出规格漏当前正文：' + chapter['section_id'] + '：' + line[:55])
    return findings


def inspect(spec_path: Path, docx_path: Path, pdf_path: Path, project_root: Path) -> dict[str, Any]:
    """Return a machine receipt; ``passed`` is false for every failed gate."""
    spec_path, docx_path, pdf_path, project_root = map(
        Path, (spec_path, docx_path, pdf_path, project_root)
    )
    findings: list[str] = []
    try:
        project_id = _project_id(project_root)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        project_id = None
        findings = [f"项目身份文件无效：{exc}"]
    if not project_id:
        findings.append("缺少work/project.json中的project_id，无法绑定项目身份")
    try:
        spec = _read_json(spec_path)
        expected = _expected(spec)
        expected["image_bindings"] = []
        for relative in expected["image_sources"]:
            path = (project_root / relative).resolve()
            if not path.is_relative_to(project_root.resolve()) or not path.is_file():
                raise ValueError("图片来源缺失或路径越界：" + str(relative))
            expected["image_bindings"].append({"path": relative, "sha256": _sha256(path),
                                                "pixels": _pixels(path.read_bytes())})
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        return {
            "schema_version": SCHEMA_VERSION,
            "project_id": project_id,
            "input_spec_path": _relative(spec_path, project_root),
            "input_spec_sha256": _sha256(spec_path) if spec_path.is_file() else None,
            "artifacts": {},
            "expected": {},
            "counts": {},
            "coverage": {},
            "findings": [f"document-spec无效：{exc}"],
            "mechanical_check": "failed",
            "semantic_check": "NOT_RUN",
            "semantic_review": "NOT_RUN",
            "artifact_status": "candidate_only",
            "release_authorization": "NOT_RUN",
            "visual_qa": "NOT_RUN",
            "passed": False,
            "tooling": {
                "checker": "export_checks.py",
                "engine": "python-docx + PyMuPDF",
                "versions": _tool_versions(),
            },
        }
    try:
        findings.extend(_check_current_chapters(spec, project_root))
    except (OSError, ValueError, KeyError, TypeError, AttributeError) as exc:
        findings.append('当前导出源产物结构无效：' + str(exc))
    spec_project_id = spec.get("project_id")
    if spec_project_id is not None and spec_project_id != project_id:
        findings.append("document-spec的project_id与项目身份不一致")

    artifacts: dict[str, Any] = {}
    snapshots: dict[str, Any] = {}
    if not docx_path.is_file():
        findings.append("DOCX文件缺失")
    else:
        artifacts["docx"] = _artifact_meta(docx_path, project_root)
        try:
            snapshots["docx"], docx_findings = _docx_snapshot(docx_path, expected)
            findings.extend(docx_findings)
        except Exception as exc:
            findings.append(f"DOCX格式无效或无法读取：{exc}")
    if not pdf_path.is_file():
        findings.append("PDF文件缺失")
    else:
        try:
            import fitz

            with fitz.open(pdf_path) as pdf:
                page_count = len(pdf)
            artifacts["pdf"] = _artifact_meta(pdf_path, project_root, page_count)
            snapshots["pdf"], pdf_findings = _pdf_snapshot(pdf_path, expected)
            findings.extend(pdf_findings)
        except Exception as exc:
            findings.append(f"PDF格式无效或无法读取：{exc}")

    # Reuse the existing generic artifact scanner so its anomaly rules remain shared.
    if inspect_artifact is not None:
        for name, path in (("docx", docx_path), ("pdf", pdf_path)):
            if path.is_file():
                try:
                    baseline = inspect_artifact(path)
                    findings.extend(f"{name}: {note}" for note in baseline.get("findings", []))
                except Exception as exc:
                    findings.append(f"{name}基础工件检查失败：{exc}")

    coverage = {name: item.get("coverage", {}) for name, item in snapshots.items()}
    counts = {name: item.get("counts", {}) for name, item in snapshots.items()}
    unique_findings = sorted(set(findings))
    passed = not unique_findings and set(artifacts) == {"docx", "pdf"}
    return {
        "schema_version": SCHEMA_VERSION,
        "project_id": project_id,
        "input_spec_path": _relative(spec_path, project_root),
            "input_spec_sha256": _sha256(spec_path),
        "artifacts": artifacts,
        "expected": expected,
        "counts": counts,
        "coverage": coverage,
        "findings": unique_findings,
        "mechanical_check": "passed" if passed else "failed",
        "semantic_check": "NOT_RUN",
        "semantic_review": "NOT_RUN",
        "artifact_status": "candidate_only",
        "release_authorization": "NOT_RUN",
        "visual_qa": "NOT_RUN",
        "passed": passed,
        "tooling": {
            "checker": "export_checks.py",
            "engine": "python-docx + PyMuPDF",
            "versions": _tool_versions(),
        },
    }


def verify_receipt(
    receipt: dict[str, Any], spec_path: Path, docx_path: Path, pdf_path: Path, project_root: Path
) -> dict[str, Any]:
    """Re-read all inputs and reject a receipt whose bytes or spec have drifted."""
    current = inspect(spec_path, docx_path, pdf_path, project_root)
    if not current.get("passed"):
        raise ValueError(
            "导出收据对应的当前文件未通过检查：" + "；".join(current.get("findings", []))
        )
    for field in ("schema_version", "project_id", "input_spec_path", "input_spec_sha256"):
        if receipt.get(field) != current.get(field):
            raise ValueError(f"导出收据{field}已漂移")
    if receipt.get("expected") != current.get("expected"):
        raise ValueError("导出收据的document-spec期望内容已漂移")
    for field in (
        "counts",
        "coverage",
        "findings",
        "mechanical_check",
        "semantic_check",
        "semantic_review",
        "artifact_status",
        "release_authorization",
        "visual_qa",
        "tooling",
    ):
        if receipt.get(field) != current.get(field):
            raise ValueError(f"导出收据{field}已漂移")
    for name in ("docx", "pdf"):
        saved = receipt.get("artifacts", {}).get(name, {})
        actual = current["artifacts"][name]
        for field in ("path", "sha256", "bytes", "pages"):
            if saved.get(field) != actual.get(field):
                raise ValueError(f"导出收据{ name }的{ field }已漂移")
    return {"valid": True, "receipt": current}


def _write_receipt(path: Path, value: dict[str, Any]) -> None:
    if path.exists():
        raise ValueError("拒绝覆盖已有收据，请使用新版本文件名")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--spec", type=Path, required=True)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--receipt", type=Path)
    args = parser.parse_args()
    result = inspect(args.spec, args.docx, args.pdf, args.project_root)
    if args.receipt:
        _write_receipt(args.receipt, result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
