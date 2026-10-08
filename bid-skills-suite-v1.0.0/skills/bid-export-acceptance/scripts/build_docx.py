#!/usr/bin/env python3
"""Build a conservative editable working draft from a document specification.

Not a universal template-preserving engine. Fixed forms/complex attachments require host tools.
"""
from __future__ import annotations
import argparse, json, math, sys
from pathlib import Path


_EXPORT_FIELDS = {'style', 'cover', 'toc', 'header', 'footer'}
_STYLE_FIELDS = {
    'body_size_pt', 'body_line_spacing', 'body_east_asia_font',
    'body_latin_font', 'body_first_line_indent_chars',
    'body_space_before_pt', 'body_space_after_pt', 'body_alignment',
    'heading_sizes_pt', 'heading_east_asia_font', 'title_size_pt',
    'subtitle_size_pt', 'heading_space_before_pt', 'heading_space_after_pt',
    'table_size_pt', 'caption_size_pt', 'page_margins_cm',
}
_COVER_FIELDS = {'enabled', 'title', 'subtitle', 'metadata_rows', 'alignment', 'own_page', 'image', 'typography'}
_COVER_TYPOGRAPHY_FIELDS = {
    'title_east_asia_font', 'subtitle_east_asia_font', 'metadata_east_asia_font',
    'title_size_pt', 'subtitle_size_pt', 'metadata_size_pt',
}
_TOC_FIELDS = {'enabled', 'title', 'levels', 'typography'}
_TOC_TYPOGRAPHY_FIELDS = {
    'east_asia_font', 'latin_font', 'size_pt', 'line_spacing',
    'space_before_pt', 'space_after_pt', 'alignment',
}
_HEADER_FOOTER_FIELDS = {
    'text', 'alignment', 'page_field', 'page_prefix', 'page_suffix',
}


def _reject_duplicate_json_keys(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f'JSON对象包含重复字段：{key}')
        result[key] = value
    return result


def _read_json(path: Path):
    try:
        return json.loads(path.read_text(encoding='utf-8'), object_pairs_hook=_reject_duplicate_json_keys)
    except json.JSONDecodeError as exc:
        raise ValueError(f'JSON格式无效：{path}') from exc


def _check_fields(value, allowed, label):
    if not isinstance(value, dict):
        raise ValueError(f'{label}必须是对象')
    unknown = set(value) - allowed
    if unknown:
        raise ValueError(f'{label}存在不支持的字段：{sorted(unknown)}')


def _check_text(value, label, *, allow_empty=True):
    if not isinstance(value, str) or (not allow_empty and not value.strip()):
        raise ValueError(f'{label}必须是文本')
    if len(value) > 500:
        raise ValueError(f'{label}过长')
    return value


def _check_bool(value, label):
    if not isinstance(value, bool):
        raise ValueError(f'{label}必须是布尔值')
    return value


def _check_number(value, label, limit, *, allow_zero=False):
    if (isinstance(value, bool) or not isinstance(value, (int, float))
            or not math.isfinite(value) or value > limit
            or (value < 0 if allow_zero else value <= 0)):
        raise ValueError(f'{label}无效')
    return value


def _alignment(value, label):
    if value not in ('left', 'center', 'right', 'justify'):
        raise ValueError(f'{label}须为left、center、right或justify')
    return value


def _validate_cover_image(image, asset_root: Path, content_width: float, usable_height: float):
    _check_fields(image, {'path', 'width_cm', 'height_cm'}, 'cover.image')
    path_value = image.get('path')
    if not isinstance(path_value, str) or not path_value.strip():
        raise ValueError('cover.image.path必须是非空相对路径')
    relative = Path(path_value)
    root = asset_root.resolve()
    path = (root / relative).resolve()
    if relative.is_absolute() or '..' in relative.parts or not path.is_relative_to(root):
        raise ValueError('cover.image路径越界')
    if not path.is_file():
        raise ValueError('cover.image图片缺失：' + str(relative))
    if path.suffix.lower() not in {'.png', '.jpg', '.jpeg'}:
        raise ValueError('此基础生成器仅直接插入PNG/JPEG，SVG须先受控渲染')
    width = _check_number(image.get('width_cm', content_width), 'cover.image.width_cm', content_width)
    height = image.get('height_cm')
    if height is not None:
        _check_number(height, 'cover.image.height_cm', usable_height)
    return path, width, height


def _validate_export_settings(raw, *, spec_title, spec_subtitle, style, asset_root, content_width):
    """Normalize and validate the optional export-only configuration."""
    _check_fields(raw, _EXPORT_FIELDS, 'export_settings')
    body_font = style.get('body_east_asia_font', 'SimSun')
    latin_font = style.get('body_latin_font', 'Arial')
    heading_font = style.get('heading_east_asia_font', 'SimHei')
    body_size = style.get('body_size_pt', 11)
    spacing = style.get('body_line_spacing', 1.25)
    before = style.get('body_space_before_pt', 0)
    after = style.get('body_space_after_pt', 7)
    margins = style.get('page_margins_cm', {'top': 2.3, 'bottom': 2.3, 'left': 2.5, 'right': 2.5})
    usable_height = 29.7 - margins['top'] - margins['bottom']
    if usable_height <= 0:
        raise ValueError('页边距使可用页面高度无效')

    cover_value = raw.get('cover', {})
    _check_fields(cover_value, _COVER_FIELDS, 'export_settings.cover')
    cover = dict(cover_value)
    cover_typography_value = cover.get('typography', {})
    _check_fields(cover_typography_value, _COVER_TYPOGRAPHY_FIELDS, 'export_settings.cover.typography')
    cover_typography = dict(cover_typography_value)
    cover_title = cover.get('title', spec_title)
    cover_subtitle = cover.get('subtitle', spec_subtitle or '')
    _check_text(cover_title, 'cover.title', allow_empty=False)
    _check_text(cover_subtitle, 'cover.subtitle')
    enabled = cover.get('enabled', True)
    own_page = cover.get('own_page', False)
    _check_bool(enabled, 'cover.enabled')
    _check_bool(own_page, 'cover.own_page')
    cover_alignment = _alignment(cover.get('alignment', 'center'), 'cover.alignment')
    rows = cover.get('metadata_rows', [])
    if not isinstance(rows, list):
        raise ValueError('cover.metadata_rows必须是数组')
    metadata = []
    labels = set()
    for index, row in enumerate(rows):
        _check_fields(row, {'label', 'value'}, f'cover.metadata_rows[{index}]')
        label = _check_text(row.get('label'), f'cover.metadata_rows[{index}].label', allow_empty=False)
        value = _check_text(row.get('value'), f'cover.metadata_rows[{index}].value')
        if label in labels:
            raise ValueError('cover.metadata_rows不允许重复label：' + label)
        labels.add(label)
        metadata.append((label, value))
    cover_font_defaults = {
        'title_east_asia_font': heading_font,
        'subtitle_east_asia_font': heading_font,
        'metadata_east_asia_font': body_font,
        'title_size_pt': style.get('title_size_pt', 22),
        'subtitle_size_pt': style.get('subtitle_size_pt', 13),
        'metadata_size_pt': body_size,
    }
    cover_font = {**cover_font_defaults, **cover_typography}
    for key in ('title_east_asia_font', 'subtitle_east_asia_font', 'metadata_east_asia_font'):
        _check_text(cover_font[key], f'cover.typography.{key}', allow_empty=False)
    for key in ('title_size_pt', 'subtitle_size_pt', 'metadata_size_pt'):
        _check_number(cover_font[key], f'cover.typography.{key}', 72)
    cover_image = None
    if 'image' in cover:
        if cover['image'] is None:
            raise ValueError('cover.image不能为null；不需要图片时请删除该字段')
        cover_image = _validate_cover_image(cover['image'], asset_root, content_width, usable_height)

    toc_value = raw.get('toc', {})
    _check_fields(toc_value, _TOC_FIELDS, 'export_settings.toc')
    toc = dict(toc_value)
    toc_enabled = toc.get('enabled', False)
    _check_bool(toc_enabled, 'toc.enabled')
    toc_title = _check_text(toc.get('title', '目录'), 'toc.title', allow_empty=False)
    levels = toc.get('levels', 3)
    if isinstance(levels, bool) or not isinstance(levels, int) or not 1 <= levels <= 9:
        raise ValueError('toc.levels必须是1到9的整数')
    toc_typography = dict(toc.get('typography', {}))
    _check_fields(toc_typography, _TOC_TYPOGRAPHY_FIELDS, 'export_settings.toc.typography')
    toc_font = {
        'east_asia_font': body_font,
        'latin_font': latin_font,
        'size_pt': body_size,
        'line_spacing': spacing,
        'space_before_pt': before,
        'space_after_pt': after,
        'alignment': 'left',
        **toc_typography,
    }
    for key in ('east_asia_font', 'latin_font'):
        _check_text(toc_font[key], f'toc.typography.{key}', allow_empty=False)
    _check_number(toc_font['size_pt'], 'toc.typography.size_pt', 72)
    _check_number(toc_font['line_spacing'], 'toc.typography.line_spacing', 5)
    _check_number(toc_font['space_before_pt'], 'toc.typography.space_before_pt', 72, allow_zero=True)
    _check_number(toc_font['space_after_pt'], 'toc.typography.space_after_pt', 72, allow_zero=True)
    _alignment(toc_font['alignment'], 'toc.typography.alignment')

    def validate_header_footer(name, default):
        value_raw = raw.get(name, {})
        _check_fields(value_raw, _HEADER_FOOTER_FIELDS, f'export_settings.{name}')
        value = dict(value_raw)
        text = _check_text(value.get('text', default), f'{name}.text')
        alignment = _alignment(value.get('alignment', 'left' if name == 'header' else 'right'), f'{name}.alignment')
        page_field = value.get('page_field', name == 'footer')
        _check_bool(page_field, f'{name}.page_field')
        prefix = _check_text(value.get('page_prefix', '第 '), f'{name}.page_prefix')
        suffix = _check_text(value.get('page_suffix', ' 页'), f'{name}.page_suffix')
        return {'text': text, 'alignment': alignment, 'page_field': page_field,
                'page_prefix': prefix, 'page_suffix': suffix}

    return {
        'cover': {'enabled': enabled, 'title': cover_title, 'subtitle': cover_subtitle,
                  'metadata': metadata, 'alignment': cover_alignment, 'own_page': own_page,
                  'image': cover_image, 'typography': cover_font},
        'toc': {'enabled': toc_enabled, 'title': toc_title, 'levels': levels,
                'typography': toc_font},
        'header': validate_header_footer('header', '内部工作稿｜需人工复核与签署'),
        'footer': validate_header_footer('footer', ''),
    }


def _resolve_export_settings(spec, export_settings_path, *, style, asset_root, content_width):
    raw = _read_export_settings(spec, export_settings_path)
    return _validate_export_settings(raw, spec_title=spec.get('title'),
                                     spec_subtitle=spec.get('subtitle'), style=style,
                                     asset_root=asset_root, content_width=content_width)


def _read_export_settings(spec, export_settings_path):
    if 'export_settings' in spec and export_settings_path is not None:
        raise ValueError('不能同时提供外部export-settings和spec内嵌导出配置')
    if export_settings_path is not None:
        path = Path(export_settings_path)
        if not path.is_file():
            raise ValueError('导出配置文件不存在：' + str(path))
        raw = _read_json(path)
    else:
        raw = spec.get('export_settings', {})
    _check_fields(raw, _EXPORT_FIELDS, 'export_settings')
    return raw


def _validate_style(raw):
    if not isinstance(raw, dict):
        raise ValueError('style必须是对象')
    _check_fields(raw, _STYLE_FIELDS, 'style')
    body_size = raw.get('body_size_pt', 11)
    spacing = raw.get('body_line_spacing', 1.25)

    def size(value, limit, allow_zero=False):
        if (isinstance(value, bool) or not isinstance(value, (int, float))
                or not math.isfinite(value) or value > limit
                or (value < 0 if allow_zero else value <= 0)):
            raise ValueError('字号、间距、缩进或页边距无效')
        return value

    size(body_size, 72)
    size(spacing, 5)
    font = raw.get('body_east_asia_font', 'SimSun')
    latin_font = raw.get('body_latin_font', 'Arial')
    heading_font = raw.get('heading_east_asia_font', 'SimHei')
    for family in (font, latin_font, heading_font):
        if not isinstance(family, str) or not family.strip():
            raise ValueError('字体名称无效')
    indent = size(raw.get('body_first_line_indent_chars', 0), 4, True)
    before = size(raw.get('body_space_before_pt', 0), 72, True)
    after = size(raw.get('body_space_after_pt', 7), 72, True)
    alignment = raw.get('body_alignment', 'left')
    if alignment not in ('left', 'justify'):
        raise ValueError('正文对齐须为left或justify')
    heading_before = size(raw.get('heading_space_before_pt', 0), 72, True)
    heading_after = size(raw.get('heading_space_after_pt', 7), 72, True)
    heading_sizes = raw.get('heading_sizes_pt', [13, 13, 13])
    if not isinstance(heading_sizes, list) or len(heading_sizes) != 3:
        raise ValueError('heading_sizes_pt须为三级标题的三个字号')
    for value in heading_sizes:
        size(value, 72)
    title_size = size(raw.get('title_size_pt', 22), 72)
    subtitle_size = size(raw.get('subtitle_size_pt', 13), 72)
    table_size = size(raw.get('table_size_pt', body_size), 72)
    caption_size = size(raw.get('caption_size_pt', body_size), 72)
    margins = raw.get('page_margins_cm', {'top': 2.3, 'bottom': 2.3, 'left': 2.5, 'right': 2.5})
    if not isinstance(margins, dict) or set(margins) != {'top', 'bottom', 'left', 'right'}:
        raise ValueError('page_margins_cm须含且仅含top/bottom/left/right')
    for value in margins.values():
        size(value, 5)
    content_width = 21 - margins['left'] - margins['right']
    if content_width <= 0:
        raise ValueError('页边距使正文宽度无效')
    if 29.7 - margins['top'] - margins['bottom'] <= 0:
        raise ValueError('页边距使可用页面高度无效')
    return {
        'body_size_pt': body_size, 'body_line_spacing': spacing,
        'body_east_asia_font': font, 'body_latin_font': latin_font,
        'body_first_line_indent_chars': indent,
        'body_space_before_pt': before, 'body_space_after_pt': after,
        'body_alignment': alignment, 'heading_sizes_pt': heading_sizes,
        'heading_east_asia_font': heading_font,
        'title_size_pt': title_size, 'subtitle_size_pt': subtitle_size,
        'heading_space_before_pt': heading_before,
        'heading_space_after_pt': heading_after, 'table_size_pt': table_size,
        'caption_size_pt': caption_size, 'page_margins_cm': margins,
    }, content_width


def _set_run_font(run, *, east_asia_font, latin_font, size_pt):
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run.font.name = latin_font
    run.font.size = Pt(size_pt)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn('w:ascii'), latin_font)
    rfonts.set(qn('w:hAnsi'), latin_font)
    rfonts.set(qn('w:cs'), latin_font)
    rfonts.set(qn('w:eastAsia'), east_asia_font)
    _set_black_rpr(rpr)
    complex_size = rpr.find(qn('w:szCs'))
    if complex_size is None:
        complex_size = OxmlElement('w:szCs')
        rpr.append(complex_size)
    complex_size.set(qn('w:val'), str(round(size_pt * 2)))


def _set_black_rpr(rpr):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    color = rpr.find(qn('w:color'))
    if color is None:
        color = OxmlElement('w:color')
        rpr.append(color)
    color.set(qn('w:val'), '000000')
    for attribute in (qn('w:themeColor'), qn('w:themeTint'), qn('w:themeShade')):
        color.attrib.pop(attribute, None)


def _set_paragraph_alignment(paragraph, value):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph.alignment = {
        'left': WD_ALIGN_PARAGRAPH.LEFT,
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
        'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[value]


def _insert_bookmark(paragraph, bookmark_id, name):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bookmark_id))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bookmark_id))
    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def _configure_toc_styles(document, toc_settings):
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    typography = toc_settings['typography']
    for level in range(1, 10):
        # Word matches its built-in TOC styles by this exact name/identity.
        # A custom "TOC 1" is renamed on import and ignored when fields refresh.
        name = f'toc {level}'
        try:
            toc_style = document.styles[name]
        except KeyError:
            toc_style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
        toc_style.element.set(qn('w:styleId'), f'TOC{level}')
        toc_style.element.attrib.pop(qn('w:customStyle'), None)
        toc_style.base_style = document.styles['Normal']
        toc_style.font.name = typography['latin_font']
        toc_style.font.size = Pt(typography['size_pt'])
        rpr = toc_style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        for script in ('ascii', 'hAnsi', 'cs'):
            rfonts.set(qn(f'w:{script}'), typography['latin_font'])
        rfonts.set(qn('w:eastAsia'), typography['east_asia_font'])
        _set_black_rpr(rpr)
        complex_size = rpr.find(qn('w:szCs'))
        if complex_size is None:
            complex_size = OxmlElement('w:szCs')
            rpr.append(complex_size)
        complex_size.set(qn('w:val'), str(round(typography['size_pt'] * 2)))
        toc_style.paragraph_format.space_before = Pt(typography['space_before_pt'])
        toc_style.paragraph_format.space_after = Pt(typography['space_after_pt'])
        toc_style.paragraph_format.line_spacing = typography['line_spacing']
        toc_style.paragraph_format.left_indent = Cm((level - 1) * 0.6)
        toc_style.paragraph_format.first_line_indent = Pt(0)
        _set_paragraph_alignment(toc_style.paragraph_format, typography['alignment'])


def _add_toc_field(document, toc_settings, style):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    _configure_toc_styles(document, toc_settings)
    title = document.add_paragraph(toc_settings['title'], 'Title')
    title.paragraph_format.first_line_indent = Pt(0)
    _set_paragraph_alignment(title, 'center')
    field_paragraph = document.add_paragraph()
    typography = toc_settings['typography']
    field_paragraph.paragraph_format.first_line_indent = Pt(0)
    field_paragraph.paragraph_format.space_before = Pt(typography['space_before_pt'])
    field_paragraph.paragraph_format.space_after = Pt(typography['space_after_pt'])
    field_paragraph.paragraph_format.line_spacing = typography['line_spacing']
    _set_paragraph_alignment(field_paragraph, typography['alignment'])
    def wrapped(element):
        run = OxmlElement('w:r')
        run.append(element)
        return run

    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    begin.set(qn('w:dirty'), 'true')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = f' TOC \\o "1-{toc_settings["levels"]}" \\h \\z \\u '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    field_paragraph._p.append(wrapped(begin))
    field_paragraph._p.append(wrapped(instruction))
    field_paragraph._p.append(wrapped(separate))
    placeholder = field_paragraph.add_run('目录字段将在支持的Word交付引擎中更新；当前工作稿未刷新页码。')
    _set_run_font(placeholder, east_asia_font=typography['east_asia_font'],
                  latin_font=typography['latin_font'], size_pt=typography['size_pt'])
    field_paragraph._p.append(wrapped(end))


def _configure_header_footer(section, export_settings, style):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for name in ('header', 'footer'):
        settings = export_settings[name]
        container = getattr(section, name)
        paragraph = container.paragraphs[0]
        paragraph.text = ''
        _set_paragraph_alignment(paragraph, settings['alignment'])
        if settings['text']:
            run = paragraph.add_run(settings['text'])
            _set_run_font(run, east_asia_font=style.get('body_east_asia_font', 'SimSun'),
                          latin_font=style.get('body_latin_font', 'Arial'),
                          size_pt=style.get('body_size_pt', 11))
        if settings['page_field']:
            prefix = paragraph.add_run(settings['page_prefix'])
            _set_run_font(prefix, east_asia_font=style.get('body_east_asia_font', 'SimSun'),
                          latin_font=style.get('body_latin_font', 'Arial'),
                          size_pt=style.get('body_size_pt', 11))
            field = OxmlElement('w:fldSimple')
            field.set(qn('w:instr'), 'PAGE')
            paragraph._p.append(field)
            suffix = paragraph.add_run(settings['page_suffix'])
            _set_run_font(suffix, east_asia_font=style.get('body_east_asia_font', 'SimSun'),
                          latin_font=style.get('body_latin_font', 'Arial'),
                          size_pt=style.get('body_size_pt', 11))

def build(spec_path: Path, out_path: Path, asset_root: Path, export_settings_path: Path | None = None):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    spec=_read_json(spec_path)
    if out_path.exists():raise ValueError('拒绝覆盖已有文档，请使用新版本文件名')
    raw_export_settings = _read_export_settings(spec, export_settings_path)
    base_style = spec.get('style', {})
    if not isinstance(base_style, dict):
        raise ValueError('style必须是对象')
    export_style = raw_export_settings.get('style', {})
    if not isinstance(export_style, dict):
        raise ValueError('export_settings.style必须是对象')
    merged_style = {**base_style, **export_style}
    typography, content_width = _validate_style(merged_style)
    body_size = typography['body_size_pt']
    spacing = typography['body_line_spacing']
    font = typography['body_east_asia_font']
    latin_font = typography['body_latin_font']
    heading_font = typography['heading_east_asia_font']
    indent = typography['body_first_line_indent_chars']
    before = typography['body_space_before_pt']
    after = typography['body_space_after_pt']
    alignment = typography['body_alignment']
    heading_before = typography['heading_space_before_pt']
    heading_after = typography['heading_space_after_pt']
    heading_sizes = typography['heading_sizes_pt']
    title_size = typography['title_size_pt']
    subtitle_size = typography['subtitle_size_pt']
    table_size = typography['table_size_pt']
    caption_size = typography['caption_size_pt']
    margins = typography['page_margins_cm']
    usable_height = 29.7 - margins['top'] - margins['bottom']
    asset_root = Path(asset_root)
    export_settings = _resolve_export_settings(spec, export_settings_path, style=typography,
                                                asset_root=asset_root, content_width=content_width)
    document=Document();section=document.sections[0]
    # Explicit template fonts must not compete with inherited theme fonts.
    for fonts in document.styles.element.iter(qn('w:rFonts')):
        for attribute in list(fonts.attrib):
            if attribute.lower().endswith('theme'):del fonts.attrib[attribute]
    defaults=document.styles.element.find('.//' + qn('w:docDefaults') +
                                          '/' + qn('w:rPrDefault') + '/' +
                                          qn('w:rPr') + '/' + qn('w:rFonts'))
    if defaults is not None:
        for script in ('ascii','hAnsi','cs'):defaults.set(qn('w:'+script),latin_font)
        defaults.set(qn('w:eastAsia'),font)
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    for edge,value in margins.items():setattr(section,edge+'_margin',Cm(value))
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        font_size=(body_size if name=='Normal' else title_size if name=='Title' else
                   subtitle_size if name=='Subtitle' else heading_sizes[int(name[-1])-1])
        style=document.styles[name];style.font.name=latin_font;style.font.size=Pt(font_size)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),font if name=='Normal' else heading_font)
        style.element.rPr.rFonts.set(qn('w:cs'),latin_font)
        _set_black_rpr(style.element.rPr)
        complex_size=style.element.rPr.find(qn('w:szCs'))
        if complex_size is None:
            complex_size=OxmlElement('w:szCs');style.element.rPr.append(complex_size)
        complex_size.set(qn('w:val'),str(round(font_size*2)))
        style.paragraph_format.space_before=Pt(before if name=='Normal' else heading_before)
        style.paragraph_format.space_after=Pt(after if name=='Normal' else heading_after)
        style.paragraph_format.line_spacing=spacing
        if name in {'Title', 'Subtitle'}:
            paragraph_properties = style.element.find(qn('w:pPr'))
            if paragraph_properties is not None:
                for border in list(paragraph_properties.findall(qn('w:pBdr'))):
                    paragraph_properties.remove(border)
    document.core_properties.title=spec['title'];document.core_properties.author='';document.core_properties.last_modified_by=''
    _configure_header_footer(section, export_settings, typography)
    cover = export_settings['cover']
    if cover['enabled']:
        cover_title = document.add_heading(cover['title'], 0)
        _set_paragraph_alignment(cover_title, cover['alignment'])
        for run in cover_title.runs:
            _set_run_font(run, east_asia_font=cover['typography']['title_east_asia_font'],
                          latin_font=typography.get('body_latin_font', 'Arial'),
                          size_pt=cover['typography']['title_size_pt'])
        if cover['subtitle']:
            subtitle = document.add_paragraph(cover['subtitle'], 'Subtitle')
            _set_paragraph_alignment(subtitle, cover['alignment'])
            subtitle.paragraph_format.first_line_indent = Pt(0)
            for run in subtitle.runs:
                _set_run_font(run, east_asia_font=cover['typography']['subtitle_east_asia_font'],
                              latin_font=typography.get('body_latin_font', 'Arial'),
                              size_pt=cover['typography']['subtitle_size_pt'])
        if cover['metadata']:
            metadata_table = document.add_table(rows=0, cols=2)
            metadata_table.style = 'Table Grid'
            metadata_table.autofit = False
            metadata_table.columns[0].width = Cm(min(4, content_width / 3))
            metadata_table.columns[1].width = Cm(content_width - min(4, content_width / 3))
            for label, value in cover['metadata']:
                cells = metadata_table.add_row().cells
                cells[0].text = label
                cells[1].text = value
                for cell in cells:
                    cell.width = metadata_table.columns[0].width if cell is cells[0] else metadata_table.columns[1].width
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.first_line_indent = Pt(0)
                        _set_paragraph_alignment(paragraph, cover['alignment'])
                        for run in paragraph.runs:
                            _set_run_font(run, east_asia_font=cover['typography']['metadata_east_asia_font'],
                                          latin_font=typography.get('body_latin_font', 'Arial'),
                                          size_pt=cover['typography']['metadata_size_pt'])
        if cover['image'] is not None:
            image_path, image_width, image_height = cover['image']
            kwargs = {'width': Cm(image_width)}
            if image_height is not None:
                kwargs['height'] = Cm(image_height)
            document.add_picture(str(image_path), **kwargs)
            image_paragraph = document.paragraphs[-1]
            _set_paragraph_alignment(image_paragraph, cover['alignment'])
        notice = document.add_paragraph('本文件由基础排版脚本生成，不代表真实投标通过审核。客户指定版式、目录域与签署要求需在正式交付前另行落实。')
        notice.paragraph_format.first_line_indent = Pt(0)
        _set_paragraph_alignment(notice, cover['alignment'])
        if cover['own_page']:
            document.add_page_break()
    if export_settings['toc']['enabled']:
        if cover['enabled'] and not cover['own_page']:
            document.add_page_break()
        _add_toc_field(document, export_settings['toc'], typography)
        document.add_page_break()
    elif cover['enabled'] and cover['own_page']:
        # The cover's explicit page break already starts the body on a new page.
        pass
    bookmark_id = 1
    for item in spec['sections']:
        level=max(1,min(int(item.get('level',1)),3))
        heading = document.add_heading(item['title'], level=level)
        if export_settings['toc']['enabled']:
            _insert_bookmark(heading, bookmark_id, f'_Toc{bookmark_id}')
            bookmark_id += 1
        for text in item.get('paragraphs',[]):
            if not str(text).strip():continue
            paragraph=document.add_paragraph(str(text))
            paragraph.paragraph_format.first_line_indent=Pt(body_size*indent)
            paragraph.paragraph_format.alignment=(WD_ALIGN_PARAGRAPH.JUSTIFY
                if alignment=='justify' else WD_ALIGN_PARAGRAPH.LEFT)
        for table_spec in item.get('tables',[]):
            headers=table_spec['headers'];rows=table_spec['rows']
            if not headers or any(len(row)!=len(headers) for row in rows):raise ValueError('表格列数不一致')
            widths = table_spec.get('widths_cm')
            if widths is not None:
                if (not isinstance(widths, list) or len(widths) != len(headers)
                        or any(isinstance(w, bool) or not isinstance(w, (int, float))
                               or not math.isfinite(w) or w <= 0 for w in widths)
                        or sum(widths) > content_width):
                    raise ValueError('表格列宽必须为正数，与列数一致且不超过当前页面正文宽度')
            table=document.add_table(rows=1,cols=len(headers));table.style='Table Grid'
            if widths is not None:
                table.autofit = False
                for column, width in zip(table.columns, widths):column.width = Cm(width)
            trPr=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
            for i,v in enumerate(headers):table.rows[0].cells[i].text=str(v)
            for values in rows:
                cells=table.add_row().cells
                for i,v in enumerate(values):cells[i].text=str(v)
            for row_index, row in enumerate(table.rows):
                row_tr_pr = row._tr.get_or_add_trPr()
                row_tr_pr.append(OxmlElement('w:cantSplit'))
                if row_index == 0:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.keep_with_next = True
            if widths is not None:
                for row in table.rows:
                    for cell, width in zip(row.cells, widths):cell.width = Cm(width)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.first_line_indent=Pt(0)
                        for run in paragraph.runs:run.font.size=Pt(table_size)
        for image in item.get('images',[]):
            if not isinstance(image, dict):
                raise ValueError('图片配置必须是对象')
            rel=Path(image.get('path', ''))
            image_path, image_width, image_height = _validate_cover_image(
                {'path': str(rel), **{key: image[key] for key in ('width_cm', 'height_cm') if key in image}},
                asset_root, content_width, usable_height)
            kwargs = {'width': Cm(image_width)}
            if image_height is not None:
                kwargs['height'] = Cm(image_height)
            document.add_picture(str(image_path), **kwargs)
            caption=document.add_paragraph(image.get('caption',''))
            caption.paragraph_format.first_line_indent=Pt(0)
            for run in caption.runs:run.font.size=Pt(caption_size)
    out_path.parent.mkdir(parents=True,exist_ok=True);document.save(out_path)
    return {'output':str(out_path),'status':'draft_only','visual_qa':'NOT_RUN','notes':['字体文件未随包分发，请在运行环境配置合法字体。','不自动更新目录，不代替复杂原始模板填报和证明材料插页。']}

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--spec',type=Path,required=True);p.add_argument('--out',type=Path,required=True);p.add_argument('--asset-root',type=Path);p.add_argument('--export-settings',type=Path,help='可选项目导出配置JSON，亦可内嵌到spec.export_settings');a=p.parse_args()
    try:print(json.dumps(build(a.spec,a.out,a.asset_root or a.spec.parent, a.export_settings),ensure_ascii=False,indent=2));return 0
    except (ImportError,OSError,ValueError,KeyError) as e:print('DOCX生成失败：'+str(e),file=sys.stderr);return 2
if __name__=='__main__':raise SystemExit(main())
