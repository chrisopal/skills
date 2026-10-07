#!/usr/bin/env python3
"""Build a conservative editable working draft from a document specification.

Not a universal template-preserving engine. Fixed forms/complex attachments require host tools.
"""
from __future__ import annotations
import argparse, json, math, sys
from pathlib import Path

def build(spec_path: Path,out_path: Path,asset_root: Path):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    spec=json.loads(spec_path.read_text(encoding='utf-8'))
    if out_path.exists():raise ValueError('拒绝覆盖已有文档，请使用新版本文件名')
    typography = spec.get('style', {})
    if not isinstance(typography, dict):raise ValueError('style必须是对象')
    if set(typography)-{'body_size_pt','body_line_spacing','body_east_asia_font',
                       'body_latin_font','body_first_line_indent_chars',
                       'body_space_before_pt','body_space_after_pt','body_alignment','heading_sizes_pt',
                       'heading_east_asia_font','title_size_pt','subtitle_size_pt',
                       'heading_space_before_pt','heading_space_after_pt',
                       'table_size_pt','caption_size_pt','page_margins_cm'}:
        raise ValueError('不支持的正文样式字段')
    body_size = typography.get('body_size_pt', 11)
    spacing = typography.get('body_line_spacing', 1.25)
    font = typography.get('body_east_asia_font', 'SimSun')
    def size(value, limit, allow_zero=False):
        if (isinstance(value, bool) or not isinstance(value, (int, float))
                or not math.isfinite(value) or value > limit
                or (value < 0 if allow_zero else value <= 0)):
            raise ValueError('字号、间距、缩进或页边距无效')
        return value
    size(body_size, 72);size(spacing, 5)
    latin_font=typography.get('body_latin_font','Arial')
    heading_font=typography.get('heading_east_asia_font','SimHei')
    for family in (font,latin_font,heading_font):
        if not isinstance(family,str) or not family.strip():raise ValueError('字体名称无效')
    indent=size(typography.get('body_first_line_indent_chars',0),4,True)
    before=size(typography.get('body_space_before_pt',0),72,True)
    after=size(typography.get('body_space_after_pt',7),72,True)
    alignment=typography.get('body_alignment','left')
    if alignment not in ('left','justify'):raise ValueError('正文对齐须为left或justify')
    heading_before=size(typography.get('heading_space_before_pt',0),72,True)
    heading_after=size(typography.get('heading_space_after_pt',7),72,True)
    heading_sizes=typography.get('heading_sizes_pt',[13,13,13])
    if not isinstance(heading_sizes,list) or len(heading_sizes)!=3:
        raise ValueError('heading_sizes_pt须为三级标题的三个字号')
    for value in heading_sizes:size(value,72)
    title_size=size(typography.get('title_size_pt',22),72)
    subtitle_size=size(typography.get('subtitle_size_pt',13),72)
    table_size=size(typography.get('table_size_pt',body_size),72)
    caption_size=size(typography.get('caption_size_pt',body_size),72)
    margins=typography.get('page_margins_cm',{'top':2.3,'bottom':2.3,'left':2.5,'right':2.5})
    if not isinstance(margins,dict) or set(margins)!={'top','bottom','left','right'}:
        raise ValueError('page_margins_cm须含且仅含top/bottom/left/right')
    for value in margins.values():size(value,5)
    content_width=21-margins['left']-margins['right']
    document=Document();section=document.sections[0]
    # Explicit template fonts must not compete with inherited theme fonts.
    for fonts in document.styles.element.iter(qn('w:rFonts')):
        for attribute in list(fonts.attrib):
            if attribute.lower().endswith('theme'):del fonts.attrib[attribute]
    defaults=document.styles.element.find('.//' + qn('w:docDefaults') +
                                          '/' + qn('w:rPrDefault') + '/' +
                                          qn('w:rPr') + '/' + qn('w:rFonts'))
    if defaults is not None:
        for script in ('ascii','hAnsi','cs'):defaults.set(qn('w:'+script),latin_font)
        defaults.set(qn('w:eastAsia'),font)
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    for edge,value in margins.items():setattr(section,edge+'_margin',Cm(value))
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        font_size=(body_size if name=='Normal' else title_size if name=='Title' else
                   subtitle_size if name=='Subtitle' else heading_sizes[int(name[-1])-1])
        style=document.styles[name];style.font.name=latin_font;style.font.size=Pt(font_size)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),font if name=='Normal' else heading_font)
        style.element.rPr.rFonts.set(qn('w:cs'),latin_font)
        complex_size=style.element.rPr.find(qn('w:szCs'))
        if complex_size is None:
            complex_size=OxmlElement('w:szCs');style.element.rPr.append(complex_size)
        complex_size.set(qn('w:val'),str(round(font_size*2)))
        style.paragraph_format.space_before=Pt(before if name=='Normal' else heading_before)
        style.paragraph_format.space_after=Pt(after if name=='Normal' else heading_after)
        style.paragraph_format.line_spacing=spacing
    document.core_properties.title=spec['title'];document.core_properties.author='';document.core_properties.last_modified_by=''
    section.header.paragraphs[0].text='内部工作稿｜需人工复核与签署'
    footer=section.footer.paragraphs[0];footer.alignment=2
    footer.add_run('第 ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    footer.add_run(' 页')
    document.add_heading(spec['title'],0)
    if spec.get('subtitle'):document.add_paragraph(spec['subtitle'],'Subtitle')
    document.add_paragraph('本文件由基础排版脚本生成，不代表真实投标通过审核。客户指定版式、目录域与签署要求需在正式交付前另行落实。')
    for item in spec['sections']:
        document.add_heading(item['title'],level=max(1,min(int(item.get('level',1)),3)))
        for text in item.get('paragraphs',[]):
            if not str(text).strip():continue
            paragraph=document.add_paragraph(str(text))
            paragraph.paragraph_format.first_line_indent=Pt(body_size*indent)
            paragraph.paragraph_format.alignment=(WD_ALIGN_PARAGRAPH.JUSTIFY
                if alignment=='justify' else WD_ALIGN_PARAGRAPH.LEFT)
        for table_spec in item.get('tables',[]):
            headers=table_spec['headers'];rows=table_spec['rows']
            if not headers or any(len(row)!=len(headers) for row in rows):raise ValueError('表格列数不一致')
            widths = table_spec.get('widths_cm')
            if widths is not None:
                if (not isinstance(widths, list) or len(widths) != len(headers)
                        or any(isinstance(w, bool) or not isinstance(w, (int, float))
                               or not math.isfinite(w) or w <= 0 for w in widths)
                        or sum(widths) > content_width):
                    raise ValueError('表格列宽必须为正数，与列数一致且不超过当前页面正文宽度')
            table=document.add_table(rows=1,cols=len(headers));table.style='Table Grid'
            if widths is not None:
                table.autofit = False
                for column, width in zip(table.columns, widths):column.width = Cm(width)
            trPr=table.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
            for i,v in enumerate(headers):table.rows[0].cells[i].text=str(v)
            for values in rows:
                cells=table.add_row().cells
                for i,v in enumerate(values):cells[i].text=str(v)
            if widths is not None:
                for row in table.rows:
                    for cell, width in zip(row.cells, widths):cell.width = Cm(width)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.first_line_indent=Pt(0)
                        for run in paragraph.runs:run.font.size=Pt(table_size)
        for image in item.get('images',[]):
            rel=Path(image['path']);path=(asset_root/rel).resolve()
            if rel.is_absolute() or '..' in rel.parts or not path.is_relative_to(asset_root.resolve()):raise ValueError('图片路径越界')
            if not path.is_file():raise ValueError('图片缺失：'+str(rel))
            if path.suffix.lower() not in {'.png','.jpg','.jpeg'}:raise ValueError('此基础生成器仅直接插入PNG/JPEG，SVG须先受控渲染')
            document.add_picture(str(path),width=Cm(min(float(image.get('width_cm',15)),content_width)))
            caption=document.add_paragraph(image.get('caption',''))
            caption.paragraph_format.first_line_indent=Pt(0)
            for run in caption.runs:run.font.size=Pt(caption_size)
    out_path.parent.mkdir(parents=True,exist_ok=True);document.save(out_path)
    return {'output':str(out_path),'status':'draft_only','visual_qa':'NOT_RUN','notes':['字体文件未随包分发，请在运行环境配置合法字体。','不自动更新目录，不代替复杂原始模板填报和证明材料插页。']}

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--spec',type=Path,required=True);p.add_argument('--out',type=Path,required=True);p.add_argument('--asset-root',type=Path);a=p.parse_args()
    try:print(json.dumps(build(a.spec,a.out,a.asset_root or a.spec.parent),ensure_ascii=False,indent=2));return 0
    except (ImportError,OSError,ValueError,KeyError) as e:print('DOCX生成失败：'+str(e),file=sys.stderr);return 2
if __name__=='__main__':raise SystemExit(main())
