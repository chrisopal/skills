import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_SECTION

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import bidkit
import extract_sources
import review_checks


class DocxSourceTextTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.project = Path(self.tmp.name) / "project"
        bidkit.init_project(self.project, "DOCX-TEST")

        document = Document()
        document.add_paragraph("DOCX body paragraph one")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "row one col one"
        table.cell(0, 1).text = "row one col two"
        table.cell(1, 0).text = "row two col one"
        table.cell(1, 1).text = "DOCX table quote"
        document.add_paragraph("DOCX body paragraph two")
        document.sections[0].header.paragraphs[0].text = "DOCX header quote"
        document.sections[0].footer.paragraphs[0].text = "DOCX footer quote"

        source = self.project / "work" / "source.docx"
        document.save(source)
        registered = bidkit.register_source(self.project, source, "supplier")
        self.registered = registered

    def tearDown(self):
        self.tmp.cleanup()

    def source(self, location, quote, registered=None, **overrides):
        registered = registered or self.registered
        value = {
            "source_id": registered["source_id"],
            "revision": registered["revision"],
            "sha256": registered["sha256"],
            "location": location,
            "quote": quote,
        }
        value.update(overrides)
        return value

    def register_variant(self, stem, *, all_parts=False, variant=None):
        document = Document()
        section = document.sections[0]
        if all_parts or variant == "default":
            section.header.paragraphs[0].text = "variant default header"
            section.footer.paragraphs[0].text = "variant default footer"
        if all_parts or variant == "first":
            section.different_first_page_header_footer = True
            section.first_page_header.paragraphs[0].text = "variant first header"
            section.first_page_footer.paragraphs[0].text = "variant first footer"
        if all_parts or variant == "even":
            document.settings.odd_and_even_pages_header_footer = True
            section.even_page_header.paragraphs[0].text = "variant even header"
            section.even_page_footer.paragraphs[0].text = "variant even footer"
        path = self.project / "work" / f"{stem}.docx"
        document.save(path)
        return bidkit.register_source(self.project, path, "supplier")

    def test_reads_exact_body_paragraph_and_table_locations(self):
        for location, quote in (
            ("body[1]/paragraph", "DOCX body paragraph one"),
            ("body[2]/table/r2/c2", "DOCX table quote"),
            ("body[3]/paragraph", "DOCX body paragraph two"),
        ):
            with self.subTest(location=location):
                review_checks.source_text(self.project, self.source(location, quote), {})

    def test_reads_header_and_footer_locations(self):
        review_checks.source_text(
            self.project, self.source("header/p1", "DOCX header quote"), {}
        )
        review_checks.source_text(
            self.project, self.source("footer/p1", "DOCX footer quote"), {}
        )
        review_checks.source_text(
            self.project, self.source("hdr/p1", "DOCX header quote"), {}
        )
        review_checks.source_text(
            self.project, self.source("ftr/p1", "DOCX footer quote"), {}
        )

    def test_reads_single_first_and_even_page_parts_without_default_parts(self):
        for variant in ("first", "even"):
            registered = self.register_variant(variant, variant=variant)
            with self.subTest(variant=variant):
                review_checks.source_text(
                    self.project,
                    self.source("hdr/p1", f"variant {variant} header", registered),
                    {},
                )
                review_checks.source_text(
                    self.project,
                    self.source("ftr/p1", f"variant {variant} footer", registered),
                    {},
                )

    def test_short_location_rejects_distinct_default_first_even_parts(self):
        registered = self.register_variant("all-parts", all_parts=True)
        for location, quote in (
            ("hdr/p1", "variant default header"),
            ("ftr/p1", "variant default footer"),
        ):
            with self.subTest(location=location):
                with self.assertRaisesRegex(ValueError, "不唯一"):
                    review_checks.source_text(
                        self.project,
                        self.source(location, quote, registered),
                        {},
                    )

    def test_missing_header_or_footer_part_is_not_created_or_read(self):
        registered = self.register_variant("no-parts")
        for location in ("hdr/p1", "ftr/p1"):
            with self.subTest(location=location):
                with self.assertRaisesRegex(ValueError, "不存在"):
                    review_checks.source_text(
                        self.project,
                        self.source(location, "missing part", registered),
                        {},
                    )

    def test_extract_sources_reads_default_first_even_parts_by_identity(self):
        registered = self.register_variant("extract-all-parts", all_parts=True)
        result = extract_sources.extract(
            self.project,
            "DOCX-TEST",
            source_ids=[registered["source_id"]],
        )
        blocks = result["data"]["blocks"]
        headers = [block for block in blocks if block["location"] == "hdr/p1"]
        footers = [block for block in blocks if block["location"] == "ftr/p1"]
        self.assertEqual(len(headers), 3)
        self.assertEqual(len(footers), 3)
        self.assertEqual({block["text"] for block in headers}, {
            "variant default header",
            "variant first header",
            "variant even header",
        })
        self.assertEqual({block["text"] for block in footers}, {
            "variant default footer",
            "variant first footer",
            "variant even footer",
        })

    def test_rejects_wrong_location_and_quote_without_searching_whole_file(self):
        cases = (
            ("body[9]/paragraph", "DOCX body paragraph one", "超出原件范围"),
            ("body[2]/paragraph", "DOCX table quote", "定位类型"),
            ("body[1]/paragraph", "DOCX table quote", "引文不在原件指定位置"),
            ("unknown/p1", "DOCX header quote", "定位"),
            ("body[2]/table/r0/c1", "row one col one", "行号超出"),
            ("body[2]/table/r1/c9", "row one col one", "列号超出"),
        )
        for location, quote, message in cases:
            with self.subTest(location=location):
                with self.assertRaisesRegex(ValueError, message):
                    review_checks.source_text(self.project, self.source(location, quote), {})

        cached = {(self.registered["source_id"], "body[1]/paragraph"): "DOCX table quote"}
        source = self.source("body[1]/paragraph", "DOCX table quote")
        with self.assertRaisesRegex(ValueError, "引文不在原件指定位置"):
            review_checks.source_text(self.project, source, cached)

    def test_rejects_old_hash_before_reading_docx(self):
        stale = self.source("body[1]/paragraph", "DOCX body paragraph one", sha256="0" * 64)
        with self.assertRaisesRegex(ValueError, "哈希过期"):
            review_checks.source_text(self.project, stale, {})

    def test_rejects_corrupt_docx_package(self):
        path = self.project / self.registered["relative_path"]
        path.write_bytes(b"not a zip package")
        digest = bidkit.digest(path)
        registry_path = self.project / "inputs" / "source-registry.json"
        registry = bidkit.read(registry_path)
        registry["sources"][0]["sha256"] = digest
        bidkit.atomic(registry_path, registry)
        corrupted = self.source("body[1]/paragraph", "DOCX body paragraph one", sha256=digest)
        with self.assertRaisesRegex(ValueError, "DOCX"):
            review_checks.source_text(self.project, corrupted, {})

    def test_missing_docx_dependency_is_explicitly_unreviewed(self):
        source = self.source("body[1]/paragraph", "DOCX body paragraph one")
        with patch.dict(sys.modules, {"docx": None}):
            with self.assertRaisesRegex(ValueError, "python-docx"):
                review_checks.source_text(self.project, source, {})

    def test_multiple_distinct_headers_make_short_location_ambiguous(self):
        document = Document()
        document.add_paragraph("body")
        document.sections[0].header.paragraphs[0].text = "first header"
        second = document.add_section(WD_SECTION.NEW_PAGE)
        second.header.is_linked_to_previous = False
        second.header.paragraphs[0].text = "second header"
        source_path = self.project / "work" / "ambiguous.docx"
        document.save(source_path)
        registered = bidkit.register_source(self.project, source_path, "supplier")
        source = {
            "source_id": registered["source_id"],
            "revision": registered["revision"],
            "sha256": registered["sha256"],
            "location": "header/p1",
            "quote": "first header",
        }
        with self.assertRaisesRegex(ValueError, "不唯一|歧义"):
            review_checks.source_text(self.project, source, {})

    def test_readable_quote_only_validates_the_evidence_location(self):
        source = self.source("body[1]/paragraph", "DOCX body paragraph one")
        cache = {}
        self.assertIsNone(review_checks.source_text(self.project, source, cache))
        self.assertEqual(cache, {})

    def test_empty_quote_cannot_prove_any_source(self):
        for quote in ("", " \n\t"):
            with self.subTest(quote=quote), self.assertRaisesRegex(ValueError, "引文不能为空"):
                review_checks.source_text(self.project, self.source("body[1]/paragraph", quote), {})

    def test_linked_headers_share_one_unambiguous_part(self):
        path = self.project / self.registered['relative_path']
        document = Document(path)
        document.add_section(WD_SECTION.NEW_PAGE)
        updated_path = self.project / 'work/linked.docx'
        document.save(updated_path)
        registered = bidkit.register_source(self.project, updated_path, 'supplier')
        source = dict(self.source('hdr/p1', 'DOCX header quote'), **{
            key: registered[key] for key in ('source_id', 'revision', 'sha256')})
        review_checks.source_text(self.project, source, {})

    def test_standalone_review_consumers_declare_existing_docx_dependency(self):
        from refresh_distribution import DISTRIBUTED_SCRIPTS
        for skill_id, scripts in DISTRIBUTED_SCRIPTS.items():
            if 'review_checks.py' not in scripts:
                continue
            with self.subTest(skill_id=skill_id):
                requirements = (ROOT / 'skills' / skill_id / 'requirements.txt').read_text()
                self.assertIn('python-docx>=1.1.2,<2', requirements)


if __name__ == "__main__":
    unittest.main()
