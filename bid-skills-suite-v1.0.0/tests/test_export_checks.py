"""Regression tests for the byte-bound DOCX/PDF export acceptance receipt."""

import base64
import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import build_docx  # noqa: E402
import export_checks  # noqa: E402

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aZ1kAAAAASUVORK5CYII="
)


@unittest.skipUnless(
    importlib.util.find_spec("docx") and importlib.util.find_spec("fitz"),
    "python-docx and PyMuPDF are required",
)
class ExportChecksTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        (self.root / "work").mkdir()
        (self.root / "work/project.json").write_text(
            '{"project_id":"EXPORT-TEST"}', encoding="utf-8"
        )
        (self.root / "figure.png").write_bytes(PNG)

    def tearDown(self):
        self.directory.cleanup()

    def spec(self, *, toc=False, expected_fact=True):
        value = {
            "title": "Export Acceptance",
            "subtitle": "Small fixture",
            "expected_facts": (
                [{"id": "deadline", "value": "Deadline: 30 days"}] if expected_fact else []
            ),
            "export_settings": {
                "header": {"text": "", "page_field": False},
                "footer": {"text": "", "page_field": False},
                "cover": {"metadata_rows": [{"label": "Project", "value": "EXPORT-TEST"}]},
                "toc": {"enabled": toc, "title": "Contents", "levels": 2},
            },
            "sections": [
                {
                    "title": "1. Delivery Plan",
                    "level": 1,
                    "paragraphs": ["Body paragraph", "Deadline: 30 days"],
                    "tables": [{"headers": ["Field", "Value"], "rows": [["Status", "Ready"]]}],
                    "images": [{"path": "figure.png", "width_cm": 2, "caption": "Plan figure"}],
                }
            ],
        }
        path = self.root / "document-spec.json"
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return path

    def build_docx(self, spec_path, name="draft.docx"):
        output = self.root / name
        build_docx.build(spec_path, output, self.root)
        return output

    def build_pdf(self, *, toc=False, include_fact=True, name="draft.pdf", include_image=True):
        import fitz

        output = self.root / name
        document = fitz.open()
        page = document.new_page()
        values = ["Export Acceptance", "Small fixture", "Project", "EXPORT-TEST"]
        if toc:
            values.append("Contents")
            values.append("1. Delivery Plan")
        values += ["1. Delivery Plan", "Body paragraph"]
        if include_fact:
            values.append("Deadline: 30 days")
        values += ["Field", "Value", "Status", "Ready", "Plan figure"]
        y = 60
        for value in values:
            page.insert_text((50, y), value)
            y += 18
        if include_image:
            page.insert_image(fitz.Rect(50, y, 70, y + 20), stream=PNG)
        document.save(output)
        document.close()
        return output

    def test_real_docx_and_pdf_pass_with_counts_and_identity(self):
        spec = self.spec()
        docx = self.build_docx(spec)
        pdf = self.build_pdf()
        receipt = export_checks.inspect(spec, docx, pdf, self.root)
        self.assertTrue(receipt["passed"], receipt["findings"])
        self.assertEqual(receipt["project_id"], "EXPORT-TEST")
        self.assertEqual(receipt["counts"]["docx"]["images"], 1)
        self.assertEqual(receipt["counts"]["pdf"]["pages"], 1)
        self.assertEqual(receipt["visual_qa"], "NOT_RUN")
        self.assertEqual(receipt["coverage"]["docx"]["facts"]["absent"], [])

    def test_fixed_form_docx_acceptance_checks_merge_spans_and_missing_highlight(self):
        from docx import Document
        from docx.oxml.ns import qn

        spec = self.spec()
        value = json.loads(spec.read_text(encoding="utf-8"))
        value["sections"][0]["images"] = []
        value["sections"][0]["tables"] = [{
            "headers": [["事项", "", "响应"], ["编号", "名称", "结果"]],
            "header_rows": 2,
            "rows": [["", "", ""]],
            "widths_cm": [3, 5, 5],
            "merges": [
                {"row": 0, "column": 0, "row_span": 1, "col_span": 2},
                {"row": 1, "column": 0, "row_span": 2, "col_span": 1},
            ],
            "missing_cells": [{"row": 2, "column": 2, "label": "响应结论"}],
        }]
        spec.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        docx = self.build_docx(spec, "fixed-form.docx")
        expected = export_checks._expected(value)
        expected["image_bindings"] = []
        _, findings = export_checks._docx_snapshot(docx, expected)
        self.assertEqual(findings, [])

        tampered = self.root / "fixed-form-no-highlight.docx"
        document = Document(docx)
        cell = document.tables[1].cell(2, 2)
        highlight = cell.paragraphs[0].runs[0]._r.rPr.find(qn("w:highlight"))
        highlight.getparent().remove(highlight)
        document.save(tampered)
        _, findings = export_checks._docx_snapshot(tampered, expected)
        self.assertTrue(any("待补充高亮" in finding for finding in findings), findings)

    def test_declared_fixed_form_sources_are_bound_and_match_section_tables(self):
        import hashlib

        table = {
            "headers": [["字段", ""], ["名称", "值"]],
            "header_rows": 2,
            "rows": [["", ""]],
            "merges": [{"row": 0, "column": 0, "row_span": 1, "col_span": 2}],
            "missing_cells": [{"row": 2, "column": 1, "label": "填报值"}],
        }
        paths = {}
        for name, payload in (("template.docx", b"template"), ("filled.docx", b"filled")):
            path = self.root / name
            path.write_bytes(payload)
            paths[name] = hashlib.sha256(payload).hexdigest()
        table_path = self.root / "table.json"
        table_path.write_text(json.dumps({"paragraphs": [], "tables": [table]}, ensure_ascii=False))
        paths["table.json"] = hashlib.sha256(table_path.read_bytes()).hexdigest()
        commercial_path = self.root / "artifacts/12-commercial-content.json"
        commercial_path.parent.mkdir()
        commercial_path.write_text(json.dumps({"project_id": "EXPORT-TEST", "data": {"forms": [{
            "id": "FORM-1", "section_id": "FORM-1",
            "template_path": "template.docx", "filled_template_path": "filled.docx",
            "table_spec_path": "table.json", "fields": [],
        }]}}, ensure_ascii=False), encoding="utf-8")
        paths["artifacts/12-commercial-content.json"] = hashlib.sha256(
            commercial_path.read_bytes()).hexdigest()
        spec = {
            "sections": [{
                "section_id": "FORM-1", "paragraphs": [], "tables": [table],
            }],
            "inputs": [{"relative_path": name, "sha256": digest}
                       for name, digest in paths.items()],
        }
        self.assertEqual(export_checks._check_current_chapters(spec, self.root), [])

        table_path.write_text(json.dumps({"paragraphs": [], "tables": [{
            **table, "merges": []
        }]}), encoding="utf-8")
        findings = export_checks._check_current_chapters(spec, self.root)
        self.assertTrue(any("table.json" in finding and "未绑定" in finding for finding in findings))

    def test_declared_fixed_form_source_path_cannot_escape_or_be_missing(self):
        base = {"sections": [{"paragraphs": [], "tables": []}], "inputs": []}
        for field, value, marker in (
            ("template_path", "../outside.docx", "路径越界"),
            ("filled_template_path", "missing.docx", "文件缺失"),
            ("table_spec_path", "missing.json", "文件缺失"),
        ):
            with self.subTest(field=field):
                spec = json.loads(json.dumps(base))
                spec["sections"][0][field] = value
                findings = export_checks._check_current_chapters(spec, self.root)
                self.assertTrue(any(marker in finding for finding in findings), findings)

    def test_multiple_forms_match_ordered_subsequences_inside_shared_section(self):
        import hashlib

        def table(name, *, merged=False):
            value = {
                "headers": [[name, ""]] if merged else ["字段", "值"],
                "rows": [[name + "-row", ""]] if merged else [[name, "响应"]],
                "paragraphs_before": [name + "-before"],
                "paragraphs_after": [name + "-after"],
            }
            if merged:
                value["widths_cm"] = [3, 5]
                value["merges"] = [{"row": 0, "column": 0, "row_span": 1, "col_span": 2}]
            return value

        technical = table("技术表")
        form_one_a = table("表单一-A", merged=True)
        form_one_b = table("表单一-B")
        form_two = table("表单二")
        tail = table("章节尾表")
        section = {
            "section_id": "SHARED",
            "paragraphs": ["技术正文前", "表单一段落", "技术正文中", "表单二段落", "技术正文后"],
            "tables": [technical, form_one_a, form_one_b, form_two, tail],
        }
        owners = [
            ("form-one.json", {
                "section_id": "SHARED", "table_spec_path": "form-one.json",
                "table_spec": {"paragraphs": ["表单一段落"], "tables": [form_one_a, form_one_b]},
            }),
            ("form-two.json", {
                "section_id": "SHARED", "table_spec_path": "form-two.json",
                "table_spec": {"paragraphs": ["表单二段落"], "tables": [form_two]},
            }),
        ]
        bindings = {}
        for filename, owner in owners:
            path = self.root / filename
            path.write_text(json.dumps(owner["table_spec"], ensure_ascii=False), encoding="utf-8")
            bindings[filename] = hashlib.sha256(path.read_bytes()).hexdigest()
        sections_by_id = {"SHARED": section}
        cursors = {}
        for filename, owner in owners:
            self.assertEqual(
                export_checks._check_declared_fixed_form_paths(
                    owner, filename, sections_by_id=sections_by_id, bindings=bindings,
                    project_root=self.root, sequence_cursors=cursors),
                [],
            )

        missing_paragraph = dict(section, paragraphs=["技术正文前", "表单二段落", "技术正文后"])
        findings = export_checks._check_declared_fixed_form_paths(
            owners[0][1], "form-one.json", sections_by_id={"SHARED": missing_paragraph},
            bindings=bindings, project_root=self.root, sequence_cursors={})
        self.assertTrue(any("paragraphs" in finding for finding in findings), findings)

        missing_table = dict(section, tables=[technical, form_one_a, form_two, tail])
        findings = export_checks._check_declared_fixed_form_paths(
            owners[0][1], "form-one.json", sections_by_id={"SHARED": missing_table},
            bindings=bindings, project_root=self.root, sequence_cursors={})
        self.assertTrue(any("tables" in finding for finding in findings), findings)

        swapped = dict(section, tables=[technical, form_one_b, form_one_a, form_two, tail])
        findings = export_checks._check_declared_fixed_form_paths(
            owners[0][1], "form-one.json", sections_by_id={"SHARED": swapped},
            bindings=bindings, project_root=self.root, sequence_cursors={})
        self.assertTrue(any("tables" in finding for finding in findings), findings)

        bad_merge = json.loads(json.dumps(owners[0][1]["table_spec"]))
        bad_merge["tables"][0]["merges"] = []
        bad_path = self.root / "form-one-bad-merge.json"
        bad_path.write_text(json.dumps(bad_merge, ensure_ascii=False), encoding="utf-8")
        bad_owner = dict(owners[0][1], table_spec_path="form-one-bad-merge.json")
        bad_bindings = dict(bindings, **{
            "form-one-bad-merge.json": hashlib.sha256(bad_path.read_bytes()).hexdigest()
        })
        findings = export_checks._check_declared_fixed_form_paths(
            bad_owner, "form-one-bad-merge.json", sections_by_id=sections_by_id,
            bindings=bad_bindings, project_root=self.root, sequence_cursors={})
        self.assertTrue(any("tables" in finding for finding in findings), findings)

    def test_docx_merge_snapshot_ignores_vertical_horizontal_continuation(self):
        from docx import Document

        spec = self.spec()
        value = json.loads(spec.read_text(encoding="utf-8"))
        value["sections"][0]["images"] = []
        value["sections"][0]["tables"] = [{
            "headers": [["合并标题", ""]],
            "rows": [["", ""]],
            "merges": [{"row": 0, "column": 0, "row_span": 2, "col_span": 2}],
        }]
        spec.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        docx = self.build_docx(spec, "vertical-horizontal-merge.docx")
        document = Document(docx)
        self.assertEqual(
            export_checks._docx_table_merges(document.tables[1]),
            [{"row": 0, "column": 0, "row_span": 2, "col_span": 2}],
        )

    def test_page_reference_in_body_is_not_stripped_as_running_footer(self):
        import fitz

        expected = export_checks._expected({
            "title": "Index", "export_settings": {
                "cover": {"enabled": False},
                "header": {"text": "", "page_field": False},
                "footer": {"page_field": True, "page_prefix": "Page ", "page_suffix": ""},
            },
            "sections": [{"title": "1. Index", "paragraphs": ["See", "Page 7", "after"]}],
        })
        path = self.root / "body-page-reference.pdf"
        with fitz.open() as document:
            page = document.new_page()
            for y, text in ((80, "1. Index"), (100, "See"), (120, "Page 7"),
                            (140, "after"), (820, "Page 1")):
                page.insert_text((50, y), text)
            document.save(path)
        _, findings = export_checks._pdf_snapshot(path, expected)
        self.assertEqual(findings, [])

    def test_body_page_reference_cannot_prove_missing_running_footer(self):
        import fitz

        expected = export_checks._expected({
            "title": "Index", "export_settings": {
                "cover": {"enabled": False},
                "header": {"text": "", "page_field": False},
                "footer": {"page_field": True, "page_prefix": "Page ", "page_suffix": ""},
            },
            "sections": [{"title": "1. Index", "paragraphs": ["Page 1"]}],
        })
        path = self.root / "missing-running-footer.pdf"
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text((50, 80), "1. Index")
            page.insert_text((50, 120), "Page 1")
            document.save(path)
        _, findings = export_checks._pdf_snapshot(path, expected)
        self.assertTrue(any("PDF第1页缺少声明的页码" in value for value in findings))

    def test_pdf_expected_order_includes_table_level_before_and_after(self):
        import fitz

        expected = export_checks._expected({
            "title": "Table order",
            "export_settings": {
                "cover": {"enabled": False},
                "header": {"text": "", "page_field": False},
                "footer": {"text": "", "page_field": False},
            },
            "sections": [{
                "title": "Fixed form",
                "paragraphs": [],
                "tables": [
                    {"headers": ["字段", "值"], "rows": [["A", "1"]],
                     "paragraphs_before": ["表一脚注"],
                     "paragraphs_after": ["表一签署区"]},
                    {"headers": ["字段", "值"], "rows": [["B", "2"]],
                     "paragraphs_before": ["表二脚注"],
                     "paragraphs_after": ["表二签署区"]},
                ],
            }],
        })
        path = self.root / "table-order-wrong.pdf"
        with fitz.open() as document:
            page = document.new_page()
            # Every string is present, but table one’s after text comes after
            # table two’s before text, so the fixed form order is invalid.
            values = ["Fixed form", "表一脚注", "字段", "值", "A", "1",
                      "表二脚注", "表一签署区", "字段", "值", "B", "2", "表二签署区"]
            for y, value in enumerate(values, start=1):
                page.insert_text((50, 40 + y * 18), value)
            document.save(path)
        _, findings = export_checks._pdf_snapshot(path, expected)
        self.assertTrue(any("正文/标题/表格/图注或顺序不一致" in value for value in findings), findings)

    def test_pdf_expected_order_expands_merge_for_docx_but_reads_anchor_once(self):
        import fitz

        expected = export_checks._expected({
            "title": "Merged form",
            "export_settings": {
                "cover": {"enabled": False},
                "header": {"text": "", "page_field": False},
                "footer": {"text": "", "page_field": False},
            },
            "sections": [{
                "title": "Fixed form",
                "paragraphs": [],
                "tables": [{
                    "headers": [["事项", ""], ["编号", "名称"]],
                    "header_rows": 2,
                    "rows": [["", ""]],
                    "merges": [{"row": 0, "column": 0, "row_span": 1, "col_span": 2}],
                    "missing_cells": [{"row": 2, "column": 1, "label": "填报值"}],
                }],
            }],
        })
        path = self.root / "merged-form.pdf"
        with fitz.open() as document:
            page = document.new_page()
            for y, value in enumerate(("Fixed form", "事项", "编号", "名称", "【待补充：填报值】"), start=1):
                page.insert_text((50, 40 + y * 18), value, fontname="china-s")
            document.save(path)
        _, findings = export_checks._pdf_snapshot(path, expected)
        self.assertEqual(findings, [])

    def test_missing_heading_image_and_fact_cannot_pass(self):
        spec = self.spec()
        actual_spec = json.loads(spec.read_text(encoding="utf-8"))
        actual_spec["sections"][0]["title"] = "1. Other Plan"
        actual_spec["sections"][0]["images"] = []
        actual_spec["sections"][0]["paragraphs"] = ["Body paragraph"]
        actual = self.root / "actual.json"
        actual.write_text(json.dumps(actual_spec), encoding="utf-8")
        docx = self.build_docx(actual, "missing.docx")
        pdf = self.build_pdf(include_fact=False, include_image=False, name="missing.pdf")
        receipt = export_checks.inspect(spec, docx, pdf, self.root)
        self.assertFalse(receipt["passed"])
        joined = "\n".join(receipt["findings"])
        self.assertIn("章节标题", joined)
        self.assertIn("插图数量", joined)
        self.assertIn("关键事实", joined)

    def test_toc_placeholder_is_explicitly_rejected(self):
        spec = self.spec(toc=True)
        docx = self.build_docx(spec, "toc.docx")
        pdf = self.build_pdf(toc=True, name="toc.pdf")
        receipt = export_checks.inspect(spec, docx, pdf, self.root)
        self.assertFalse(receipt["passed"])
        self.assertTrue(any("目录" in finding for finding in receipt["findings"]))

    def test_verify_receipt_rejects_changed_output_bytes_and_spec(self):
        spec = self.spec()
        docx = self.build_docx(spec)
        pdf = self.build_pdf()
        receipt = export_checks.inspect(spec, docx, pdf, self.root)
        self.assertTrue(receipt["passed"], receipt["findings"])
        pdf.write_bytes(pdf.read_bytes() + b"drift")
        with self.assertRaisesRegex(ValueError, "当前文件未通过|漂移"):
            export_checks.verify_receipt(receipt, spec, docx, pdf, self.root)

    def test_wrong_picture_with_same_count_fails(self):
        import fitz
        spec = self.spec()
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
        pix.clear_with(0)
        (self.root / 'other.png').write_bytes(pix.tobytes('png'))
        actual = json.loads(spec.read_text())
        actual['sections'][0]['images'][0]['path'] = 'other.png'
        actual_path = self.root / 'actual.json'
        actual_path.write_text(json.dumps(actual))
        receipt = export_checks.inspect(spec, self.build_docx(actual_path), self.build_pdf(), self.root)
        self.assertFalse(receipt['passed'])
        self.assertTrue(any('插图内容' in f for f in receipt['findings']))

    def test_incomplete_spec_cannot_hide_current_outline(self):
        spec = self.spec()
        (self.root / 'artifacts').mkdir()
        (self.root / 'artifacts/08-outline.json').write_text(json.dumps({'data': {'sections': [{'id': 'A'}, {'id': 'B'}]}}))
        (self.root / 'artifacts/11-technical-content.json').write_text(json.dumps({'data': {'chapters': [{'section_id': 'A', 'body_markdown': 'Actual text'}]}}))
        receipt = export_checks.inspect(spec, self.build_docx(spec), self.build_pdf(), self.root)
        self.assertFalse(receipt['passed'])
        self.assertTrue(any('当前目录' in f for f in receipt['findings']))
        self.assertTrue(any('漏当前正文' in f for f in receipt['findings']))

    def test_markdown_quote_markers_do_not_require_literal_export_but_text_does(self):
        import hashlib

        spec = self.spec()
        value = json.loads(spec.read_text())
        value['sections'][0]['section_id'] = 'A'
        (self.root / 'artifacts').mkdir()
        writing = self.root / 'artifacts/11-technical-content.json'
        for body, missing in [('> Body paragraph', False),
                              ('> > **Body paragraph**', False),
                              ('> Body paragraph\n> Required safety condition', True),
                              ('Threshold > 5', True)]:
            with self.subTest(body=body):
                writing.write_text(json.dumps({'project_id': 'EXPORT-TEST', 'data': {
                    'chapters': [{'section_id': 'A', 'body_markdown': body}]}}))
                value['inputs'] = [{'relative_path': 'artifacts/11-technical-content.json',
                                    'sha256': hashlib.sha256(writing.read_bytes()).hexdigest()}]
                findings = export_checks._check_current_chapters(value, self.root)
                self.assertEqual(bool(findings), missing, findings)

    def test_other_project_artifacts_cannot_pass_receipt(self):
        import hashlib
        spec = self.spec()
        value = json.loads(spec.read_text())
        value['sections'][0]['section_id'] = 'A'
        (self.root / 'artifacts').mkdir()
        inputs = {'08-outline.json': {'data': {'sections': [{'id': 'A'}]}},
                  '11-technical-content.json': {'data': {'chapters': [{'section_id': 'A', 'body_markdown': 'Body paragraph'}]}}}
        value['inputs'] = []
        for name, payload in inputs.items():
            payload['project_id'] = 'OTHER-PROJECT'
            path = self.root / 'artifacts' / name
            path.write_text(json.dumps(payload))
            value['inputs'].append({'relative_path': 'artifacts/'+name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
        spec.write_text(json.dumps(value))
        receipt = export_checks.inspect(spec, self.build_docx(spec), self.build_pdf(), self.root)
        self.assertFalse(receipt['passed'])
        self.assertTrue(any('项目身份不一致' in f for f in receipt['findings']))

    def test_missing_declared_header_footer_and_page_field_fails(self):
        spec = self.spec()
        expected = json.loads(spec.read_text())
        expected['export_settings']['header'] = {'text': 'CONFIDENTIAL HEADER'}
        expected['export_settings']['footer'] = {'text': 'Page ', 'page_field': True, 'page_prefix': '', 'page_suffix': ''}
        docx = self.build_docx(spec)
        spec.write_text(json.dumps(expected))
        receipt = export_checks.inspect(spec, docx, self.build_pdf(), self.root)
        self.assertFalse(receipt['passed'])
        self.assertTrue(any('DOCX缺少声明的header' in f for f in receipt['findings']))
        self.assertTrue(any('PDF第1页缺少声明的页码' in f for f in receipt['findings']))

    def test_internal_signing_trace_cannot_pass_even_when_in_spec(self):
        import fitz
        spec = self.spec()
        value = json.loads(spec.read_text())
        trace = 'confirmation_ref: work/test-authorisation.md'
        value['sections'][0]['paragraphs'].append(trace)
        spec.write_text(json.dumps(value))
        docx = self.build_docx(spec)
        pdf = self.build_pdf()
        with fitz.open(pdf) as document:
            document[0].insert_text((50, 700), trace)
            document.saveIncr()
        receipt = export_checks.inspect(spec, docx, pdf, self.root)
        self.assertFalse(receipt['passed'])
        self.assertTrue(any('内部执行信息' in finding for finding in receipt['findings']))

    def test_current_commercial_value_cannot_be_omitted_from_spec(self):
        import hashlib
        spec = self.spec()
        value = json.loads(spec.read_text())
        value['sections'][0]['section_id'] = 'A'
        (self.root / 'artifacts').mkdir()
        inputs = {
            '08-outline.json': {'data': {'sections': [{'id': 'A'}]}},
            '11-technical-content.json': {'data': {'chapters': [{'section_id': 'A', 'body_markdown': 'Body paragraph'}]}},
            '12-commercial-content.json': {'data': {'forms': [{'id': 'F1', 'fields': [{'key': 'Amount', 'value': '650000 yuan'}]}]}},
        }
        value['inputs'] = []
        for name, payload in inputs.items():
            payload['project_id'] = 'EXPORT-TEST'
            path = self.root / 'artifacts' / name
            path.write_text(json.dumps(payload))
            value['inputs'].append({'relative_path': 'artifacts/'+name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
        spec.write_text(json.dumps(value))
        receipt = export_checks.inspect(spec, self.build_docx(spec), self.build_pdf(), self.root)
        self.assertFalse(receipt['passed'])
        self.assertTrue(any('漏商务字段' in finding for finding in receipt['findings']))

    def test_each_partial_current_artifact_is_checked(self):
        import hashlib
        import itertools
        directory = self.root/'artifacts'
        directory.mkdir()
        names = ('08-outline.json', '11-technical-content.json', '12-commercial-content.json')
        payloads = (
            {'sections': [{'id': 'B'}]},
            {'chapters': [{'section_id': 'A', 'body_markdown': 'Required current body'}]},
            {'forms': [{'id': 'F1', 'fields': [{'key': 'Amount', 'value': '650000 yuan'}]}]},
        )
        for count in range(1, 4):
            for indices in itertools.combinations(range(3), count):
                with self.subTest(indices=indices):
                    for old in directory.glob('*.json'):
                        old.unlink()
                    spec = self.spec()
                    value = json.loads(spec.read_text())
                    value['sections'][0]['section_id'] = 'A'
                    value['inputs'] = []
                    for i in indices:
                        path = directory/names[i]
                        path.write_text(json.dumps({'project_id': 'EXPORT-TEST', 'data': payloads[i]}))
                        value['inputs'].append({'relative_path': 'artifacts/'+names[i], 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
                    spec.write_text(json.dumps(value))
                    docx = self.build_docx(spec, 'partial-'+''.join(map(str, indices))+'.docx')
                    pdf = self.build_pdf(name='partial-'+''.join(map(str, indices))+'.pdf')
                    receipt = export_checks.inspect(spec, docx, pdf, self.root)
                    self.assertFalse(receipt['passed'])
                    for i in indices:
                        self.assertTrue(any(('当前目录', '漏当前正文', '漏商务字段')[i] in f for f in receipt['findings']), receipt['findings'])

    def test_malformed_current_artifact_returns_failed_receipt(self):
        spec = self.spec()
        (self.root/'artifacts').mkdir()
        (self.root/'artifacts/08-outline.json').write_text('{"data":{}}')
        receipt = export_checks.inspect(spec, self.build_docx(spec), self.build_pdf(), self.root)
        self.assertFalse(receipt['passed'])
        self.assertTrue(any('结构' in f for f in receipt['findings']))

    def test_omitted_page_field_uses_builder_default(self):
        spec = self.spec()
        value = json.loads(spec.read_text())
        value['export_settings']['footer'] = {'text': 'CONF', 'page_field': False}
        spec.write_text(json.dumps(value))
        docx = self.build_docx(spec)
        value['export_settings']['footer'].pop('page_field')
        spec.write_text(json.dumps(value))
        receipt = export_checks.inspect(spec, docx, self.build_pdf(), self.root)
        self.assertFalse(receipt['passed'])
        self.assertTrue(any('PAGE' in f for f in receipt['findings']), receipt['findings'])

    def test_internal_paths_are_rejected_regardless_of_extension(self):
        for i, path in enumerate(('work/attempt.yaml', 'reviews/a.csv', 'artifacts/proof.png', 'profiles/layout.docx')):
            with self.subTest(path=path):
                spec = self.spec()
                value = json.loads(spec.read_text())
                value['sections'][0]['paragraphs'].append(path)
                spec.write_text(json.dumps(value))
                receipt = export_checks.inspect(spec, self.build_docx(spec, 'path'+str(i)+'.docx'), self.build_pdf(name='path'+str(i)+'.pdf'), self.root)
                self.assertFalse(receipt['passed'])
                self.assertTrue(any('内部执行信息' in f for f in receipt['findings']))

    def test_malformed_files_fail_with_machine_findings(self):
        spec = self.spec()
        docx = self.root / "bad.docx"
        pdf = self.root / "bad.pdf"
        docx.write_bytes(b"not a zip")
        pdf.write_bytes(b"not a pdf")
        receipt = export_checks.inspect(spec, docx, pdf, self.root)
        self.assertFalse(receipt["passed"])
        self.assertTrue(any("DOCX格式无效" in finding for finding in receipt["findings"]))
        self.assertTrue(any("PDF格式无效" in finding for finding in receipt["findings"]))


if __name__ == "__main__":
    unittest.main()
