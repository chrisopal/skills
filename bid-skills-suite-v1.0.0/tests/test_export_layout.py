"""Export configuration must survive into a real, inspectable DOCX."""

import base64
import importlib.util
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import build_docx  # noqa: E402


@unittest.skipUnless(importlib.util.find_spec('docx'), 'Declared optional python-docx not installed')
class ExportLayoutTests(unittest.TestCase):
    def _write_spec(self, root, *, embedded=None, images=None):
        root.mkdir(parents=True, exist_ok=True)
        spec = {
            'title': '原始项目标题',
            'subtitle': '原始副标题',
            'style': {'body_size_pt': 12, 'body_line_spacing': 1.5,
                      'body_east_asia_font': 'Songti SC',
                      'heading_east_asia_font': 'Heiti SC'},
            'sections': [
                {'title': '一、技术方案', 'level': 1, 'paragraphs': ['正文一。']},
                {'title': '1.1 实施安排', 'level': 2, 'paragraphs': ['正文二。'],
                 'images': images or [],
                 'tables': [{'headers': ['要求', '响应'], 'rows': [['原文要求', '待核验响应']]}]},
            ],
        }
        if embedded is not None:
            spec['export_settings'] = embedded
        path = root / 'spec.json'
        path.write_text(json.dumps(spec, ensure_ascii=False), encoding='utf-8')
        return path

    def _config(self):
        return {
            'style': {
                'body_size_pt': 11,
                'body_line_spacing': 1.25,
                'body_east_asia_font': 'Songti SC',
                'body_latin_font': 'Times New Roman',
                'body_first_line_indent_chars': 2,
                'body_space_before_pt': 0,
                'body_space_after_pt': 0,
                'body_alignment': 'justify',
                'heading_sizes_pt': [15, 13, 11],
                'heading_east_asia_font': 'Heiti SC',
                'title_size_pt': 24,
                'subtitle_size_pt': 14,
                'heading_space_before_pt': 12,
                'heading_space_after_pt': 6,
                'table_size_pt': 10,
                'caption_size_pt': 10,
                'page_margins_cm': {'top': 2.4, 'bottom': 2.4, 'left': 2.2, 'right': 2.2},
            },
            'cover': {
                'title': '配置后的封面标题',
                'subtitle': '配置后的封面副标题',
                'alignment': 'center',
                'own_page': True,
                'metadata_rows': [
                    {'label': '项目名称', 'value': '示例项目'},
                    {'label': '文件状态', 'value': '内部工作稿'},
                ],
                'typography': {
                    'title_east_asia_font': 'Heiti SC',
                    'subtitle_east_asia_font': 'Songti SC',
                    'metadata_east_asia_font': 'Songti SC',
                },
            },
            'toc': {
                'enabled': True,
                'title': '目录',
                'levels': 2,
                'typography': {
                    'east_asia_font': 'Songti SC',
                    'latin_font': 'Times New Roman',
                    'size_pt': 10.5,
                    'line_spacing': 1.2,
                    'space_before_pt': 0,
                    'space_after_pt': 3,
                    'alignment': 'left',
                },
            },
            'header': {'text': '内部工作稿｜配置页眉', 'alignment': 'left'},
            'footer': {'text': '候选版', 'alignment': 'right', 'page_field': True,
                       'page_prefix': '第 ', 'page_suffix': ' 页'},
        }

    def test_external_settings_create_cover_toc_bookmarks_and_fields(self):
        from docx import Document
        from docx.oxml.ns import qn

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            spec_path = self._write_spec(root)
            settings_path = root / 'profiles-export-settings.json'
            settings_path.write_text(json.dumps(self._config(), ensure_ascii=False), encoding='utf-8')
            output = root / 'draft.docx'

            result = build_docx.build(spec_path, output, root, settings_path)

            self.assertEqual(result['status'], 'draft_only')
            self.assertEqual(result['visual_qa'], 'NOT_RUN')
            document = Document(output)
            self.assertEqual(document.paragraphs[0].text, '配置后的封面标题')
            self.assertEqual(document.paragraphs[1].text, '配置后的封面副标题')
            self.assertEqual([[cell.text for cell in row.cells] for row in document.tables[0].rows], [
                ['项目名称', '示例项目'], ['文件状态', '内部工作稿']
            ])
            body_table = document.tables[1]
            self.assertTrue(all(paragraph.paragraph_format.keep_with_next
                                for paragraph in body_table.rows[0].cells[0].paragraphs))
            self.assertTrue(all(row._tr.get_or_add_trPr().find(qn('w:cantSplit')) is not None
                                for row in body_table.rows))
            self.assertEqual(document.sections[0].header.paragraphs[0].text, '内部工作稿｜配置页眉')
            self.assertEqual(document.sections[0].footer.paragraphs[0].text, '候选版第  页')
            header_run = next(run for run in document.sections[0].header.paragraphs[0].runs if run.text)
            header_color = header_run._r.rPr.find(qn('w:color'))
            self.assertEqual(header_color.get(qn('w:val')), '000000')
            self.assertEqual(document.styles['Normal'].font.size.pt, 11)
            self.assertEqual(document.styles['Normal'].paragraph_format.line_spacing, 1.25)
            self.assertEqual(document.styles['Heading 1'].font.size.pt, 15)
            for style_name in ('Title', 'Subtitle'):
                self.assertFalse(document.styles[style_name].element.xpath('./w:pPr/w:pBdr'))
            for style_name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'toc 1', 'toc 2'):
                color = document.styles[style_name].element.rPr.find(qn('w:color'))
                self.assertEqual(color.get(qn('w:val')), '000000', style_name)
            self.assertAlmostEqual(document.sections[0].left_margin.cm, 2.2, places=2)
            body = next(paragraph for paragraph in document.paragraphs if paragraph.text == '正文一。')
            self.assertEqual(body.paragraph_format.first_line_indent.pt, 22)
            self.assertEqual(body.alignment, 3)
            self.assertTrue(any(paragraph.text == '目录' for paragraph in document.paragraphs))
            toc_paragraph = next(paragraph for paragraph in document.paragraphs
                                 if '目录字段将在' in paragraph.text)
            toc_placeholder = next(run for run in toc_paragraph.runs if run.text)
            self.assertEqual(toc_placeholder.font.size.pt, 10.5)
            self.assertEqual(toc_placeholder._r.rPr.find(qn('w:color')).get(qn('w:val')), '000000')
            self.assertEqual(toc_paragraph._p.pPr.jc.get(qn('w:val')), 'left')
            self.assertEqual(document.styles['toc 1'].font.size.pt, 10.5)
            self.assertEqual(document.styles['toc 1'].style_id, 'TOC1')
            self.assertTrue(document.styles['toc 1'].builtin)
            self.assertIsNone(document.styles['toc 1'].element.get(qn('w:customStyle')))
            self.assertEqual(document.styles['toc 1'].paragraph_format.left_indent.cm, 0)
            self.assertAlmostEqual(document.styles['toc 2'].paragraph_format.left_indent.cm, 0.6, places=2)
            field_tags = {qn('w:fldChar'), qn('w:instrText')}
            self.assertFalse(any(child.tag in field_tags for child in toc_paragraph._p))
            field_nodes = toc_paragraph._p.xpath('.//w:fldChar') + toc_paragraph._p.xpath('.//w:instrText')
            self.assertTrue(field_nodes)
            self.assertTrue(all(node.getparent().tag == qn('w:r') for node in field_nodes))
            self.assertGreaterEqual(sum(1 for paragraph in document.paragraphs
                                        if paragraph._p.xpath('.//w:br[@w:type="page"]')), 2)

            with zipfile.ZipFile(output) as archive:
                document_xml = archive.read('word/document.xml').decode('utf-8')
                header_xml = archive.read('word/header1.xml').decode('utf-8')
                footer_xml = archive.read('word/footer1.xml').decode('utf-8')
            self.assertIn('TOC \\o "1-2" \\h \\z \\u', document_xml)
            self.assertIn('<w:r><w:fldChar', document_xml)
            self.assertIn('w:dirty="true"', document_xml)
            self.assertIn('w:name="_Toc1"', document_xml)
            self.assertIn('w:name="_Toc2"', document_xml)
            self.assertEqual(document_xml.count('w:tblHeader'), 1)
            self.assertEqual(document_xml.count('w:cantSplit'), 2)
            self.assertIn('w:instr="PAGE"', footer_xml)
            self.assertIn('内部工作稿｜配置页眉', header_xml)

    def test_header_page_field_is_not_silently_ignored(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self._write_spec(root, embedded={'header': {'text': '页眉', 'page_field': True,
                                                        'page_prefix': ' 第 ', 'page_suffix': ' 页'}})
            output = root / 'header-page.docx'
            build_docx.build(root / 'spec.json', output, root)
            with zipfile.ZipFile(output) as archive:
                header = archive.read('word/header1.xml').decode('utf-8')
            self.assertIn('w:instr="PAGE"', header)
            self.assertIn(' 第 ', header)

    def test_embedded_settings_and_legacy_call_keep_conservative_defaults(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            embedded = {'style': {'body_size_pt': 10},
                        'cover': {'title': '内嵌标题', 'own_page': False},
                        'toc': {'enabled': False},
                        'header': {'text': '内嵌页眉'},
                        'footer': {'text': '内嵌页脚', 'page_field': False}}
            self._write_spec(root, embedded=embedded)
            output = root / 'embedded.docx'
            result = build_docx.build(root / 'spec.json', output, root)
            document = Document(output)
            self.assertEqual(result['status'], 'draft_only')
            self.assertEqual(document.paragraphs[0].text, '内嵌标题')
            self.assertEqual(document.styles['Normal'].font.size.pt, 10)
            self.assertFalse(any('目录字段将在' in paragraph.text for paragraph in document.paragraphs))
            self.assertEqual(document.sections[0].header.paragraphs[0].text, '内嵌页眉')
            self.assertEqual(document.sections[0].footer.paragraphs[0].text, '内嵌页脚')
            with zipfile.ZipFile(output) as archive:
                self.assertNotIn(b'TOC \\o', archive.read('word/document.xml'))

            legacy_root = root / 'legacy'
            legacy_root.mkdir()
            self._write_spec(legacy_root)
            legacy_output = legacy_root / 'legacy.docx'
            build_docx.build(legacy_root / 'spec.json', legacy_output, legacy_root)
            legacy = Document(legacy_output)
            self.assertEqual(legacy.paragraphs[0].text, '原始项目标题')
            self.assertEqual(legacy.sections[0].header.paragraphs[0].text, '内部工作稿｜需人工复核与签署')
            with zipfile.ZipFile(legacy_output) as archive:
                self.assertIn(b'w:instr="PAGE"', archive.read('word/footer1.xml'))

    def test_inline_figure_does_not_inherit_body_line_spacing(self):
        # Image paragraphs must not inherit prose leading; pagination also needs
        # the height preflight and a real saved/rendered document check.
        from docx import Document
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aZ1kAAAAASUVORK5CYII=')
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'figure.png').write_bytes(png)
            spec = self._write_spec(root, images=[{
                'path': 'figure.png', 'width_cm': 4, 'caption': '图1 设计示意'}])
            output = root / 'figure.docx'
            build_docx.build(spec, output, root)
            document = Document(output)
            index = next(i for i, p in enumerate(document.paragraphs)
                         if p._p.xpath('.//w:drawing'))
            image = document.paragraphs[index].paragraph_format
            self.assertEqual(image.line_spacing, 1)
            self.assertTrue(image.keep_together)
            self.assertTrue(image.keep_with_next)
            self.assertEqual(document.paragraphs[index + 1].text, '图1 设计示意')

    def test_tall_figure_reserves_inline_line_box_and_full_width_caption(self):
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aZ1kAAAAASUVORK5CYII=')
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'figure.png').write_bytes(png)
            spec = self._write_spec(root, images=[{
                'path': 'figure.png', 'width_cm': 16, 'height_cm': 23.4,
                'caption': '完整中文图注' * 16}])
            output = root / 'too-tall.docx'
            with self.assertRaisesRegex(ValueError, '排版组超过'):
                build_docx.build(spec, output, root)
            self.assertFalse(output.exists())

    def test_fixed_form_table_preserves_merges_missing_highlight_and_widths(self):
        from docx import Document
        from docx.oxml.ns import qn

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            spec = self._write_spec(root)
            value = json.loads(spec.read_text(encoding='utf-8'))
            value['sections'][1]['tables'] = [{
                'headers': [['事项', '', '响应'], ['编号', '名称', '结果']],
                'header_rows': 2,
                'rows': [['', '', '']],
                'widths_cm': [3, 5, 5],
                'merges': [
                    {'row': 0, 'column': 0, 'row_span': 1, 'col_span': 2},
                    {'row': 1, 'column': 0, 'row_span': 2, 'col_span': 1},
                ],
                'missing_cells': [{'row': 2, 'column': 2, 'label': '响应结论'}],
            }]
            spec.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
            output = root / 'fixed-form.docx'
            build_docx.build(spec, output, root)
            document = Document(output)
            table = document.tables[0]
            self.assertEqual([[cell.text.strip() for cell in row.cells] for row in table.rows], [
                ['事项', '事项', '响应'], ['编号', '名称', '结果'],
                ['编号', '', '【待补充：响应结论】'],
            ])
            self.assertTrue(all(row._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None
                                for row in table.rows[:2]))
            self.assertTrue(table.rows[0]._tr.xpath('.//w:gridSpan[@w:val="2"]'))
            self.assertTrue(table.rows[1]._tr.xpath('.//w:vMerge[@w:val="restart"]'))
            missing_run = next(run for run in table.cell(2, 2).paragraphs[0].runs
                               if run.text)
            highlight = missing_run._r.rPr.find(qn('w:highlight'))
            self.assertEqual(highlight.get(qn('w:val')), 'yellow')
            self.assertAlmostEqual(table.columns[0].width.cm, 3, places=2)

    def test_fixed_form_table_rejects_unsafe_merge_inputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for merges, rows in [
                ([{'row': 0, 'column': 0, 'row_span': 2, 'col_span': 2},
                  {'row': 1, 'column': 1, 'row_span': 1, 'col_span': 1}], [['x', '']]),
                ([{'row': 0, 'column': 0, 'row_span': 1, 'col_span': 2}], [['x', 'occupied']]),
                ([{'row': 5, 'column': 0, 'row_span': 1, 'col_span': 1}], [['x', '']]),
            ]:
                with self.subTest(merges=merges):
                    spec = self._write_spec(root)
                    value = json.loads(spec.read_text(encoding='utf-8'))
                    value['sections'][1]['tables'] = [{
                        'headers': ['A', 'B'], 'rows': rows,
                        'merges': merges,
                    }]
                    spec.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
                    with self.assertRaises(ValueError):
                        build_docx.build(spec, root / ('invalid-' + str(len(merges)) + '.docx'), root)

    def test_merged_anchor_rebuild_keeps_newlines_and_span_width(self):
        from docx import Document
        from docx.oxml.ns import qn

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            spec = self._write_spec(root)
            value = json.loads(spec.read_text(encoding='utf-8'))
            value['sections'][1]['tables'] = [{
                'headers': [['标题\n第二行', '', '结果']],
                'rows': [['', '', '通过']],
                'widths_cm': [3, 5, 4],
                'merges': [{'row': 0, 'column': 0, 'row_span': 1, 'col_span': 2}],
            }]
            spec.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
            output = root / 'merged-anchor.docx'
            build_docx.build(spec, output, root)
            table = Document(output).tables[0]
            anchor = table.cell(0, 0)
            self.assertEqual(anchor.text, '标题\n第二行')
            self.assertEqual(len(anchor.paragraphs), 1)
            tcw = anchor._tc.tcPr.find(qn('w:tcW'))
            self.assertIsNotNone(tcw)
            self.assertAlmostEqual(anchor.width.cm, 8, places=1)
            self.assertEqual(tcw.get(qn('w:type')), 'dxa')
            self.assertTrue(table.rows[0]._tr.xpath('.//w:gridSpan[@w:val="2"]'))

    def test_fixed_form_table_can_omit_header_and_serialize_markdown(self):
        from docx import Document
        from fixed_form_tables import effective_grid, table_to_markdown

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            spec = self._write_spec(root)
            value = json.loads(spec.read_text(encoding='utf-8'))
            table = {
                'headers': [], 'header_rows': 0,
                'rows': [['字段', ''], ['', '']],
                'merges': [{'row': 0, 'column': 0, 'row_span': 2, 'col_span': 1}],
                'missing_cells': [{'row': 1, 'column': 1, 'label': '填写值'}],
            }
            value['sections'][1]['tables'] = [table]
            spec.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
            output = root / 'no-header.docx'
            build_docx.build(spec, output, root)
            document = Document(output)
            self.assertFalse(document.tables[0].rows[0]._tr.xpath('.//w:tblHeader'))
            self.assertEqual(effective_grid(table), [['字段', ''], ['字段', '【待补充：填写值】']])
            markdown = table_to_markdown(table)
            self.assertIn('rowspan="2"', markdown)
            self.assertIn('<mark>【待补充：填写值】</mark>', markdown)

            pipe = table_to_markdown({'headers': ['A', 'B'], 'rows': [['x|y', '<unsafe>']]})
            self.assertIn('| x&#124;y | &lt;unsafe&gt; |', pipe)
            self.assertNotIn('<unsafe>', pipe)

    def test_sections_set_a4_orientation_and_use_local_content_width(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            spec = self._write_spec(root)
            value = json.loads(spec.read_text(encoding='utf-8'))
            value['sections'][0]['orientation'] = 'portrait'
            value['sections'][0]['tables'] = []
            value['sections'][1]['orientation'] = 'landscape'
            value['sections'][1]['tables'][0]['widths_cm'] = [12, 12]
            spec.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
            output = root / 'orientations.docx'
            build_docx.build(spec, output, root)
            document = Document(output)
            self.assertEqual(len(document.sections), 2)
            self.assertTrue(document.sections[0].page_height > document.sections[0].page_width)
            self.assertTrue(document.sections[1].page_width > document.sections[1].page_height)
            self.assertAlmostEqual(document.sections[1].page_width.cm, 29.7, places=1)
            self.assertAlmostEqual(document.tables[0].columns[0].width.cm, 12, places=1)

    def test_table_level_footnotes_and_signatures_stay_between_two_tables(self):
        from docx.oxml.ns import qn

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            spec = self._write_spec(root)
            value = json.loads(spec.read_text(encoding='utf-8'))
            value['sections'][1]['paragraphs'] = []
            value['sections'][1]['images'] = []
            value['sections'][1]['tables'] = [
                {'headers': ['字段', '值'], 'rows': [['A', '1']],
                 'paragraphs_before': ['表一脚注'],
                 'paragraphs_after': ['表一签署区']},
                {'headers': ['字段', '值'], 'rows': [['B', '2']],
                 'paragraphs_before': ['表二脚注'],
                 'paragraphs_after': ['表二签署区']},
            ]
            spec.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
            output = root / 'table-order.docx'
            build_docx.build(spec, output, root)
            from docx import Document
            document = Document(output)
            body_order = []
            for child in document.element.body.iterchildren():
                if child.tag == qn('w:tbl'):
                    body_order.append('TABLE')
                elif child.tag == qn('w:p'):
                    text = ''.join(node.text or '' for node in child.iter(qn('w:t')))
                    if text in {'表一脚注', '表一签署区', '表二脚注', '表二签署区'}:
                        body_order.append(text)
            self.assertEqual(body_order, ['表一脚注', 'TABLE', '表一签署区',
                                          '表二脚注', 'TABLE', '表二签署区'])

    def test_invalid_export_options_and_image_bounds_fail_before_output(self):
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aZ1kAAAAASUVORK5CYII=')
        invalid_configs = [
            {'unexpected': {}},
            {'export': {}},
            {'style': {'unknown_style_option': 1}},
            {'style': {'body_size_pt': 0}},
            {'cover': None},
            {'toc': None},
            {'footer': None},
            {'cover': {'typography': None}},
            {'toc': {'enabled': True, 'levels': 10}},
            {'cover': {'metadata_rows': [{'label': '重复', 'value': '1'},
                                         {'label': '重复', 'value': '2'}]}},
            {'header': {'alignment': 'diagonal'}},
            {'toc': {'typography': {'unknown': 1}}},
        ]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            spec_path = self._write_spec(root)
            for index, config in enumerate(invalid_configs):
                with self.subTest(config=config):
                    settings = root / f'invalid-{index}.json'
                    settings.write_text(json.dumps(config), encoding='utf-8')
                    output = root / f'invalid-{index}.docx'
                    with self.assertRaises(ValueError):
                        build_docx.build(spec_path, output, root, settings)
                    self.assertFalse(output.exists())

            duplicate = root / 'duplicate.json'
            duplicate.write_text('{"toc": {}, "toc": {}}', encoding='utf-8')
            with self.assertRaises(ValueError):
                build_docx.build(spec_path, root / 'duplicate.docx', root, duplicate)

            (root / 'figure.png').write_bytes(png)
            too_wide = self._config()
            too_wide['cover']['image'] = {'path': 'figure.png', 'width_cm': 99}
            wide_path = root / 'wide.json'
            wide_path.write_text(json.dumps(too_wide), encoding='utf-8')
            with self.assertRaises(ValueError):
                build_docx.build(spec_path, root / 'wide.docx', root, wide_path)

            image_spec = self._write_spec(root / 'image-spec', images=[{'path': '../figure.png', 'width_cm': 2}])
            with self.assertRaises(ValueError):
                build_docx.build(image_spec, root / 'image.docx', root / 'image-spec')

            embedded = self._config()
            embedded_spec = self._write_spec(root / 'embedded-conflict', embedded=embedded)
            conflict = root / 'conflict.json'
            conflict.write_text('{}', encoding='utf-8')
            with self.assertRaises(ValueError):
                build_docx.build(embedded_spec, root / 'conflict.docx', root / 'embedded-conflict', conflict)


if __name__ == '__main__':
    unittest.main()
