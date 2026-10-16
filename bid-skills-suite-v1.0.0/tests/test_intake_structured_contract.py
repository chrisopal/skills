"""Exercise the actual structured extractor against the public source contract."""
import base64
import copy
import io
import json
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import bidkit  # noqa: E402
import extract_sources  # noqa: E402
from validate_output import validate  # noqa: E402
from test_knowledge_formats import _pptx_package  # noqa: E402


def _schema():
    return json.loads((ROOT / "skills/bid-source-intake/assets/output.schema.json").read_text())


def _payload(tmp_path, pptx=False):
    tmp_path.mkdir(parents=True, exist_ok=True)
    image = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    source = tmp_path / ("source.pptx" if pptx else "source.docx")
    if pptx:
        _pptx_package(source, image)
    else:
        doc = Document()
        doc.add_paragraph("Source with structured content")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
        table.cell(1, 0).text = "Field"
        table.cell(1, 1).text = "Value"
        doc.add_picture(io.BytesIO(image))
        doc.save(source)
    project = tmp_path / "project"
    bidkit.init_project(project, "CONTRACT-TEST")
    bidkit.register_source(project, source, "main")
    return extract_sources.extract(project, "CONTRACT-TEST")


def test_docx_table_and_embedded_image_match_contract(tmp_path):
    payload = _payload(tmp_path)
    assert any("table" in b for b in payload["data"]["blocks"])
    assert any("asset" in b for b in payload["data"]["blocks"])
    assert validate(payload, _schema()) == []


def test_pptx_notes_tables_and_images_match_contract(tmp_path):
    payload = _payload(tmp_path, pptx=True)
    assert any(b.get("is_note") for b in payload["data"]["blocks"])
    assert validate(payload, _schema()) == []


def test_structured_metadata_still_rejects_invalid_hash_and_unknown_fields(tmp_path):
    payload = _payload(tmp_path)
    corrupt = copy.deepcopy(payload)
    asset = next(b["asset"] for b in corrupt["data"]["blocks"] if "asset" in b)
    asset["sha256"] = "not-a-hash"
    assert any("sha256" in error for error in validate(corrupt, _schema()))
    corrupt = copy.deepcopy(payload)
    corrupt["data"]["blocks"][0]["unsafe_instruction"] = "untrusted source data"
    assert any("unsafe_instruction" in error for error in validate(corrupt, _schema()))


def test_structured_metadata_rejects_inconsistent_kind_fields(tmp_path):
    payload = _payload(tmp_path)

    corrupt = copy.deepcopy(payload)
    table = next(block for block in corrupt["data"]["blocks"] if "table" in block)
    table["content_kind"] = "text"
    assert validate(corrupt, _schema())

    corrupt = copy.deepcopy(payload)
    asset = next(block for block in corrupt["data"]["blocks"] if "asset" in block)
    asset["content_kind"] = "table"
    assert validate(corrupt, _schema())

    notes_payload = _payload(tmp_path / "pptx", pptx=True)
    corrupt = copy.deepcopy(notes_payload)
    note = next(block for block in corrupt["data"]["blocks"] if block.get("is_note"))
    note["is_note"] = False
    assert validate(corrupt, _schema())
