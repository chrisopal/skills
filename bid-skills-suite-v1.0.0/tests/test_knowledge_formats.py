import hashlib
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import bidkit  # noqa: E402
import extract_sources  # noqa: E402
import knowledge_formats  # noqa: E402


def _pptx_package(path: Path, image: bytes) -> None:
    slide = """<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld><p:spTree><p:nvGrpSpPr/><p:grpSpPr/><p:sp><p:nvSpPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>Slide body</a:t></a:r></a:p></p:txBody></p:sp><p:graphicFrame><p:nvGraphicFramePr/><p:xfrm/><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table"><a:tbl><a:tblGrid><a:gridCol/><a:gridCol/></a:tblGrid><a:tr><a:tc><a:txBody><a:p><a:r><a:t>Head</a:t></a:r></a:p></a:txBody></a:tc><a:tc><a:txBody><a:p><a:r><a:t>Value</a:t></a:r></a:p></a:txBody></a:tc></a:tr><a:tr><a:tc><a:txBody><a:p><a:r><a:t>A</a:t></a:r></a:p></a:txBody></a:tc><a:tc><a:txBody><a:p><a:r><a:t>B</a:t></a:r></a:p></a:txBody></a:tc></a:tr></a:tbl></a:graphicData></a:graphic></p:graphicFrame></p:spTree></p:cSld></p:sld>"""
    rels = """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide" Target="../notesSlides/notesSlide1.xml"/></Relationships>"""
    presentation = """<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst></p:presentation>"""
    presentation_rels = """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/></Relationships>"""
    notes = """<p:notes xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:sp><p:txBody><a:p><a:r><a:t>Speaker note</a:t></a:r></a:p></p:txBody></p:sp></p:notes>"""
    with zipfile.ZipFile(path, "w") as package:
        package.writestr("ppt/presentation.xml", presentation)
        package.writestr("ppt/_rels/presentation.xml.rels", presentation_rels)
        package.writestr("ppt/slides/slide1.xml", slide)
        package.writestr("ppt/slides/_rels/slide1.xml.rels", rels)
        package.writestr("ppt/notesSlides/notesSlide1.xml", notes)
        package.writestr("ppt/media/image1.png", image)


def _ordered_pptx(path: Path) -> None:
    def slide(text: str) -> str:
        return f'''<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld><p:spTree><p:nvGrpSpPr/><p:grpSpPr/><p:sp><p:nvSpPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>'''
    presentation = '''<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="rId2"/><p:sldId id="257" r:id="rId1"/></p:sldIdLst></p:presentation>'''
    relationships = '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide2.xml"/></Relationships>'''
    with zipfile.ZipFile(path, "w") as package:
        package.writestr("ppt/presentation.xml", presentation)
        package.writestr("ppt/_rels/presentation.xml.rels", relationships)
        package.writestr("ppt/slides/slide1.xml", slide("file one"))
        package.writestr("ppt/slides/slide2.xml", slide("file two"))


class KnowledgeFormatTests(unittest.TestCase):
    def test_grouped_shapes_are_read_and_chart_coverage_is_partial(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / 'grouped.pptx'
            _ordered_pptx(source)
            with zipfile.ZipFile(source) as package:
                files = {name: package.read(name) for name in package.namelist()}
            slide = files['ppt/slides/slide2.xml'].decode()
            group = '''<p:grpSp><p:grpSp><p:sp><p:txBody><a:p><a:r><a:t>Grouped critical clause</a:t></a:r></a:p></p:txBody></p:sp></p:grpSp></p:grpSp><p:graphicFrame><a:graphic><a:graphicData uri="chart"/></a:graphic></p:graphicFrame>'''
            files['ppt/slides/slide2.xml'] = slide.replace('</p:spTree>', group + '</p:spTree>').encode()
            with zipfile.ZipFile(source, 'w') as package:
                for name, data in files.items():
                    package.writestr(name, data)
            result = knowledge_formats.extract_pptx(source, Path(tmp) / 'assets/knowledge', 'SRC-001')
            self.assertIn('Grouped critical clause', [row['text'] for row in result['blocks']])
            self.assertTrue(result['review'])
            self.assertTrue(any('图形对象' in row for row in result['warnings']))

    def test_pptx_uses_presentation_relationship_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "ordered.pptx"
            _ordered_pptx(source)
            parsed = knowledge_formats.extract_pptx(source, Path(tmp) / "assets" / "knowledge", "SRC-001")
            self.assertEqual([item["text"] for item in parsed["blocks"]], ["file two", "file one"])
            self.assertEqual(parsed["warnings"], [])

    def test_pptx_missing_presentation_metadata_is_explicit_partial(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "partial.pptx"
            with zipfile.ZipFile(source, "w") as package:
                package.writestr("ppt/slides/slide1.xml", '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld><p:spTree/></p:cSld></p:sld>')
            parsed = knowledge_formats.extract_pptx(source, Path(tmp) / "assets" / "knowledge", "SRC-001")
            self.assertTrue(parsed["review"])
            self.assertTrue(any("presentation.xml" in warning for warning in parsed["warnings"]))

    def test_pptx_hmerge_and_vmerge_are_retained(self):
        element = ElementTree.fromstring('''
            <p:graphicFrame xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
              xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <a:graphic><a:graphicData><a:tbl>
                <a:tblGrid><a:gridCol/><a:gridCol/></a:tblGrid>
                <a:tr><a:tc><a:txBody><a:p><a:r><a:t>H</a:t></a:r></a:p></a:txBody></a:tc><a:tc hMerge="1"/></a:tr>
                <a:tr><a:tc rowSpan="2"><a:txBody><a:p><a:r><a:t>Left</a:t></a:r></a:p></a:txBody></a:tc><a:tc><a:txBody><a:p><a:r><a:t>R1</a:t></a:r></a:p></a:txBody></a:tc></a:tr>
                <a:tr><a:tc vMerge="1"/><a:tc><a:txBody><a:p><a:r><a:t>R2</a:t></a:r></a:p></a:txBody></a:tc></a:tr>
              </a:tbl></a:graphicData></a:graphic>
            </p:graphicFrame>''')
        table, _ = knowledge_formats._pptx_table(element)
        self.assertEqual(table["merges"], [
            {"row": 0, "column": 0, "row_span": 1, "col_span": 2},
            {"row": 1, "column": 0, "row_span": 2, "col_span": 1},
        ])

    def test_docx_table_keeps_cells_and_merge_shape(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "merged"
        table.cell(0, 1).text = ""
        table.cell(1, 0).text = "left"
        table.cell(1, 1).text = "right"
        table.cell(0, 0).merge(table.cell(0, 1))
        spec, text = knowledge_formats.docx_table_spec(table)
        self.assertEqual(spec["headers"], [["merged", ""]])
        self.assertEqual(spec["rows"], [["left", "right"]])
        self.assertEqual(spec["merges"], [{"row": 0, "column": 0, "row_span": 1, "col_span": 2}])
        self.assertIn("merged", text)

    def test_pptx_preserves_slide_table_notes_and_image_hash(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "sample.pptx"
            image = b"PNG fixture bytes"
            _pptx_package(source, image)
            parsed = knowledge_formats.extract_pptx(source, root / "assets" / "knowledge", "SRC-001")
            self.assertEqual(parsed["warnings"], [])
            self.assertEqual([item["content_kind"] for item in parsed["blocks"]], ["text", "table", "asset", "notes"])
            table = next(item for item in parsed["blocks"] if item["content_kind"] == "table")["table"]
            self.assertEqual(table["headers"], [["Head", "Value"]])
            asset = next(item for item in parsed["blocks"] if item["content_kind"] == "asset")["asset"]
            self.assertEqual(asset["sha256"], hashlib.sha256(image).hexdigest())
            self.assertTrue((root / asset["relative_path"]).is_file())
            note = next(item for item in parsed["blocks"] if item["content_kind"] == "notes")
            self.assertTrue(note["is_note"])

    def test_asset_path_rejects_symlink_and_svg_is_file_link_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            outside = root / "outside"
            outside.mkdir()
            (root / "assets").symlink_to(outside, target_is_directory=True)
            with self.assertRaisesRegex(ValueError, "符号链接"):
                knowledge_formats._write_asset(root / "assets" / "knowledge", "SRC-001", "x", "x.png", b"x")

            safe = root / "safe"
            (safe / "assets" / "knowledge").mkdir(parents=True)
            asset = knowledge_formats._write_asset(safe / "assets" / "knowledge", "SRC-001", "x", "x.svg", b"<svg/>")
            self.assertFalse(asset["safe_bitmap_preview"])
            self.assertEqual(asset["delivery"], "file_link")

    def test_extract_sources_assigns_legacy_fields_to_structured_blocks(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            project = root / "project"
            bidkit.init_project(project, "FORMAT-TEST")
            source = root / "sample.pptx"
            _pptx_package(source, b"image")
            registered = bidkit.register_source(project, source, "supplier")
            result = extract_sources.extract(project, "FORMAT-TEST", source_ids=[registered["source_id"]])
            structured = [item for item in result["data"]["blocks"] if item.get("content_kind") == "table"]
            self.assertEqual(len(structured), 1)
            self.assertTrue(all(item["source_id"] == registered["source_id"] and item["text"] for item in structured))
            self.assertEqual(result["data"]["coverage"], "partial")


if __name__ == "__main__":
    unittest.main()
