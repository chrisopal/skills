"""Customer body typography must survive actual DOCX generation."""
import importlib.util
import base64
import json
import struct
import sys
import tempfile
import unittest
import zlib
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import build_docx


def _solid_png(width, height):
    def chunk(kind, payload):
        return (struct.pack('>I', len(payload)) + kind + payload
                + struct.pack('>I', zlib.crc32(kind + payload) & 0xffffffff))

    rows = b''.join(b'\x00' + b'\x22\x88\xaa' * width for _ in range(height))
    return (b'\x89PNG\r\n\x1a\n'
            + chunk(b'IHDR', struct.pack('>IIBBBBB', width, height, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(rows))
            + chunk(b'IEND', b''))


@unittest.skipUnless(importlib.util.find_spec('docx'), 'Declared optional python-docx not installed')
class WorkingDocxTests(unittest.TestCase):
    def test_adjacent_tables_have_a_separator_for_office_roundtrip(self):
        from docx import Document
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            table = {'headers': ['Field'], 'rows': [['Value']]}
            spec = {'title': 'Two tables', 'sections': [{'title': 'Section', 'tables': [table, table]}]}
            (root/'spec.json').write_text(json.dumps(spec))
            build_docx.build(root/'spec.json', root/'draft.docx', root)
            document = Document(root/'draft.docx')
            self.assertEqual(len(document.tables), 2)
            self.assertEqual(document.tables[0]._tbl.getnext().tag, qn('w:p'))
            self.assertEqual(document.tables[0]._tbl.getnext().getnext().tag, qn('w:tbl'))

    def test_layout_template_survives_real_document(self):
        from docx import Document
        from docx.oxml.ns import qn
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = json.loads((ROOT/'assets/layout-templates/general-bid-v1.json').read_text())
            style = template['style']
            (root/'figure.png').write_bytes(base64.b64decode(
                'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aZ1kAAAAASUVORK5CYII='))
            spec = {'title': '版式测试', 'style': style, 'sections': [{
                'title': '一级标题', 'paragraphs': ['中文正文第一段。', '  ', '', '中文正文第二段。'],
                'tables': [{'headers': ['要求'], 'rows': [['待确认']]}],
                'images': [{'path': 'figure.png', 'width_cm': 4, 'caption': '方案示意图'}]}]}
            (root/'spec.json').write_text(json.dumps(spec))
            build_docx.build(root/'spec.json', root/'draft.docx', root)
            doc = Document(root/'draft.docx')
            body = next(p for p in doc.paragraphs if p.text == '中文正文第一段。')
            self.assertEqual(body.paragraph_format.first_line_indent.pt, 24)
            self.assertEqual(body.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
            self.assertEqual(doc.styles['Normal'].paragraph_format.space_before.pt, 0)
            self.assertEqual(doc.styles['Normal'].paragraph_format.space_after.pt, 0)
            self.assertEqual(doc.styles['Normal'].font.size.pt, 12)
            self.assertEqual(doc.styles['Normal'].paragraph_format.line_spacing, 1.5)
            self.assertEqual(doc.styles['Normal'].element.rPr.rFonts.get(qn('w:eastAsia')), 'SimSun')
            self.assertEqual(doc.styles['Heading 1'].font.size.pt, style['heading_sizes_pt'][0])
            self.assertEqual(doc.styles['Heading 1'].paragraph_format.space_before.pt, 12)
            self.assertEqual(doc.styles['Heading 1'].paragraph_format.space_after.pt, 6)
            cell = doc.tables[0].cell(1, 0).paragraphs[0]
            self.assertEqual(cell.paragraph_format.first_line_indent.pt, 0)
            self.assertEqual(cell.runs[0].font.size.pt, style['table_size_pt'])
            self.assertEqual(cell.paragraph_format.space_before.pt, 0)
            self.assertEqual(cell.paragraph_format.space_after.pt, 0)
            self.assertAlmostEqual(cell.paragraph_format.line_spacing,
                                   style.get('table_line_spacing', style['body_line_spacing']), places=2)
            self.assertEqual(cell.runs[0]._r.rPr.rFonts.get(qn('w:eastAsia')), style['body_east_asia_font'])
            table_xml = doc.tables[0]._tbl.xml
            self.assertIn('w:tblBorders', table_xml)
            self.assertIn('w:fill="D9EAF7"', table_xml)
            caption = next(p for p in doc.paragraphs if p.text == '方案示意图')
            self.assertEqual(caption.paragraph_format.first_line_indent.pt, 0)
            self.assertEqual(caption.runs[0].font.size.pt, style['caption_size_pt'])
            self.assertEqual(caption.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertTrue(caption.paragraph_format.keep_together)
            image_paragraph = next(p for p in doc.paragraphs
                                   if p._p.findall('.//' + qn('w:drawing')))
            self.assertEqual(image_paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertTrue(image_paragraph.paragraph_format.keep_with_next)
            self.assertAlmostEqual(doc.sections[0].left_margin.cm, 2.5, places=2)
            self.assertFalse([p for p in doc.paragraphs if not p.text.strip()
                              and not p._p.findall('.//' + qn('w:drawing'))])

    def test_customer_table_visual_overrides_and_spacing(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root/'figure.png').write_bytes(_solid_png(2, 2))
            style = {'body_size_pt': 11, 'body_line_spacing': 1.5,
                     'table_line_spacing': 1.2, 'caption_line_spacing': 1.3,
                     'table_border_color': None, 'table_header_fill': None,
                     'table_alternate_fill': None, 'table_cell_padding': None}
            spec = {'title': '自定义表格', 'style': style, 'sections': [{
                'title': '表格', 'tables': [{'headers': ['字段'], 'rows': [['内容']]}],
                'images': [{'path': 'figure.png', 'width_cm': 2, 'caption': '图注'}]}]}
            (root/'spec.json').write_text(json.dumps(spec))
            build_docx.build(root/'spec.json', root/'draft.docx', root)
            doc = Document(root/'draft.docx')
            cell_paragraph = doc.tables[0].cell(1, 0).paragraphs[0]
            self.assertAlmostEqual(cell_paragraph.paragraph_format.line_spacing, 1.2, places=2)
            caption = next(p for p in doc.paragraphs if p.text == '图注')
            self.assertAlmostEqual(caption.paragraph_format.line_spacing, 1.3, places=2)
            table_xml = doc.tables[0]._tbl.xml
            self.assertNotIn('w:tblBorders', table_xml)
            self.assertNotIn('w:shd', table_xml)
            self.assertNotIn('w:tcMar', table_xml)

    def test_inline_image_group_rejects_tall_portrait(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root/'portrait.png').write_bytes(_solid_png(2, 100))
            spec = {'title': '高图', 'sections': [{'title': '图',
                    'images': [{'path': 'portrait.png', 'width_cm': 4, 'caption': '图注'}]}]}
            (root/'spec.json').write_text(json.dumps(spec))
            with self.assertRaises(ValueError):
                build_docx.build(root/'spec.json', root/'draft.docx', root)
            self.assertFalse((root/'draft.docx').exists())

    def test_inline_image_group_rejects_oversize_explicit_height(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root/'figure.png').write_bytes(_solid_png(10, 10))
            spec = {'title': '高图', 'sections': [{'title': '图',
                    'images': [{'path': 'figure.png', 'width_cm': 4, 'height_cm': 26.2,
                                'caption': '图注'}]}]}
            (root/'spec.json').write_text(json.dumps(spec))
            with self.assertRaises(ValueError):
                build_docx.build(root/'spec.json', root/'draft.docx', root)
            self.assertFalse((root/'draft.docx').exists())

    def test_near_full_page_image_reserves_heading_height(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root/'figure.png').write_bytes(_solid_png(10, 10))
            spec = {'title': '图示', 'sections': [{'title': '本节架构图',
                    'images': [{'path': 'figure.png', 'width_cm': 4,
                                'height_cm': 24, 'caption': '图注'}]}]}
            (root/'spec.json').write_text(json.dumps(spec))
            with self.assertRaisesRegex(ValueError, '标题'):
                build_docx.build(root/'spec.json', root/'draft.docx', root)
            self.assertFalse((root/'draft.docx').exists())

    def test_tender_can_override_all_text_font_sizes(self):
        from docx import Document
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template = json.loads((ROOT/'assets/layout-templates/general-bid-v1.json').read_text())
            style = {**template['style'], 'title_size_pt': 12, 'subtitle_size_pt': 12,
                     'heading_sizes_pt': [12, 12, 12], 'heading_east_asia_font': 'SimSun',
                     'table_size_pt': 12, 'caption_size_pt': 12}
            spec = {'title': '投标文件', 'subtitle': '候选工作稿', 'style': style,
                    'sections': [{'title': '方案', 'paragraphs': ['正文']} ]}
            (root/'spec.json').write_text(json.dumps(spec))
            build_docx.build(root/'spec.json', root/'draft.docx', root)
            doc = Document(root/'draft.docx')
            for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
                self.assertEqual(doc.styles[name].font.size.pt, 12, name)
                self.assertEqual(doc.styles[name].element.rPr.find(qn('w:szCs')).get(qn('w:val')), '24')
            self.assertFalse([attribute for fonts in doc.styles.element.iter(qn('w:rFonts'))
                              for attribute in fonts.attrib if attribute.lower().endswith('theme')])

    def test_invalid_template_fields_rejected_without_output(self):
        invalid = [{'body_first_line_indent_chars': True}, {'body_space_after_pt': float('nan')},
                   {'heading_sizes_pt': [16, 14]}, {'table_size_pt': -1},
                   {'table_line_spacing': 0}, {'table_header_fill': 'blue'},
                   {'table_cell_padding': {'top': 1, 'start': 1, 'bottom': -1, 'end': 1}},
                   {'page_margins_cm': {'top': 2.5, 'bottom': 2.5, 'left': 2.5, 'right': 2.5, 'extra': 1}},
                   {'page_margins_cm': {'top': 2.5, 'bottom': 2.5, 'left': True, 'right': 2.5}}]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for i, style in enumerate(invalid):
                with self.subTest(style=style):
                    (root/'spec.json').write_text(json.dumps({'title': '测试', 'sections': [], 'style': style}))
                    with self.assertRaises(ValueError):
                        build_docx.build(root/'spec.json', root/f'draft-{i}.docx', root)
                    self.assertFalse((root/f'draft-{i}.docx').exists())

    def test_customer_font_and_spacing(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            spec = {'title': 'Test draft', 'style': {'body_size_pt': 12,
                    'body_line_spacing': 1.5, 'body_east_asia_font': 'SimSun'},
                    'sections': [{'title': 'Chapter', 'paragraphs': ['中文正文']} ]}
            (root / 'spec.json').write_text(json.dumps(spec))
            build_docx.build(root / 'spec.json', root / 'draft.docx', root)
            doc = Document(root / 'draft.docx')
            self.assertEqual(doc.styles['Normal'].font.size.pt, 12)
            self.assertEqual(doc.styles['Normal'].paragraph_format.line_spacing, 1.5)
            self.assertIn('中文正文', '\n'.join(p.text for p in doc.paragraphs))

    def test_invalid_style_fails_before_output(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root/'spec.json').write_text(json.dumps({'title':'draft','sections':[],
                                                     'style':{'body_size_pt': -1}}))
            with self.assertRaises(ValueError):
                build_docx.build(root/'spec.json', root/'draft.docx', root)
            self.assertFalse((root/'draft.docx').exists())

    def test_response_column_widths_survive_in_grid_and_cells(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            spec = {'title': '响应表', 'sections': [{'title': '逐项响应', 'tables': [
                {'headers': ['要求', '响应', '状态'], 'rows': [['原文', '较长响应', '待核验']],
                 'widths_cm': [4, 9, 3]}]}]}
            (root/'spec.json').write_text(json.dumps(spec))
            build_docx.build(root/'spec.json', root/'draft.docx', root)
            table = Document(root/'draft.docx').tables[0]
            self.assertFalse(table.autofit)
            for row in table.rows:
                self.assertGreater(row.cells[1].width, row.cells[0].width * 2)
            spec['sections'][0]['tables'][0]['widths_cm'] = [4, 20, 3]
            (root/'invalid.json').write_text(json.dumps(spec))
            with self.assertRaises(ValueError):
                build_docx.build(root/'invalid.json', root/'invalid.docx', root)
            self.assertFalse((root/'invalid.docx').exists())


if __name__ == '__main__':
    unittest.main()
