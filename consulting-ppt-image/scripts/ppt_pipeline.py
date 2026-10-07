#!/usr/bin/env python3
"""Local, manifest-driven production control for one-image-per-slide PPTX.

Does NOT generate images, read image text, browse sources or authenticate reviewers.
A recorded visual review is an attestation, not an automated semantic guarantee.
Requires Python 3.10+, Pillow, python-pptx and jsonschema.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import shutil
import sys
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
CHECKS = ("single_page", "text_fidelity", "visual_style", "layout_readable",
          "evidence_and_labels", "page_number", "no_invented_facts")

class PipelineError(Exception):
    pass

def now() -> str:
    return datetime.now(timezone.utc).isoformat()

def canonical(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True,
                      separators=(",", ":")).encode("utf-8")

def digest(value: Any) -> str:
    return hashlib.sha256(canonical(value)).hexdigest()

def file_digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def load_json(path: Path) -> dict:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise PipelineError(f"Cannot read JSON {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise PipelineError(f"Expected JSON object: {path}")
    return value

def save_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n",
                   encoding="utf-8")
    tmp.replace(path)

def safe_path(base: Path, rel: str) -> Path:
    p = (base / rel).resolve()
    if not p.is_relative_to(base.resolve()):
        raise PipelineError(f"Path escapes project: {rel}")
    return p

def schema_errors(pack: dict) -> list[str]:
    try:
        from jsonschema import Draft202012Validator
    except ImportError as exc:
        raise PipelineError("Install dependencies: pip install -r requirements.txt") from exc
    schema = load_json(ROOT / "schemas/slide-content-pack.schema.json")
    return [f"schema {'.'.join(map(str, e.absolute_path))}: {e.message}"
            for e in Draft202012Validator(schema).iter_errors(pack)]

def check_pack(pack: dict, stage: str = "structure", only: set[str] | None = None) -> dict:
    errors = schema_errors(pack)
    warnings: list[str] = []
    if errors:
        return {"errors": errors, "warnings": warnings}
    slides = sorted(pack["slides"], key=lambda s: s["order"])
    n = pack["page_budget"]
    if len(slides) != n:
        errors.append(f"page_budget={n}, but slides={len(slides)}")
    if [s["order"] for s in slides] != list(range(1, n + 1)):
        errors.append("Slide orders must be exactly 1..page_budget, without gaps or duplicates")
    ids = [s["slide_id"] for s in slides]
    if len(set(ids)) != len(ids):
        errors.append("Duplicate stable slide_id")
    section_ids = [s["id"] for s in pack["sections"]]
    if len(set(section_ids)) != len(section_ids):
        errors.append("Duplicate section id")
    observed = []
    for s in slides:
        if s["section"] not in section_ids:
            errors.append(f"{s['slide_id']}: undeclared section {s['section']}")
        if not observed or observed[-1] != s["section"]:
            observed.append(s["section"])
    if len(observed) != len(set(observed)):
        errors.append("A section is interrupted by another section (A-B-A)")
    expected = [sid for sid in section_ids if any(s["section"] == sid for s in slides)]
    if observed != expected:
        errors.append("Slide section order differs from declared section order")
    purposes: dict[str, str] = {}
    titles: dict[str, str] = {}
    for s in slides:
        for key, registry in [("purpose_id", purposes), ("title", titles)]:
            normalized = re.sub(r"\s+", "", s[key]).casefold()
            if normalized in registry:
                errors.append(f"Duplicate {key}: {registry[normalized]} / {s['slide_id']}")
            registry[normalized] = s["slide_id"]
        if s.get("story_role") == "closing" and s["order"] != n:
            errors.append(f"{s['slide_id']}: closing page occurs before the deck ends")
    evidence = {e["id"]: e for e in pack["citation_index"]}
    if len(evidence) != len(pack["citation_index"]):
        errors.append("Duplicate evidence id")
    for s in slides:
        sid = s["slide_id"]
        for eid in s["evidence_ids"]:
            if eid not in evidence:
                errors.append(f"{sid}: unknown evidence {eid}")
        node_ids = {node["id"] for node in s.get("diagram_spec", {}).get("nodes", [])}
        if len(node_ids) != len(s.get("diagram_spec", {}).get("nodes", [])):
            errors.append(f"{sid}: duplicate diagram node id")
        for edge in s.get("diagram_spec", {}).get("edges", []):
            if edge["from"] not in node_ids or edge["to"] not in node_ids:
                errors.append(f"{sid}: diagram edge points to an undeclared node")
        for dependency in s.get("depends_on_slides", []):
            if dependency not in ids or dependency == sid:
                errors.append(f"{sid}: invalid slide dependency {dependency}")
        if s.get("content_kind") == "proposal" and not s.get("visible_qualifier"):
            warnings.append(f"{sid}: proposal should have a visible qualifier, e.g. 规划建议")
        if s.get("illustrative_values") and not s.get("visible_qualifier"):
            errors.append(f"{sid}: illustrative values require a visible label")
        active = only is None or sid in only
        if stage in ("render", "final") and active:
            if s["review_status"] != "approved" or not pack.get("approval_ref"):
                errors.append(f"{sid}: content not approved / missing approval_ref")
            if s.get("content_complete", True) is False:
                errors.append(f"{sid}: outline-only content cannot be rendered as a complete slide")
            for eid in s["evidence_ids"]:
                e = evidence.get(eid, {})
                if e.get("verification_status") == "unverified":
                    errors.append(f"{sid}: unresolved evidence {eid}")
        benchmark = s.get("benchmark")
        if benchmark:
            related = [evidence[x] for x in s["evidence_ids"] if x in evidence]
            wrong = [e["id"] for e in related if e.get("entity_id") != benchmark["entity_id"]]
            if wrong:
                errors.append(f"{sid}: benchmark entity does not match evidence: {wrong}")
            screenshots = [e for e in related if e.get("type") == "official_screenshot"]
            if stage in ("render", "final") and active:
                if not screenshots or not any(e.get("verification_status") == "verified" for e in screenshots):
                    errors.append(f"{sid}: benchmark needs a verified original screenshot")
        points_len = sum(len(p) for p in s["supporting_points"])
        if points_len > 900:
            warnings.append(f"{sid}: high text density; review rather than shrink font")
    # Minimal deterministic schedule contract: date-only FS dependencies.
    schedule = pack.get("schedule", {})
    tasks = schedule.get("tasks", [])
    task_map = {t["id"]: t for t in tasks}
    milestones = {m["id"]: m for m in schedule.get("milestones", [])}
    if len(task_map) != len(tasks):
        errors.append("Duplicate schedule task id")
    if len(milestones) != len(schedule.get("milestones", [])):
        errors.append("Duplicate milestone id")
    dates: dict[str, tuple[date, date]] = {}
    for t in tasks:
        try:
            start, end = date.fromisoformat(t["start"]), date.fromisoformat(t["end"])
            dates[t["id"]] = (start, end)
            if start > end:
                errors.append(f"{t['id']}: start is after end")
        except ValueError:
            errors.append(f"{t['id']}: invalid ISO date")
    for t in tasks:
        for dep in t.get("depends_on", []):
            if dep not in task_map:
                errors.append(f"{t['id']}: unknown task dependency {dep}")
            elif dep in dates and t["id"] in dates and dates[dep][1] > dates[t["id"]][0]:
                errors.append(f"{t['id']}: FS dependency {dep} finishes after successor starts")
        for mid in t.get("must_finish_before", []):
            if mid not in milestones:
                errors.append(f"{t['id']}: unknown milestone {mid}")
                continue
            try:
                deadline = date.fromisoformat(milestones[mid]["date"])
                if t["id"] in dates and dates[t["id"]][1] > deadline:
                    errors.append(f"{t['id']}: finishes after required milestone {mid}")
            except ValueError:
                errors.append(f"{mid}: invalid milestone date")
    visited: set[str] = set()
    active_tasks: set[str] = set()
    def visit(tid: str) -> None:
        if tid in active_tasks:
            errors.append(f"Schedule dependency cycle at {tid}")
            return
        if tid in visited:
            return
        active_tasks.add(tid)
        for dep in task_map[tid].get("depends_on", []):
            if dep in task_map:
                visit(dep)
        active_tasks.remove(tid)
        visited.add(tid)
    for tid in task_map:
        visit(tid)
    layouts = [s["visual_type"] for s in slides]
    for i in range(len(layouts) - 2):
        if len(set(layouts[i:i+3])) == 1:
            warnings.append(f"Orders {i+1}-{i+3}: same layout family repeated; check monotony")
    return {"errors": errors, "warnings": warnings}

def require(report: dict) -> None:
    if report["errors"]:
        raise PipelineError("\n".join(report["errors"]))

def signature(pack: dict, style: dict, slide: dict) -> str:
    by_id = {s["slide_id"]: s for s in pack["slides"]}
    deps = {i: by_id[i] for i in slide.get("depends_on_slides", []) if i in by_id}
    return digest({"deck_id": pack["deck_id"], "slide": slide,
                   "total": pack["page_budget"], "style": style, "dependencies": deps,
                   "evidence": [e for e in pack["citation_index"] if e["id"] in slide["evidence_ids"]],
                   "schedule": pack.get("schedule") if slide["visual_type"] in ("gantt", "roadmap") else None,
                   "section_name": next(x["title"] for x in pack["sections"] if x["id"] == slide["section"])})

def get_project(project: Path) -> tuple[dict, dict, dict]:
    manifest = load_json(project / "deck-manifest.json")
    pack = load_json(project / "slide-content-pack.json")
    style = load_json(project / "style-profile.json")
    if digest(pack) != manifest["pack_sha256"]:
        raise PipelineError("Content pack changed outside the approved sync workflow")
    if digest(style) != manifest["style_sha256"]:
        raise PipelineError("Locked style changed; create a separately approved style revision")
    return manifest, pack, style

def picked(pack: dict, span: str | None) -> list[dict]:
    slides = sorted(pack["slides"], key=lambda s: s["order"])
    if span is None:
        return slides
    orders: set[int] = set()
    try:
        for item in span.split(","):
            a, b = map(int, item.split("-")) if "-" in item else (int(item), int(item))
            if a < 1 or a > b or b > pack["page_budget"]:
                raise ValueError()
            orders.update(range(a, b + 1))
    except ValueError as exc:
        raise PipelineError("Invalid range; use e.g. 1-10 or 21-30") from exc
    return [s for s in slides if s["order"] in orders]

def init_project(pack_path: Path, project: Path, style_path: Path) -> dict:
    pack, style = load_json(pack_path), load_json(style_path)
    if pack.get("theme_profile") != style.get("id"):
        raise PipelineError("Content pack theme_profile does not match selected style id")
    report = check_pack(pack)
    require(report)
    if project.exists() and any(project.iterdir()):
        raise PipelineError(f"Project not empty; refusing overwrite: {project}")
    for rel in ("prompts", "renders", "reviews", "output", "history"):
        (project / rel).mkdir(parents=True, exist_ok=True)
    save_json(project / "slide-content-pack.json", pack)
    save_json(project / "style-profile.json", style)
    m = {"schema_version": "1.0", "deck_id": pack["deck_id"],
         "pack_sha256": digest(pack), "style_sha256": digest(style),
         "created_at": now(), "pages": {s["slide_id"]: {"versions": [], "selected_version": None}
                                             for s in pack["slides"]}, "exports": []}
    save_json(project / "deck-manifest.json", m)
    save_json(project / "output/structure-audit.json", report)
    return {"project": str(project), "page_count": pack["page_budget"], **report}

def make_plan(project: Path, span: str | None, draft: bool = False) -> dict:
    m, pack, style = get_project(project)
    slides = picked(pack, span)
    report = check_pack(pack, "structure" if draft else "render", {s["slide_id"] for s in slides})
    require(report)
    queue = []
    ordered = sorted(pack["slides"], key=lambda s: s["order"])
    for s in slides:
        sid = s["slide_id"]
        version = max((x["version"] for x in m["pages"][sid]["versions"]), default=0) + 1
        sig = signature(pack, style, s)
        path = project / f"prompts/{sid}_v{version:03d}_{sig[:10]}.md"
        prev = ordered[s["order"] - 2]["title"] if s["order"] > 1 else "封面/开场"
        nxt = ordered[s["order"]]["title"] if s["order"] < len(ordered) else "全稿结束"
        src = [e for e in pack["citation_index"] if e["id"] in s["evidence_ids"]]
        prompt = f'''# {'草稿预览，不得执行图像生成' if draft else '单页生成任务'}：{sid}

只生成一张独立的横向PPT页面；不要多页拼图、联系表、长图或多画布。
本次仅处理实际第{s['order']}页 / 共{pack['page_budget']}页。
稳定页ID：{sid}；版本：v{version:03d}；内容签名：{sig}

## 本轮锁定样式
```json
{json.dumps(style, ensure_ascii=False, indent=2)}
```
样式参考只取版式、配色与图形语法，不复制其他项目名称、数值、企业、产品或工艺。
使用当前宿主真实的图像生成/编辑工具及其有效参数。工具支持上下文式指令时，将本任务置于上下文；不要臆造prompt字段、模型名、图片ID或API。

## 唯一页面内容
```json
{json.dumps(s, ensure_ascii=False, indent=2)}
```

## 在全稿中的作用
前一页：{prev}
本页唯一结论：{s['key_message']}
下一页：{nxt}
只用相邻页建立衔接，不把它们的内容一起画进本页。

## 证据与素材
```json
{json.dumps(src, ensure_ascii=False, indent=2)}
```
对标截图须读取对应原始素材，不能用生成图片冒充官方截图。概念图显著标“概念示意”。不得虚构ROI、百分比、政策结果、产能、认证或已实现成效。

## 画面验收
纯白页面背景；墨绿色标题、同一绿色系强调、深灰正文；无山水花纹、角落标语、装饰水印、四色轮换。
根据visual_type/layout_hint选择结构，不把所有页面强制改成卡片。
中文清楚，连线方向明确，结构和数字严格忠于内容；如需增删核心对象，退回内容评审。
完成后立即保存为独立图片，再用register登记，不复用imagegen.png作为成品路径。
'''
        if not path.exists():
            path.write_text(prompt, encoding="utf-8")
        queue.append({"slide_id": sid, "order": s["order"], "version": version,
                      "prompt": str(path), "render_signature": sig,
                      "status": "draft_only" if draft else "ready_for_host_image_tool"})
    save_json(project / "output/work-queue.json", {"created_at": now(), "items": queue})
    return {"count": len(queue), "queue": queue, **report}

def find_slide(pack: dict, sid: str) -> dict:
    for s in pack["slides"]:
        if s["slide_id"] == sid:
            return s
    raise PipelineError(f"Unknown slide_id {sid}; use manifest IDs, not printed page labels")

def image_info(path: Path, style: dict) -> tuple[int, int, str]:
    from PIL import Image
    try:
        with Image.open(path) as image:
            w, h = image.size
            fmt = image.format
            if getattr(image, "n_frames", 1) != 1:
                raise PipelineError("Animated/multiframe images are not supported")
            if fmt not in ("PNG", "JPEG"):
                raise PipelineError("Use a PNG or JPEG raster page, not SVG/PDF/GIF")
            image.verify()
    except (OSError, ValueError) as exc:
        raise PipelineError(f"Unreadable image {path}: {exc}") from exc
    cw, ch = style["canvas"]["aspect_ratio"]
    if abs((w/h)/(cw/ch) - 1) > style["canvas"].get("aspect_tolerance", 0.005):
        raise PipelineError(f"Image aspect ratio {w}x{h} differs from locked {cw}:{ch}")
    if w < style["canvas"].get("minimum_width", 1600):
        raise PipelineError(f"Image too small: width={w}. Do not upscale to claim native resolution")
    return w, h, fmt

def register_image(project: Path, sid: str, path: Path, version: int | None,
                   producer: str) -> dict:
    m, pack, style = get_project(project)
    s = find_slide(pack, sid)
    require(check_pack(pack, "render", {sid}))
    w, h, fmt = image_info(path, style)
    records = m["pages"][sid]["versions"]
    v = version if version is not None else max((x["version"] for x in records), default=0) + 1
    if v < 1 or any(x["version"] == v for x in records):
        raise PipelineError("Version already exists / invalid; original versions are immutable")
    rel = f"renders/{sid}_v{v:03d}.{'png' if fmt == 'PNG' else 'jpg'}"
    target = safe_path(project, rel)
    if target.exists():
        raise PipelineError(f"Refusing to overwrite existing raster {rel}")
    shutil.copyfile(path, target)
    record = {"version": v, "path": rel, "sha256": file_digest(target),
              "width": w, "height": h, "format": fmt,
              "render_signature": signature(pack, style, s), "producer": producer,
              "created_at": now(), "review": None, "approval": None}
    records.append(record)
    save_json(project / "deck-manifest.json", m)
    return {"slide_id": sid, **record}

def find_version(m: dict, sid: str, version: int) -> dict:
    for v in m["pages"].get(sid, {}).get("versions", []):
        if v["version"] == version:
            return v
    raise PipelineError(f"No registered artifact {sid} v{version}")

def record_review(project: Path, review_path: Path) -> dict:
    m, pack, style = get_project(project)
    r = load_json(review_path)
    for key in ("slide_id", "version", "artifact_sha256", "reviewer", "checks", "notes"):
        if key not in r:
            raise PipelineError(f"Review missing {key}")
    if not isinstance(r["reviewer"], str) or not r["reviewer"].strip():
        raise PipelineError("Reviewer identity/role must be recorded")
    v = find_version(m, r["slide_id"], r["version"])
    s = find_slide(pack, r["slide_id"])
    if v["sha256"] != r["artifact_sha256"] or file_digest(safe_path(project, v["path"])) != v["sha256"]:
        raise PipelineError("Review targets a different or modified image hash")
    if v["render_signature"] != signature(pack, style, s):
        raise PipelineError("Image has an obsolete content/style/order signature")
    if not isinstance(r["checks"], dict) or not all(isinstance(r["checks"].get(c), bool) for c in CHECKS):
        raise PipelineError("Review must contain an explicit boolean for all seven checks")
    passed = all(r["checks"][c] for c in CHECKS)
    if not passed:
        v["approval"] = None
        if m["pages"][r["slide_id"]]["selected_version"] == r["version"]:
            m["pages"][r["slide_id"]]["selected_version"] = None
    r["recorded_at"] = now()
    v["review"] = r
    save_json(project / f"reviews/{r['slide_id']}_v{r['version']:03d}.json", r)
    save_json(project / "deck-manifest.json", m)
    return {"status": "qa_attestation_recorded" if passed else "qa_failed",
            "slide_id": r["slide_id"], "version": r["version"],
            "errors": [] if passed else [f"Failed visual review check: {c}" for c in CHECKS if not r["checks"][c]]}

def select_image(project: Path, sid: str, version: int, confirmation: str) -> dict:
    m, pack, style = get_project(project)
    v = find_version(m, sid, version)
    if (not confirmation.strip() or not v.get("review")
            or not all(v["review"].get("checks", {}).get(c) is True for c in CHECKS)):
        raise PipelineError("Selection requires a completed visual review and confirmation reference")
    if v["render_signature"] != signature(pack, style, find_slide(pack, sid)):
        raise PipelineError("Cannot select an obsolete image version")
    if file_digest(safe_path(project, v["path"])) != v["sha256"]:
        raise PipelineError("Image bytes changed")
    v["approval"] = {"confirmation_ref": confirmation, "recorded_at": now()}
    m["pages"][sid]["selected_version"] = version
    save_json(project / "deck-manifest.json", m)
    return {"selected": sid, "version": version}

def audit(project: Path, stage: str = "final", span: str | None = None) -> dict:
    m, pack, style = get_project(project)
    slides = picked(pack, span)
    report = check_pack(pack, stage, {s["slide_id"] for s in slides})
    report.update({"stage": stage, "scope": "full_deck" if span is None else f"explicit_partial:{span}",
                   "expected_pages": len(slides), "selected_pages": 0,
                   "not_machine_verifiable": ["semantic duplicate meaning", "actual Chinese text in raster",
                       "collage vs single page", "visual style fidelity", "external evidence truth",
                       "whether recorded review/approval really occurred"]})
    hashes: dict[str, str] = {}
    for s in slides:
        sid = s["slide_id"]
        page = m["pages"].get(sid, {})
        selected = page.get("selected_version")
        if selected is None:
            if stage == "final":
                report["errors"].append(f"{sid}: no approved selected image")
            continue
        report["selected_pages"] += 1
        try:
            v = find_version(m, sid, selected)
            path = safe_path(project, v["path"])
            if not path.is_file() or file_digest(path) != v["sha256"]:
                report["errors"].append(f"{sid}: selected image missing or hash changed")
                continue
            image_info(path, style)
            if v["render_signature"] != signature(pack, style, s):
                report["errors"].append(f"{sid}: stale image after content/style/order revision")
            review = v.get("review") or {}
            if (not v.get("approval") or review.get("artifact_sha256") != v["sha256"]
                    or not all(review.get("checks", {}).get(c) is True for c in CHECKS)):
                report["errors"].append(f"{sid}: review / approval incomplete or mismatched")
            if v["sha256"] in hashes:
                report["errors"].append(f"Identical raster reused for {hashes[v['sha256']]} and {sid}")
            hashes[v["sha256"]] = sid
        except (PipelineError, OSError) as exc:
            report["errors"].append(str(exc))
    report["passed"] = not report["errors"]
    save_json(project / "output/audit-report.json", report)
    return report

def assemble(project: Path, output: Path, span: str | None = None) -> dict:
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    report = audit(project, "final", span)
    require(report)
    if output.exists():
        raise PipelineError(f"Export exists; choose a new version name: {output}")
    m, pack, style = get_project(project)
    slides = picked(pack, span)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = ("[PARTIAL] " if span else "") + pack["deck_id"]
    expected = []
    for s in slides:
        v = find_version(m, s["slide_id"], m["pages"][s["slide_id"]]["selected_version"])
        blob = safe_path(project, v["path"]).read_bytes()
        # Exact original bytes; contain-fit keeps aspect and never crops/stretches.
        scale = min(prs.slide_width/v["width"], prs.slide_height/v["height"])
        w, h = round(v["width"] * scale), round(v["height"] * scale)
        x, y = (prs.slide_width-w)//2, (prs.slide_height-h)//2
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(BytesIO(blob), x, y, width=w, height=h)
        provenance = {"slide_id": s["slide_id"], "order": s["order"],
                      "title": s["title"], "version": v["version"], "image_sha256": v["sha256"],
                      "evidence_ids": s["evidence_ids"], "speaker_notes": s.get("speaker_notes", ""),
                      "visible_qualifier": s.get("visible_qualifier", "")}
        slide.notes_slide.notes_text_frame.text = json.dumps(provenance, ensure_ascii=False, indent=2)
        expected.append(v["sha256"])
    output.parent.mkdir(parents=True, exist_ok=True)
    tmp = output.with_name(output.stem + ".building.pptx")
    try:
        prs.save(tmp)
        loaded = Presentation(tmp)
        if len(loaded.slides) != len(slides):
            raise PipelineError("Export slide count mismatch")
        for i, sl in enumerate(loaded.slides):
            shapes = list(sl.shapes)
            if len(shapes) != 1 or shapes[0].shape_type != MSO_SHAPE_TYPE.PICTURE:
                raise PipelineError(f"Export slide {i+1} is not exactly one image")
            if hashlib.sha256(shapes[0].image.blob).hexdigest() != expected[i]:
                raise PipelineError(f"Export slide {i+1} embedded bytes differ from manifest")
        tmp.replace(output)
    finally:
        if tmp.exists():
            tmp.unlink()
    receipt = {"created_at": now(), "path": str(output.resolve()),
               "sha256": file_digest(output), "slide_count": len(slides),
               "scope": report["scope"], "pack_sha256": m["pack_sha256"],
               "embedded_image_sha256": expected, "one_image_per_slide": True,
               "original_image_bytes_preserved": True}
    m["exports"].append(receipt)
    save_json(project / "deck-manifest.json", m)
    save_json(project / "output/export-receipt.json", receipt)
    return receipt

def sync_pack(project: Path, new_pack_path: Path, confirmation: str) -> dict:
    m, old, style = get_project(project)
    new = load_json(new_pack_path)
    require(check_pack(new))
    if new.get("theme_profile") != style.get("id"):
        raise PipelineError("Style changes require a new approved style project")
    if new["deck_id"] != old["deck_id"] or not confirmation.strip():
        raise PipelineError("Sync requires same deck_id and explicit change approval reference")
    if new["version"] == old["version"]:
        raise PipelineError("Increment upstream content-pack version for revisions")
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    save_json(project / f"history/{stamp}_pack.json", old)
    save_json(project / f"history/{stamp}_manifest.json", m)
    changes, pages = [], {}
    for s in new["slides"]:
        sid = s["slide_id"]
        entry = copy.deepcopy(m["pages"].get(sid, {"versions": [], "selected_version": None}))
        if entry["selected_version"] is not None:
            v = next(x for x in entry["versions"] if x["version"] == entry["selected_version"])
            if v["render_signature"] != signature(new, style, s):
                entry["selected_version"] = None
                changes.append(sid)
        elif not any(x["slide_id"] == sid for x in old["slides"]):
            changes.append(sid)
        pages[sid] = entry
    m["pages"] = pages
    m["pack_sha256"] = digest(new)
    m["last_sync"] = {"confirmation_ref": confirmation, "at": now(), "invalidated": changes}
    save_json(project / "slide-content-pack.json", new)
    save_json(project / "deck-manifest.json", m)
    return {"invalidated_or_new": changes, "note": "Old versions retained; changed order/total requires re-review and usually re-rendering of printed page numbers."}

def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("init")
    p.add_argument("--pack", type=Path, required=True)
    p.add_argument("--project", type=Path, required=True)
    p.add_argument("--style", type=Path, default=ROOT / "assets/white-inkgreen.json")
    for cmd in ("plan", "register", "review", "select", "audit", "assemble", "sync"):
        p = sub.add_parser(cmd)
        p.add_argument("--project", type=Path, required=True)
        if cmd in ("plan", "audit", "assemble"):
            p.add_argument("--range", dest="span")
        if cmd == "plan":
            p.add_argument("--draft", action="store_true")
        if cmd in ("register", "select"):
            p.add_argument("--slide-id", required=True)
            p.add_argument("--version", type=int, required=(cmd == "select"))
        if cmd == "register":
            p.add_argument("--image", type=Path, required=True)
            p.add_argument("--producer", required=True)
        if cmd == "review":
            p.add_argument("--file", type=Path, required=True)
        if cmd in ("select", "sync"):
            p.add_argument("--confirmation", required=True)
        if cmd == "audit":
            p.add_argument("--stage", choices=("structure", "render", "final"), default="final")
        if cmd == "assemble":
            p.add_argument("--output", type=Path, required=True)
        if cmd == "sync":
            p.add_argument("--pack", type=Path, required=True)
    a = parser.parse_args()
    try:
        if a.command == "init": result = init_project(a.pack, a.project, a.style)
        elif a.command == "plan": result = make_plan(a.project, a.span, a.draft)
        elif a.command == "register": result = register_image(a.project, a.slide_id, a.image, a.version, a.producer)
        elif a.command == "review": result = record_review(a.project, a.file)
        elif a.command == "select": result = select_image(a.project, a.slide_id, a.version, a.confirmation)
        elif a.command == "audit": result = audit(a.project, a.stage, a.span)
        elif a.command == "assemble": result = assemble(a.project, a.output, a.span)
        else: result = sync_pack(a.project, a.pack, a.confirmation)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 2 if result.get("errors") else 0
    except (PipelineError, OSError, KeyError, TypeError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

if __name__ == "__main__":
    raise SystemExit(main())
