#!/usr/bin/env python3
"""Read back a complete simulation case and its generated files independently.

This audit proves local consistency and readability. Customer acceptance and
production operation remain outside its verdict, even if simulation gates pass.
"""
from __future__ import annotations

import argparse
import hashlib
from html.parser import HTMLParser
import importlib.util
import json
from pathlib import Path
import shutil
import tempfile
from urllib.parse import unquote, urlsplit


SUITE = Path(__file__).resolve().parents[1]
RUNTIME = SUITE / 'skills/delivery-orchestrator/scripts/runtime.py'


def load_runtime():
    spec = importlib.util.spec_from_file_location('full_delivery_audit_runtime', RUNTIME)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def copied_source_valid(deliverables, entry):
    """Check delivered source bytes against their pinned source artifact."""
    root = deliverables.resolve()
    artifact = (root / entry['path']).resolve()
    source = (root / entry['source_file']).resolve()
    if not all(path.is_relative_to(root) and path.is_file() for path in (artifact, source)):
        return False
    data = json.loads(artifact.read_text())
    return sha(source) == data.get('data', {}).get('sha256')


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hrefs = []

    def handle_starttag(self, tag, attrs):
        if tag == 'a':
            self.hrefs.extend(value for key, value in attrs if key == 'href' and value)


def audit(case: Path, deliverables: Path, visual_report: Path | None = None):
    from docx import Document
    from openpyxl import load_workbook

    runtime = load_runtime()
    project = case / 'project'
    report = json.loads((case / 'report.json').read_text())
    snapshot = runtime.snapshot(project)
    checks = []
    files = []

    def check(name, passed, detail):
        checks.append({'id': name, 'passed': bool(passed), 'detail': detail})

    validation = runtime.validate(project)
    check('current-and-effective-versions', validation['ok'], validation)
    check('simulation-only', snapshot['purpose'] == 'simulation', snapshot['purpose'])
    expected = {p.parent.name for p in (SUITE / 'skills').glob('delivery-*/SKILL.md')}
    runs = report.get('skill_runs', [])
    names = {item.get('skill') for item in runs}
    check('all-25-skills-executed', names == expected and len(runs) == 25,
          {'expected': len(expected), 'recorded': len(runs), 'missing': sorted(expected - names)})
    for run in runs:
        name = run['skill']
        source = SUITE / 'skills' / name / 'SKILL.md'
        refs = [ref for key in ('input_refs', 'output_refs', 'evidence_refs') for ref in run.get(key, [])]
        errors = []
        for ref in refs:
            try:
                runtime.load_artifact(project, ref)
            except (ValueError, OSError, KeyError) as exc:
                errors.append(str(exc))
        check(f'skill-evidence:{name}', bool(run.get('output_refs')) and bool(refs)
              and source.is_file() and run.get('skill_sha256') == sha(source) and not errors, errors)

    objects = snapshot['artifacts']
    by_kind = {}
    for item in objects:
        by_kind.setdefault(item['kind'], []).append(item)
    required_kinds = {'evidence', 'contract', 'kickoff', 'discovery', 'minutes', 'requirements',
                      'blueprint', 'plan', 'timesheets', 'costs', 'risks', 'change', 'weekly',
                      'milestone', 'confirmation', 'review', 'solution', 'ui_design',
                      'integration', 'work_package', 'testing', 'release', 'handover', 'closure'}
    check('complete-delivery-stage-records', required_kinds <= by_kind.keys(),
          {'missing': sorted(required_kinds - by_kind.keys())})
    req = {row['id'] for obj in by_kind.get('requirements', []) for row in obj['data']['items']}
    for kind, key in [('solution', 'requirement_ids'), ('ui_design', 'requirement_ids'),
                      ('integration', 'requirement_ids'), ('work_package', 'requirement_ids'),
                      ('testing', 'requirement_ids')]:
        covered = {rid for obj in by_kind.get(kind, []) for rid in obj['data'][key]}
        check(f'requirement-coverage:{kind}', bool(req) and req <= covered,
              {'requirements': sorted(req), 'covered': sorted(covered)})
    for obj in by_kind.get('testing', []):
        data = obj['data']
        passed = {rid for row in data['cases'] if row['state'] == 'pass' and row['evidence_refs']
                  for rid in row['requirement_ids']}
        check(f'passing-test-evidence:{obj["artifact_id"]}', req <= passed, sorted(passed))
        fixture = SUITE / 'tests/fixtures/work_order_service.py'
        check(f'actual-build-hash:{obj["artifact_id"]}', data['build_ref']['sha256'] == sha(fixture), sha(fixture))
    for obj in by_kind.get('closure', []):
        gate = snapshot['gates'][obj['artifact_id']]
        check('continuing-warranty-not-prematurely-closed', gate.get('administrative_closed')
              and gate.get('warranty_transferred') and not gate.get('warranty_ended')
              and not gate.get('fully_discharged'), gate)
    check('cost-calculation-complete', snapshot['cost_summary'].get('complete'), snapshot['cost_summary'])

    # A deliberate corruption is tested on a disposable copy, never on the case.
    with tempfile.TemporaryDirectory(prefix='delivery-audit-') as temp:
        clone = Path(temp) / 'project'
        shutil.copytree(project, clone)
        target = next(obj for obj in objects if obj['kind'] == 'requirements')
        path = clone / target['path']
        path.write_bytes(path.read_bytes() + b'\n ')
        rejected = runtime.validate(clone)
        check('tampered-registered-file-rejected', not rejected['ok'], rejected['errors'])

    docx_paths = sorted(deliverables.rglob('*.docx'))
    xlsx_paths = sorted(deliverables.rglob('*.xlsx'))
    manifest = json.loads((deliverables / 'manifest.json').read_text())
    check('manifest-project-boundary', manifest.get('project_id') == snapshot['project_id']
          and manifest.get('purpose') == 'simulation' and manifest.get('audience') == 'internal',
          {key: manifest.get(key) for key in ('project_id', 'purpose', 'audience')})
    entries = manifest.get('documents', []) + [manifest.get('workbook', {})] + manifest.get('artifacts', [])
    for entry in entries:
        path = (deliverables / entry.get('path', '')).resolve()
        safe = path.is_relative_to(deliverables.resolve()) and path.is_file()
        check(f'manifest-hash:{entry.get("path")}', safe and sha(path) == entry.get('sha256'),
              entry.get('sha256'))
        if entry.get('source_file') and safe:
            artifact = json.loads(path.read_text())
            check(f'copied-source-hash:{entry["source_file"]}',
                  copied_source_valid(deliverables, entry),
                  artifact.get('data', {}).get('sha256'))
    check('all-25-filled-task-cards', {row['skill'] for row in manifest.get('task_cards', [])} == expected
          and all(row.get('artifact_refs') and row.get('input_refs')
                  for row in manifest.get('task_cards', [])), len(manifest.get('task_cards', [])))
    check('filled-office-files-exist', bool(docx_paths) and bool(xlsx_paths),
          {'docx': len(docx_paths), 'xlsx': len(xlsx_paths)})
    for path in docx_paths:
        document = Document(path)
        text = '\n'.join([p.text for p in document.paragraphs] +
                         [cell.text for table in document.tables for row in table.rows for cell in row.cells])
        check(f'docx-readback:{path.name}', len(text) > 100 and ('模拟' in text or 'simulation' in text.lower()),
              {'characters': len(text), 'tables': len(document.tables)})
    for path in xlsx_paths:
        workbook = load_workbook(path, data_only=False)
        sheets = {sheet.title: {'rows': sheet.max_row, 'columns': sheet.max_column,
                               'nonempty': sum(v is not None for row in sheet.values for v in row)}
                  for sheet in workbook.worksheets}
        check(f'xlsx-readback:{path.name}', bool(sheets) and all(v['nonempty'] > 1 for v in sheets.values()), sheets)
        workbook.close()

    html_path = deliverables / 'index.html'
    parser = Links()
    parser.feed(html_path.read_text())
    broken = []
    for href in parser.hrefs:
        parsed = urlsplit(href)
        if parsed.scheme or parsed.netloc:
            broken.append(href)
        elif parsed.path:
            target = (html_path.parent / unquote(parsed.path)).resolve()
            if not target.is_relative_to(case.resolve()) or not target.is_file():
                broken.append(href)
    check('review-download-links-confined-and-readable', not broken, broken)
    for path in sorted(deliverables.rglob('*')):
        if path.is_file():
            files.append({'path': str(path.relative_to(case)), 'bytes': path.stat().st_size, 'sha256': sha(path)})
    if visual_report:
        visual = json.loads(visual_report.read_text())
        check('actual-visual-and-browser-review', visual.get('passed') is True, visual)
    else:
        checks.append({'id': 'actual-visual-and-browser-review', 'passed': None,
                       'detail': 'Requires an actual separately recorded render and browser review'})
    failures = [row['id'] for row in checks if row['passed'] is False]
    pending = [row['id'] for row in checks if row['passed'] is None]
    return {'project_id': snapshot['project_id'], 'purpose': snapshot['purpose'],
            'verdict': 'fail' if failures else 'conditional' if pending else 'pass',
            'checks': checks, 'failures': failures, 'pending': pending, 'files': files,
            'boundary': {'production_executed': False, 'real_customer_accepted': False,
                         'real_erp_verified': False, 'warranty_ended': False}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--case', type=Path, required=True)
    parser.add_argument('--deliverables', type=Path, required=True)
    parser.add_argument('--visual-report', type=Path)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    result = audit(args.case.resolve(), args.deliverables.resolve(), args.visual_report)
    args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({key: result[key] for key in ('verdict', 'failures', 'pending')}, ensure_ascii=False))
    return 1 if result['failures'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
