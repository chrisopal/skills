#!/usr/bin/env python3
"""Create a filled, internal simulation review package from exact runtime records.

Run with bundled workspace Python. Authoring-operation marking and visual Office
QA belong to the caller; this renderer never starts LibreOffice or a desktop app.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import importlib.util
import json
import os
import shutil
import subprocess
import tempfile
from decimal import Decimal, InvalidOperation
import re
from pathlib import Path
from urllib.parse import quote

SUITE = Path(__file__).resolve().parents[1]
ORCH = SUITE / 'skills/delivery-orchestrator'
ASSET = Path(__file__).resolve().parent / 'assets/full-delivery-review.html'


def module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    obj = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(obj)
    return obj


runtime = module('full_delivery_runtime', ORCH / 'scripts/runtime.py')
review = module('full_delivery_labels', ORCH / 'scripts/render_review.py')
exporter = module('full_delivery_exporter', ORCH / 'scripts/export_templates.py')
LABELS = dict(review.LABELS, capacity='可用容量', allocations='资源分配', travel='差旅',
              trips='出差记录', risks='风险与问题', progress='进度控制', purpose='用途',
              goals='项目目标', control='统筹控制', goal='本次目标', input_refs='输入版本',
              capacities='人员容量', available_hours='可用工时', planned_hours='计划工时', remaining_work='剩余工作', remaining_hours='剩余工时',
              activity='工作内容', outsourced='外包', incurred='已发生', paid='已支付', outsourcing='外包', purchase='采购', external_ref='外包凭据', procurement_ref='采购凭据', travel_or_expense_ref='差旅凭据',
              output_refs='输出版本', evidence_refs='执行证据', current='当前', stale='需刷新', gaps='待补事项', journal='任务记录')
# Each grouped document has a distinct business responsibility.
GROUPS = [
    ('01-contract', '合同交接与项目启动', ['contract', 'kickoff'], '合同义务、范围和角色是后续交付的执行依据。'),
    ('02-discovery', '客户调研计划与会议纪要', ['discovery', 'minutes'], '按调研安排核对业务资料，会议决定与行动项进入需求处理。'),
    ('03-requirements', '需求基线与业务蓝图', ['requirements', 'blueprint'], '需求保留来源与验收标准，业务流程说明各需求如何落地。'),
    ('04-design', '详细技术设计与页面字段权限', ['solution', 'ui_design'], '技术组件、设计决策、字段约束和状态权限共同约束工程实现。'),
    ('05-integration', 'ERP接口与数据迁移设计', ['integration'], '按接口约定、字段映射与数量及业务键核对执行迁移和恢复。'),
    ('06-plan', 'WBS总计划与里程碑安排', ['plan'], '以基线日期、预测日期、负责人和依赖组织工作。'),
    ('07-work-packages', '实施工作包与构建证据', ['work_package'], '工作包绑定有效设计、构建版本和完成依据。'),
    ('08-economics', '资源工时成本与差旅台账', ['timesheets', 'costs', 'profile'], '按已登记工时和费用口径区分发生、剩余承诺和完工预测。'),
    ('09-control', '风险变更进度与项目周报', ['risks', 'change', 'weekly'], '按截止日期查看已发生事项、变更决定和后续行动。'),
    ('10-quality', '测试缺陷与UAT记录', ['testing'], '逐项核对用例、实际结果、缺陷复测和模拟确认范围。'),
    ('11-release', '上线回滚与培训运维移交', ['release', 'handover'], '执行结果与准备状态分开记录，移交以接收证据和资产版本为依据。'),
    ('12-acceptance', '里程碑验收与项目结项质保', ['milestone', 'confirmation', 'closure'], '分别核对里程碑范围、行政收尾、质保移交和剩余义务。'),
    ('13-audit', '交付成果审核与归档', ['review', 'template_mapping', 'exchange'], '审核结论只覆盖对应版本，未核验范围保留为缺口。'),
]
NAV = [
    ('overview', '总览', []), ('planning', '计划与里程碑', ['plan']),
    ('design', '调研与设计', ['contract', 'kickoff', 'discovery', 'minutes', 'requirements', 'blueprint', 'solution', 'ui_design', 'integration']),
    ('quality', '实施与质量', ['work_package', 'testing', 'release', 'handover']),
    ('control', '成本与风险', ['timesheets', 'costs', 'profile', 'risks', 'change', 'weekly']),
    ('archive', '验收与归档', ['milestone', 'confirmation', 'closure']),
    ('outputs', '成果与审核', ['review', 'template_mapping', 'exchange']),
]
SKILL_KINDS = {
    'orchestrator': [], 'source-intake': ['source'], 'contract-handover': ['contract'],
    'initiation-kickoff': ['kickoff'], 'discovery-planning': ['discovery'], 'meeting-records': ['minutes'],
    'requirements-baseline': ['requirements'], 'business-blueprint': ['blueprint'],
    'solution-design': ['solution'], 'ui-field-design': ['ui_design'], 'integration-data': ['integration'],
    'master-planning': ['plan'], 'work-package': ['work_package'], 'resource-timesheets': ['timesheets'],
    'cost-expenses': ['costs'], 'risk-issue': ['risks'], 'change-control': ['change'],
    'progress-control': ['plan', 'weekly'], 'reporting': ['weekly'], 'quality-testing': ['testing'],
    'go-live-transition': ['release'], 'training-handover': ['handover'],
    'milestone-acceptance': ['milestone', 'confirmation'], 'closure-warranty': ['closure'],
    'artifact-review': ['review'],
}
REF_KEYS = {'artifact_id', 'revision', 'sha256'}
DETAIL_KEYS = {'evidence_refs', 'source_refs', 'decision_refs', 'receipt_refs', 'backup_refs', 'rehearsal_refs', 'reconciliation_refs', 'verification_refs', 'restore_refs', 'runbook_refs', 'transfer_refs', 'terms_refs', 'budget_evidence_refs', 'remaining_labor_evidence_refs', 'other_remaining_evidence_refs'}


def esc(value):
    return html.escape(str(value), quote=True)


def label(value):
    return LABELS.get(str(value), str(value).replace('_', ' '))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def ref_of(artifact):
    return {key: artifact[key] for key in ('artifact_id', 'revision', 'sha256')}


def text(value, context=None):
    if isinstance(value, dict):
        if REF_KEYS <= value.keys():
            return f"{value['artifact_id']} r{value['revision']}"
        return '；'.join(f'{label(k)}：{text(v, k)}' for k, v in value.items())
    if isinstance(value, list):
        return '、'.join(text(x) for x in value) if value else '无记录'
    if value is True:
        return '是'
    if value is False:
        return '否'
    if value is None or value == '':
        return '未填写'
    if context in {'billing_mode', 'mode'}:
        return {'fixed': '固定价', 'time': '按工时', 'mixed': '混合'}.get(str(value), str(value))
    return str(LABELS.get(str(value), value))


def curated(data):
    """All business fields remain readable; long reference arrays live in raw data."""
    return {key: val for key, val in data.items() if key not in DETAIL_KEYS}


def blocks(data):
    """Yield section headings and readable rows, avoiding serialized JSON bodies."""
    for key, value in curated(data).items():
        if isinstance(value, list) and value and all(isinstance(x, dict) for x in value):
            rows = []
            for index, row in enumerate(value, 1):
                identity = text(row.get('id', index))
                body = '；'.join(f'{label(k)}：{text(v, k)}' for k, v in curated(row).items() if k != 'id')
                rows.append((identity, body))
            yield label(key), rows
        else:
            yield label(key), [('', text(value, key))]


def artifact_markdown(artifact):
    lines = [f"## {label(artifact['kind'])} {artifact['artifact_id']} r{artifact['revision']}", '',
             f"状态：{label(artifact['status'])}；版本健康：{label(artifact.get('health', 'unknown'))}。", '']
    for heading, rows in blocks(artifact['data']):
        lines.extend([f'### {heading}', ''])
        lines.extend(f'- {identity + "：" if identity else ""}{body}' for identity, body in rows)
        lines.append('')
    lines.extend([f"完整依据：[登记数据](../raw/{artifact['file']})", ''])
    return '\n'.join(lines)


def load_case(case):
    project = case / 'project' if (case / 'project/00-control').is_dir() else case
    snapshot = runtime.snapshot(project)
    if snapshot.get('purpose') != 'simulation':
        raise ValueError('This internal demonstration renderer requires purpose=simulation')
    selected = runtime.selection_for(project, runtime.read_registry(project))
    if set(selected) != {record['artifact_id'] for record in snapshot['artifacts']}:
        raise ValueError('Runtime snapshot omitted missing or corrupt registered artifacts')
    records = []
    for record in snapshot['artifacts']:
        ref = ref_of(record)
        obj = runtime.load_artifact(project, ref)
        aid = obj['artifact_id']
        if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]*', aid):
            raise ValueError('Unsafe artifact identifier')
        records.append(dict(obj, sha256=ref['sha256'], health=record['health'],
                            is_effective=record.get('is_effective', False), file=f"{aid}-r{obj['revision']}.json"))
    report_path = case / 'report.json'
    report = json.loads(report_path.read_text(encoding='utf-8')) if report_path.is_file() else {}
    return project, snapshot, records, report


def trace_rows(records):
    rows = []
    requirements = [item for a in records if a['kind'] == 'requirements' for item in a['data'].get('items', [])]
    for req in requirements:
        rid = req['id']
        design, cases = [], []
        for artifact in records:
            data = artifact['data']
            if artifact['kind'] in {'solution', 'ui_design', 'integration', 'work_package'} and rid in data.get('requirement_ids', []):
                design.append(f"{artifact['artifact_id']} r{artifact['revision']}")
            if artifact['kind'] == 'testing':
                cases.extend(f"{case['id']} {label(case.get('state', 'unknown'))}" for case in data.get('cases', []) if rid in case.get('requirement_ids', []))
        rows.append({'需求': rid, '需求与验收': text(req.get('statement')) + ' / ' + text(req.get('acceptance')), '设计与工作包': '、'.join(design) or '缺少关联', '测试结果': '、'.join(cases) or '未关联用例'})
    return rows


def table(rows, columns=None):
    if not rows:
        return '<p class="empty-static">暂无登记记录</p>'
    columns = columns or list(rows[0])
    head = ''.join(f'<th scope="col">{esc(col)}</th>' for col in columns)
    body = ''.join('<tr data-search-row>' + ''.join(f'<td>{row.get(col, "")}</td>' for col in columns) + '</tr>' for row in rows)
    return f'<div class="table-scroll" tabindex="0" aria-label="可横向滚动的数据表"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>'


def artifact_html(artifact):
    heading = f"{label(artifact['kind'])} · {artifact['artifact_id']} r{artifact['revision']}"
    parts = [f'<article class="record" data-search-row><h3>{esc(heading)} <span class="tag">{esc(label(artifact["status"]))}</span></h3>']
    for title, rows in blocks(artifact['data']):
        parts.append(f'<h4>{esc(title)}</h4>')
        if len(rows) == 1 and not rows[0][0]:
            parts.append(f'<p>{esc(rows[0][1])}</p>')
        else:
            parts.append(table([{'编号': esc(identity), '业务内容': esc(body)} for identity, body in rows]))
    raw = esc(json.dumps(artifact['data'], ensure_ascii=False, indent=2))
    parts.append(f'<details><summary>版本与完整登记数据</summary><p>SHA-256 <code>{artifact["sha256"]}</code></p><a class="button secondary" download href="raw/{quote(artifact["file"])}">下载原始 JSON</a><pre>{raw}</pre></details></article>')
    return ''.join(parts)


def write_docx(path, title, introduction, records, project_id):
    from docx import Document
    from docx.shared import Cm
    doc = Document()
    exporter._style_document(doc)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(2)
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(f'{project_id} 设备维保工单定制平台与 ERP 集成')
    doc.add_paragraph('模拟演练 内部审核')
    doc.add_paragraph(introduction)
    for artifact in records:
        doc.add_heading(label(artifact['kind']) + ' ' + artifact['artifact_id'] + ' r' + str(artifact['revision']), 1)
        doc.add_paragraph('状态 ' + label(artifact['status']) + '  版本健康 ' + label(artifact['health'])
                          + '\n完整登记数据 raw/' + artifact['file'])
        for heading, rows in blocks(artifact['data']):
            doc.add_heading(heading, 2)
            if len(rows) == 1 and not rows[0][0]:
                doc.add_paragraph(rows[0][1])
            else:
                # Two columns remain legible in portrait; keep long nested data out of narrow cells.
                for identity, body in rows:
                    doc.add_paragraph(identity + '  ' + body)
    if not records:
        doc.add_paragraph('当前项目尚未登记本组记录。')
    doc.save(path)


def workbook_spec(records, traces, snapshot):
    """Return typed business ledgers; document narrative stays in the DOCX groups."""
    summary = snapshot.get('cost_summary', {})
    overview = [
        ['项目编号', snapshot['project_id'], '', '运行时项目登记'],
        ['用途', '模拟演练 内部审核', '', '当前有效版本快照'],
        ['真实客户验收', '未执行', '', '模拟确认不能替代真实客户签署'],
        ['生产上线', '未执行', '', '仅本地 HTTP 与 SQLite fixture'],
        ['成本预测范围', '本次实施范围', '', '持续质保成本未估算，未纳入完工预测'],
        ['登记产物', len(records), '项', '版本清单'],
        ['需求数量', len(traces), '项', '需求设计测试追溯'],
    ]
    financial = ('budget', 'internal_labor', 'provisional_labor', 'other_incurred',
                 'committed_remaining', 'estimated_remaining', 'forecast', 'variance')
    for key in financial:
        value = summary.get(key)
        if value is not None:
            try:
                value = float(Decimal(str(value)))
            except InvalidOperation:
                value = text(value)
            overview.append([label(key), value, summary.get('currency', 'CNY'), '成本主账当前汇总'])
    overview.append(['计算与依据完整', text(summary.get('complete')), '', text(summary.get('gaps', []))])
    datasets = [{'name': '项目概览', 'headers': ['事项', '当前值', '单位', '范围与依据'], 'rows': overview, 'summary': True}]
    datasets.append({'name': '需求设计测试追溯', 'headers': list(traces[0]) if traces else ['需求', '需求与验收', '设计与工作包', '测试结果'], 'rows': [list(r.values()) for r in traces]})
    datasets.append({'name': '版本清单', 'headers': ['产物', '类型', '版本', '状态', '健康', '有效版本', 'SHA256'], 'rows': [[a['artifact_id'], label(a['kind']), a['revision'], label(a['status']), label(a['health']), a['is_effective'], a['sha256']] for a in records]})
    extracts = [
        ('需求清单', 'requirements', 'items'), ('WBS明细', 'plan', 'work_packages'),
        ('里程碑明细', 'plan', 'milestones'), ('工时明细', 'timesheets', 'items'),
        ('人员容量', 'timesheets', 'capacities'), ('人员分配', 'timesheets', 'allocations'),
        ('剩余工作', 'timesheets', 'remaining_work'), ('费用明细', 'costs', 'items'),
        ('风险问题台账', 'risks', 'items'), ('测试用例', 'testing', 'cases'),
        ('缺陷复测', 'testing', 'defects'), ('页面字段', 'ui_design', 'fields'),
        ('权限规则', 'ui_design', 'permissions'), ('接口字段映射', 'integration', 'mappings'),
    ]
    for name, kind, key in extracts:
        entries = [dict(artifact_id=a['artifact_id'], revision=a['revision'], **row)
                   for a in records if a['kind'] == kind for row in a['data'].get(key, [])]
        keys = list(dict.fromkeys(k for row in entries for k in curated(row)))
        if keys:
            values = [[text(row.get(k)) if isinstance(row.get(k), (dict, list)) else row.get(k)
                       for k in keys] for row in entries]
            datasets.append({'name': name, 'headers': [label(k) for k in keys], 'rows': values, 'keys': keys})
    return {'project_id': snapshot['project_id'], 'sheets': datasets}


def write_xlsx(path, records, traces, snapshot):
    dependencies = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies'
    node = Path(os.environ.get('DELIVERY_NODE', dependencies / 'node/bin/node'))
    modules = Path(os.environ.get('DELIVERY_NODE_MODULES', dependencies / 'node/node_modules'))
    if not node.is_file() or not (modules / '@oai/artifact-tool').is_dir():
        raise RuntimeError('Bundled Node and @oai/artifact-tool are required; no system fallback')
    builder = Path(__file__).with_name('build_full_delivery_workbook.mjs')
    with tempfile.TemporaryDirectory(prefix='delivery-workbook-') as directory:
        working = Path(directory)
        (working / 'node_modules').symlink_to(modules, target_is_directory=True)
        shutil.copyfile(builder, working / builder.name)
        spec = working / 'workbook-spec.json'
        spec.write_text(json.dumps(workbook_spec(records, traces, snapshot), ensure_ascii=False), encoding='utf-8')
        subprocess.run([str(node), str(working / builder.name), '--spec', str(spec), '--out', str(path),
                        '--preview-dir', str(path.parent / 'workbook-previews')],
                       cwd=working, check=True, capture_output=True, text=True)


def task_cards(out, records, report, templates):
    cards = []
    journal = report.get('skill_runs', report.get('skill_journal', report.get('task_journal', [])))
    if isinstance(journal, dict):
        journal = [dict(value, skill=key) for key, value in journal.items()]
    for suffix, kinds in SKILL_KINDS.items():
        skill = 'delivery-' + suffix
        selected = records if not kinds else [a for a in records if a['kind'] in kinds]
        source = SUITE / 'skills' / skill / 'SKILL.md'
        entries = [t for t in templates if t.get('skill') == skill]
        event = next((x for x in journal if x.get('skill', x.get('skill_id')) == skill), {})
        title = next((t['title'] for t in entries), label(kinds[0]) if kinds else '项目统筹')
        gaps = event.get('unresolved', event.get('gaps', []))
        if not selected:
            gaps = [*gaps, '当前未登记该技能对应产物']
        body = [f'# {title} 任务记录', '', '模拟演练 内部审核', '', '## 目标', '',
                text(event.get('goal', f'核对{title}对应产物、版本依据和未完成事项。')), '',
                '## 技能来源', '', f'- 技能：{skill}', f'- 来源 SHA-256：{sha(source)}', '', '## 输入依据', '']
        input_refs = event.get('input_refs', [r for a in selected for r in a.get('inputs', [])])
        inputs = {f"{r['artifact_id']} r{r['revision']}": r for r in input_refs if REF_KEYS <= r.keys()}
        body.extend('- ' + name + ' SHA-256 ' + value['sha256'] for name, value in inputs.items())
        if not inputs:
            body.append('- 当前登记产物没有独立输入引用。')
        body.extend(['', '## 输出与执行证据', ''])
        for category in ('output_refs', 'evidence_refs'):
            for ref in event.get(category, []):
                body.append(f"- {label(category)} [{text(ref)}](../raw/{ref['artifact_id']}-r{ref['revision']}.json) SHA-256 {ref['sha256']}")
        for artifact in selected:
            body.append(f"- [{artifact['artifact_id']} r{artifact['revision']}](../raw/{artifact['file']})：{label(artifact['status'])}；{label(artifact['health'])}；SHA-256 {artifact['sha256']}")
            body.append(artifact_markdown(artifact))
        body.extend(['', '## 模板复用', ''])
        body.extend('- ' + t['id'] + ' ' + t['path'] for t in entries)
        if not entries:
            body.append('- 本技能使用运行时主账，无单独文档模板。')
        body.extend(['', '## 待补事项', '', text(gaps) if gaps else '当前登记记录无额外待补事项；真实客户确认与生产执行不在本次模拟范围。'])
        path = out / 'tasks' / (skill + '.md')
        path.write_text('\n'.join(body) + '\n', encoding='utf-8')
        cards.append({'skill': skill, 'title': title, 'path': 'tasks/' + path.name, 'source_sha256': sha(source), 'artifact_refs': event.get('output_refs', [ref_of(a) for a in selected]), 'input_refs': list(inputs.values()), 'evidence_refs': event.get('evidence_refs', []), 'templates': [t['id'] for t in entries]})
    return cards


def copy_evidence(project, out, records, report):
    """Copy pinned historical inputs as well as selected versions without rebasing refs."""
    pending = [ref_of(a) for a in records]
    pending.extend(runtime.walk_refs(report))
    seen, result = set(), []
    while pending:
        ref = pending.pop()
        key = (ref['artifact_id'], ref['revision'], ref['sha256'])
        if key in seen:
            continue
        seen.add(key)
        obj = runtime.load_artifact(project, ref)
        filename = f"{obj['artifact_id']}-r{obj['revision']}.json"
        item = dict(obj, sha256=ref['sha256'], file=filename)
        (out / 'raw' / filename).write_bytes(runtime.artifact_path(project, ref).read_bytes())
        if obj['kind'] == 'source':
            source = runtime.confined(project, obj['data']['path'])
            source_name = f"source-{obj['artifact_id']}-r{obj['revision']}{source.suffix}"
            (out / 'raw' / source_name).write_bytes(source.read_bytes())
            item['source_file'] = 'raw/' + source_name
        result.append(item)
        pending.extend(runtime.walk_refs(obj))
    return sorted(result, key=lambda a: (a['artifact_id'], a['revision']))


def render(case, out):
    case, out = Path(case).expanduser().resolve(), Path(out).expanduser().resolve()
    if out.exists():
        raise ValueError('Output directory already exists; choose a fresh directory')
    project, snapshot, records, report = load_case(case)
    if out == project or out.is_relative_to(project / '00-control'):
        raise ValueError('Output cannot replace runtime control records')
    out.mkdir(parents=True)
    for folder in ('raw', 'documents', 'tasks'):
        (out / folder).mkdir()
    raw_records = copy_evidence(project, out, records, report)
    for artifact in records:
        raw_info = next(item for item in raw_records if ref_of(item) == ref_of(artifact))
        if raw_info.get('source_file'):
            artifact['source_file'] = raw_info['source_file']
    (out / 'snapshot.json').write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (out / 'case-report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    templates = json.loads((ORCH / 'assets/templates/index.json').read_text(encoding='utf-8'))['templates']
    docs = []
    for slug, title, kinds, intro in GROUPS:
        selected = [a for a in records if a['kind'] in kinds]
        path = out / 'documents' / (slug + '.docx')
        write_docx(path, title, intro, selected, snapshot['project_id'])
        mdpath = path.with_suffix('.md')
        mdpath.write_text('# ' + title + '\n\n模拟演练 内部审核\n\n' + intro + '\n\n' + '\n'.join(artifact_markdown(a) for a in selected), encoding='utf-8')
        docs.append({'title': title, 'path': 'documents/' + path.name, 'markdown': 'documents/' + mdpath.name, 'sha256': sha(path), 'artifact_refs': [ref_of(a) for a in selected]})
    traces = trace_rows(records)
    write_xlsx(out / 'delivery-tracking.xlsx', records, traces, snapshot)
    cards = task_cards(out, records, report, templates)
    manifest = {'project_id': snapshot['project_id'], 'purpose': 'simulation', 'audience': 'internal',
                'documents': docs, 'workbook': {'path': 'delivery-tracking.xlsx', 'sha256': sha(out / 'delivery-tracking.xlsx')},
                'task_cards': cards, 'artifacts': [{**ref_of(a), 'kind': a['kind'], 'path': 'raw/' + a['file'], 'source_file': a.get('source_file')} for a in raw_records],
                'template_index_count': len(templates), 'boundary': {'production_executed': False, 'real_customer_acceptance': False}}
    (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    nav = ''.join(f'<a href="#{anchor}">{title}</a>' for anchor, title, _ in NAV)
    section_html = []
    for anchor, title, kinds in NAV:
        content = ''
        if anchor == 'overview':
            metrics = [('登记产物', len(records)), ('当前有效版本', sum(a['is_effective'] for a in records)), ('需求', len(traces)), ('技能任务记录', len(cards))]
            content += '<div class="metrics">' + ''.join(f'<div><span>{name}</span><strong>{value}</strong></div>' for name, value in metrics) + '</div>'
            content += '<h3>当前需关注事项</h3>' + ('<ul>' + ''.join(f'<li>{esc(x)}</li>' for x in snapshot['alerts']) + '</ul>' if snapshot['alerts'] else '<p>运行时未报告版本告警。</p>')
            content += '<h3>当前阶段核对</h3>' + table([{'对象': esc(label(key)), '当前结论': esc(text(value))} for key, value in snapshot.get('gates', {}).items()])
            content += '<h3>需求到设计与测试</h3>' + table([{k: esc(v) for k, v in row.items()} for row in traces])
        content += ''.join(artifact_html(a) for a in records if a['kind'] in kinds)
        if anchor == 'outputs':
            downloads = [{'成果': esc(d['title']), '文件': f'<a class="button secondary" href="{d["path"]}" download>DOCX</a> <a class="button secondary" href="{d["markdown"]}" download>正文 MD</a>', '版本范围': esc('、'.join(text(r) for r in d['artifact_refs']))} for d in docs]
            content += '<h3>分组成果</h3>' + table(downloads)
            sources = [a for a in raw_records if a.get('source_file')]
            content += '<h3>来源与工程证据</h3>' + table([{'来源': esc(a['artifact_id']), '版本': a['revision'], '原件': f'<a class="button secondary" href="{a["source_file"]}" download>下载来源原件</a>', '文件指纹': esc(a['data']['sha256'])} for a in sources])
            content += '<h3>25项技能任务记录</h3>' + table([{'技能': esc(card['skill']), '任务': esc(card['title']), '输出': f'<a class="button secondary" href="{card["path"]}" download>下载任务记录</a>', '产物版本数': len(card['artifact_refs'])} for card in cards])
        section_html.append(f'<section id="{anchor}" class="page-section"><h2>{title}</h2>{content}<p class="filter-empty" hidden>没有匹配记录，请修改筛选词。</p></section>')
    document = ASSET.read_text(encoding='utf-8')
    audit = case / 'final-review.md'
    audit_link = ('<a class="button secondary" href="' + esc(quote(os.path.relpath(audit, out)))
                  + '">最终独立审核</a>') if audit.is_file() else ''
    replacements = {'@@PROJECT@@': esc(snapshot['project_id']), '@@NAV@@': nav,
                    '@@FINAL_AUDIT@@': audit_link, '@@SECTIONS@@': ''.join(section_html)}
    for old, new in replacements.items():
        document = document.replace(old, new)
    (out / 'index.html').write_text(document, encoding='utf-8')
    return manifest


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--case', required=True, type=Path)
    parser.add_argument('--out', required=True, type=Path)
    args = parser.parse_args(argv)
    result = render(args.case, args.out)
    print(json.dumps({'project_id': result['project_id'], 'documents': len(result['documents']), 'tasks': len(result['task_cards']), 'out': str(args.out.resolve())}, ensure_ascii=False))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
