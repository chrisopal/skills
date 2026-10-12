#!/usr/bin/env python3
"""Export six reviewable delivery documents from a runtime project snapshot.

The exporter reads the project's registered, latest runtime records.  It never
creates facts to fill a template: missing records and fields are rendered as
``待确认``.  Each document is written as both Markdown and native DOCX, with
an immutable source manifest and a discussion-draft label.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import sys
import tempfile
from pathlib import Path
from typing import Any, Iterable

LABELS = {"id":"编号", "title":"名称", "description":"说明", "steps":"步骤", "state":"状态", "status":"状态", "source":"来源", "target":"目标", "required":"必填", "length":"长度", "precision":"精度", "role":"角色", "action":"动作", "condition":"条件", "effect":"效果", "error":"错误", "name":"名称", "rule":"规则", "rules":"业务规则", "requirement_ids":"需求编号", "owner":"责任人", "date":"日期", "reason":"原因", "statement":"要求", "acceptance":"验收", "channel":"渠道", "origin":"来源性质", "read_status":"读取状态", "timeout":"超时", "retry":"重试", "compensation":"补偿", "locator":"定位", "business_key":"业务键", "unit":"单位", "default":"默认值", "enum":"可选值", "validation":"校验", "sensitivity":"敏感级别", "precondition":"前提", "success":"成功结果", "failure_recovery":"失败恢复", "audit_event":"审计事件", "implementation":"实现状态", "button":"按钮", "flow_ids":"流程编号", "page_id":"页面编号", "field_ids":"字段编号", "applicable":"适用", "rollback_plan":"回滚方案", "freeze_window":"冻结窗口", "backup_refs":"备份依据", "rehearsal_refs":"演练依据", "reconciliation_refs":"核对依据", "draft":"草稿", "ready":"作者就绪", "blocked":"阻塞", "current":"引用有效", "stale":"需复核", "simulation":"模拟", "actual":"实际", "accepted":"已记录接受", "conditional":"有条件", "rejected":"已记录拒绝"}


SCRIPT_DIR = Path(__file__).resolve().parent


def _load_runtime() -> Any:
    spec = importlib.util.spec_from_file_location("delivery_core_runtime", SCRIPT_DIR / "runtime.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("unable to load delivery runtime")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


runtime = _load_runtime()
export_templates = None


class CoreExportError(ValueError):
    """Raised when a project or output cannot be safely exported."""


DOCS = (
    ("01-initiation-kickoff", "项目启动会", "customer", "kickoff"),
    ("02-business-blueprint", "业务蓝图", "customer", "blueprint"),
    ("03-solution-design", "实施方案与设计决策", "internal", "solution"),
    ("04-master-plan", "项目总计划与 WBS", "internal", "plan"),
    ("05-integration-data", "接口数据迁移与联调", "internal", "integration"),
    ("06-milestone-acceptance", "里程碑验收包", "customer", "milestone"),
)

FORBIDDEN_MARKERS = ("password", "passwd", "secret", "privatekey", "apikey", "accesskey",
                     "credential", "bearer", "token", "internalrate", "salary", "margin")


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _ref(value: Any) -> dict[str, Any] | None:
    if not isinstance(value, dict):
        return None
    if not {"artifact_id", "revision", "sha256"} <= value.keys():
        return None
    if not isinstance(value["artifact_id"], str) or not isinstance(value["revision"], int):
        return None
    checksum = value["sha256"]
    if not isinstance(checksum, str) or not re.fullmatch(r"[0-9a-f]{64}", checksum):
        return None
    result = {"artifact_id": value["artifact_id"], "revision": value["revision"], "sha256": checksum}
    if value.get("locator") not in (None, ""):
        result["locator"] = str(value["locator"])
    return result


def _ref_key(value: dict[str, Any]) -> tuple[Any, ...]:
    return (value["artifact_id"], value["revision"], value["sha256"], value.get("locator", ""))


def _collect_refs(value: Any, result: set[tuple[Any, ...]] | None = None) -> list[dict[str, Any]]:
    result = result if result is not None else set()
    found = _ref(value)
    if found:
        result.add(_ref_key(found))
    elif isinstance(value, dict):
        for child in value.values():
            _collect_refs(child, result)
    elif isinstance(value, list):
        for child in value:
            _collect_refs(child, result)
    return [dict(artifact_id=k[0], revision=k[1], sha256=k[2], **({"locator": k[3]} if k[3] else {}))
            for k in sorted(result)]


def _sanitize(value: Any) -> Any:
    """Apply the runtime customer boundary and remove unsafe display keys."""
    if isinstance(value, dict):
        return {key: _sanitize(child) for key, child in value.items()
                if not runtime.forbidden_customer_key(key)
                and not any(marker in re.sub(r"[^a-z0-9]", "", str(key).lower()) for marker in FORBIDDEN_MARKERS)}
    if isinstance(value, list):
        return [_sanitize(child) for child in value]
    return value


def _label(key: Any) -> str:
    return LABELS.get(str(key), str(key).replace("_", " "))


def _display(value: Any, context: str | None = None) -> str:
    if value is None or value == "":
        return "待确认"
    if value is True:
        return "是"
    if value is False:
        return "否"
    found = _ref(value)
    if found:
        return _ref_text(found)
    if isinstance(value, dict):
        return "；".join(f"{_label(k)}：{_display(v, str(k))}" for k, v in value.items()) or "待确认"
    if isinstance(value, list):
        if not value:
            return "无记录"
        return "；".join(_display(child) for child in value)
    return LABELS.get(str(value), str(value))


def _ref_text(value: dict[str, Any]) -> str:
    suffix = f"，定位 {value['locator']}" if value.get("locator") else ""
    return f"{value['artifact_id']} r{value['revision']}（指纹:{value['sha256'][:12]}）{suffix}"


def _md_cell(value: Any) -> str:
    return _display(value).replace("|", "\\|").replace("\n", " ")


def _table(headers: list[str], rows: list[list[Any]]) -> list[str]:
    if any(len(row) != len(headers) for row in rows):
        raise CoreExportError('Table columns must match all business fields')
    if len(headers) > 5:
        result = []
        for start in range(1, len(headers), 4):
            positions = [0, *range(start, min(start + 4, len(headers)))]
            result.extend(_table([headers[i] for i in positions], [[row[i] for i in positions] for row in rows]))
            result.append('')
        return result
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(_md_cell(value) for value in row[:len(headers)]) + " |")
    if not rows:
        lines.append("| " + " | ".join("待确认" for _ in headers) + " |")
    return lines


def _section(lines: list[str], title: str, body: str | None = None) -> None:
    lines.extend([f"## {title}", ""])
    if body:
        lines.extend([body, ""])


def _record_ref(record: dict[str, Any]) -> dict[str, Any]:
    return {key: record[key] for key in ("artifact_id", "revision", "sha256")}


def _load_project(project: str | Path) -> tuple[Path, dict[str, Any], dict[str, dict[str, Any]]]:
    root = Path(project).expanduser().resolve()
    info = runtime.read_project(root)
    snapshot = runtime.snapshot(root)
    registry = runtime.read_registry(root)
    selection = runtime.selection_for(root, registry)
    if {x['artifact_id'] for x in snapshot['artifacts']} != set(selection):
        raise CoreExportError('Registered source is missing or corrupt; repair before exporting')
    records: dict[str, dict[str, Any]] = {}
    core_kinds = {'kickoff','contract','requirements','blueprint','plan','solution','ui_design','integration','milestone'}
    for item in snapshot["artifacts"]:
        ref = _record_ref(item)
        # Re-read the exact artifact through runtime so the export cannot use a
        # stale, hand-edited snapshot payload.
        obj = runtime.load_artifact(root, ref)
        if obj['kind'] in core_kinds:
            runtime.check_ref(root, ref, selection)
        records[item["artifact_id"]] = dict(obj, sha256=item["sha256"], health=item.get("health"),
                                             is_effective=item.get("is_effective", False))
    return root, snapshot, records


def _by_kind(records: dict[str, dict[str, Any]], kind: str) -> list[dict[str, Any]]:
    return sorted((record for record in records.values() if record.get("kind") == kind),
                  key=lambda record: (record["artifact_id"], record["revision"]))


def _first(records: dict[str, dict[str, Any]], kind: str) -> dict[str, Any] | None:
    values = _by_kind(records, kind)
    return values[-1] if values else None


def _confirmation_for(records: dict[str, dict[str, Any]], target: dict[str, Any] | None) -> dict[str, Any] | None:
    if target is None:
        return None
    target_ref = _record_ref(target)
    for candidate in _by_kind(records, "confirmation"):
        if _data(candidate).get("target") == target_ref:
            return candidate
    return None


def _data(record: dict[str, Any] | None, customer: bool = False) -> dict[str, Any]:
    if not record:
        return {}
    value = record.get("data", {})
    return _sanitize(value) if customer else value


def _source_lines(lines: list[str], records: Iterable[dict[str, Any]], customer: bool = False) -> list[dict[str, Any]]:
    refs: set[tuple[Any, ...]] = set()
    for record in records:
        for ref in _collect_refs(record.get("inputs", [])):
            refs.add(_ref_key(ref))
        for ref in _collect_refs(record.get("data", {})):
            refs.add(_ref_key(ref))
        refs.add(_ref_key(_record_ref(record)))
    values = [dict(artifact_id=k[0], revision=k[1], sha256=k[2], **({"locator": k[3]} if k[3] else {}))
              for k in sorted(refs)]
    _section(lines, "成果状态")
    lines.extend(_table(["成果", "修订", "编写状态", "引用状态", "版本角色"],
        [[r['artifact_id'], r['revision'], r.get('status'), r.get('health'), '有效基线' if r.get('is_effective') else '候选/记录'] for r in records]))
    lines.append('')
    _section(lines, "版本与来源", "下表指纹显示前12位；完整 SHA256 与精确输入见同目录 core-documents-manifest.json。未列入内容不代表已确认。")
    lines.extend(_table(["成果编号", "修订", "指纹前12位", "定位"],
                        [[x["artifact_id"], x["revision"], x["sha256"][:12], x.get("locator", "")] for x in values]))
    lines.append("")
    return values


def _software_builds(value):
    if isinstance(value, dict):
        if {"repository", "revision", "sha256"} <= value.keys() and isinstance(value["revision"], str):
            yield {key: value[key] for key in ("repository", "revision", "sha256")}
        for child in value.values():
            yield from _software_builds(child)
    elif isinstance(value, list):
        for child in value:
            yield from _software_builds(child)


def _refs_text(refs: Any) -> str:
    values = _collect_refs(refs)
    return "；".join(_ref_text(value) for value in values) if values else "待确认"


def _base(title: str, info: dict[str, Any], visibility: str) -> list[str]:
    audience = "客户可读" if visibility == "customer" else "内部工作稿"
    return ["---", "template_version: 1.0.0", f"project_id: {info['project_id']}", f"visibility: {visibility}", "status: discussion-draft", "---", "",
            f"# {title}", "", f"项目：{info['project_id']}；用途：{info.get('purpose', '待确认')}；可见范围：{audience} 待独立审核。", "",
            "本文为讨论稿，用于评审和补充事实，不构成签署、客户验收、付款触发或生产授权。", "",
            "运行时口径：正文读取登记的当前候选记录；有效基线仍按 manifest 的 runtime_effective 单独保留，候选与有效版本不混用。", ""]


def _kickoff(info: dict[str, Any], records: dict[str, dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
    recs = [x for x in (_first(records, "kickoff"), _first(records, "contract"), _first(records, "plan")) if x]
    data = _data(_first(records, "kickoff"), True)
    contract = _data(_first(records, "contract"), True)
    lines = _base("项目启动会", info, "customer")
    _section(lines, "目标与范围")
    lines.extend(_table(["项目目标", "合同范围内", "明确不包含或待确认"],
                        [["；".join(_display(x) for x in data.get("goals", [])),
                          _display(contract.get("scope_in")), _display(contract.get("scope_out"))]]))
    lines.append("")
    _section(lines, "组织与里程碑")
    lines.extend(_table(["角色/组织", "人员", "职责", "确认权限/依据", "状态"],
                        [[x.get("role", x.get("id")), x.get("actor"), x.get("responsibility", x.get("responsibilities", x.get("duty"))),
                          x.get("authority"), x.get("state")] for x in data.get("roles", [])]))
    lines.append("")
    lines.extend(_table(["里程碑 ID", "名称/说明", "来源"],
                        [[x.get("id", "待确认"), x.get("title", x.get("name")), _refs_text(x)] if isinstance(x, dict)
                         else [x, "待确认", "见登记里程碑"] for x in data.get("milestones", [])]))
    lines.append("")
    _section(lines, "客户准备与沟通")
    lines.extend(_table(["行动/材料", "责任人", "截止", "依赖/证据", "状态"],
                        [[x.get("action", x) if isinstance(x, dict) else x,
                          x.get("owner") if isinstance(x, dict) else "待确认",
                          x.get("due") if isinstance(x, dict) else "待确认",
                          _refs_text(x) if isinstance(x, dict) else "待确认",
                          x.get("state") if isinstance(x, dict) else "待确认"] for x in data.get("customer_actions", [])]))
    lines.append("")
    lines.append(f"沟通与升级：{_display(data.get('communication'))}")
    lines.append("")
    lines.append(f"待确认事项：{_display(data.get('open_items'))}")
    lines.append("")
    refs = _source_lines(lines, recs, customer=True)
    return "\n".join(lines) + "\n", refs


def _blueprint(info: dict[str, Any], records: dict[str, dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
    bp = _first(records, "blueprint")
    req = _first(records, "requirements")
    conf = _confirmation_for(records, bp)
    data = _data(bp, True)
    req_data = _data(req, True)
    recs = [x for x in (req, bp, conf) if x]
    lines = _base("业务蓝图", info, "customer")
    lines.extend(["本文只描述业务流程、页面字段、权限、数据、接口和验收场景，以及它们与需求的覆盖关系。详细实施方案、部署、安全和设计决策属于独立文档。", ""])
    _section(lines, "现状与目标边界")
    lines.extend(_table(["现状信息", "目标系统边界", "未决事项"], [[data.get("current_state"), data.get("system_boundary"), data.get("open_items")]]))
    lines.append("")
    _section(lines, "业务流程")
    lines.extend(_table(["流程 ID", "入口/触发", "正常步骤", "异常/恢复", "角色", "需求 IDs", "验收条件"],
                        [[x.get("id"), x.get("entry", x.get("trigger")), x.get("steps"), x.get("exceptions"),
                          x.get("roles"), x.get("requirement_ids"), x.get("acceptance")] for x in data.get("processes", [])]))
    lines.append("")
    _section(lines, "页面与字段")
    lines.extend(_table(["页面 ID", "页面", "需求 IDs"], [[x.get("id"), x.get("title"), x.get("requirement_ids")] for x in data.get("pages", [])]))
    lines.append("")
    lines.extend(_table(["字段 ID", "页面", "名称", "类型", "必填", "校验", "权限"],
                        [[x.get("id"), x.get("page_id"), x.get("name"), x.get("type"), x.get("required"), x.get("validation"), x.get("permission")] for x in data.get("fields", [])]))
    lines.append("")
    _section(lines, "权限、数据与接口")
    lines.extend(_table(["权限 ID", "角色", "动作", "条件", "字段 IDs"],
                        [[x.get("id"), x.get("role"), x.get("action"), x.get("condition"), x.get("field_ids")] for x in data.get("permissions", [])]))
    lines.append("")
    lines.extend(_table(["数据 ID", "名称", "业务键", "存储/边界", "字段 IDs", "规则"],
                        [[x.get("id"), x.get("name"), x.get("business_key"), x.get("storage"), x.get("field_ids"), x.get("rules")] for x in data.get("data", [])]))
    lines.append("")
    lines.extend(_table(["接口 ID", "来源系统", "目标系统", "字段映射", "验收"],
                        [[x.get("id"), x.get("source_system"), x.get("target_system"), x.get("field_mappings"), x.get("acceptance")] for x in data.get("interfaces", [])]))
    lines.append("")
    _section(lines, "验收场景")
    for scenario in data.get("acceptance_scenarios", []):
        lines.append(f"### 场景 {scenario.get('id', '待确认')}")
        lines.extend(_table(["场景规格", "要求与验证"], [[_label(key), value] for key, value in scenario.items()]))
        lines.append("")
    if not data.get("acceptance_scenarios"):
        lines.extend(["尚未登记验收场景，不能据此声称需求验收覆盖。", ""])
    _section(lines, "需求覆盖与确认范围")
    lines.extend(_table(["需求 ID", "原文/目标", "验收", "设计关联", "覆盖状态"],
                        [[x.get("id"), x.get("statement"), x.get("acceptance"), next((c.get("design_ids") for c in data.get("coverage",[]) if c.get("requirement_id")==x.get("id")),None), x.get("status")] for x in req_data.get("items", [])]))
    lines.append("")
    lines.extend(_table(["需求 ID", "流程/页面/字段/接口/验收 IDs", "权限 IDs", "数据 IDs", "覆盖状态/缺口"],
                        [[x.get("requirement_id"), x.get("design_ids"), x.get("permission_ids"), x.get("data_ids"), x.get("gap", "") or "已登记"] for x in data.get("coverage", [])]))
    lines.append("")
    if conf:
        c = _data(conf, True)
        lines.extend(_table(["确认记录", "状态", "确认人/角色", "日期/渠道", "来源", "范围", "遗留项"],
                            [[conf["artifact_id"], c.get("state"), f"{c.get('actor')} / {c.get('role')}", f"{c.get('date')} / {c.get('channel')}", c.get("origin"), c.get("scope"), c.get("open_items")]]))
        lines.append("")
    lines.append("确认状态仅反映运行时登记；讨论稿不代表真实客户签署或验收。")
    lines.append("")
    refs = _source_lines(lines, recs, customer=True)
    return "\n".join(lines) + "\n", refs


def _aggregate(records, kind):
    modules = _by_kind(records, kind)
    if not modules:
        return None
    value = dict(modules[0], data={})
    for module in modules:
        for key, child in module['data'].items():
            if isinstance(child, list):
                value['data'].setdefault(key, []).extend(child)
            elif key not in value['data']:
                value['data'][key] = child
    return value


def _solution(info: dict[str, Any], records: dict[str, dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
    sol = _aggregate(records, "solution")
    bp = _first(records, "blueprint")
    req = _first(records, "requirements")
    data = _data(sol)
    lines = _base("实施方案与设计决策", info, "internal")
    _section(lines, "设计边界与输入")
    lines.extend(_table(["需求 IDs", "组件/方案记录", "蓝图版本", "未决事项"], [[data.get("requirement_ids"), sol and sol["artifact_id"], _refs_text(bp), data.get("open_items")]]))
    lines.append("")
    _section(lines, "组件与运行设计")
    lines.extend(_table(["组件 ID", "职责/目的", "需求 IDs", "部署单元", "安全与权限", "非功能目标", "验证方法"],
                        [[x.get("id"), x.get("purpose"), x.get("requirement_ids"), x.get("deployment"), x.get("security"), x.get("nonfunctional"), x.get("verification") ] for x in data.get("components", [])]))
    lines.append("")
    _section(lines, "设计决策")
    lines.extend(_table(["决策 ID", "决策事项", "备选方案", "选定方案", "理由与约束", "状态", "证据"],
                        [[x.get("id"), x.get("question", x.get("decision")), x.get("options"), x.get("selected"), x.get("rationale"), x.get("status"), _refs_text(x.get("evidence_refs"))] for x in data.get("decisions", [])]))
    lines.append("")
    lines.append("实施设计与业务蓝图分开维护；本文件不把设计建议、ready 状态或本地验证表述为生产部署或客户确认。")
    lines.append("")
    for ui in _by_kind(records, 'ui_design'):
        detail = ui['data']
        _section(lines, '页面与交互规格 ' + ui['artifact_id'])
        lines.extend(_table(['页面 ID','页面名称','需求 IDs','业务状态'],
            [[r.get('id'),r.get('title'),r.get('requirement_ids'),r.get('states')] for r in detail.get('pages',[])]))
        lines.append('')
        for page in detail.get('pages',[]):
            _section(lines,'页面操作 ' + str(page.get('id','待确认')))
            for operation in page.get('operations',[]):
                lines.extend(_table(['操作规格','内容'],[[_label(key), value] for key,value in operation.items()]))
                lines.append('')
        _section(lines, '字段约束 ' + ui['artifact_id'])
        columns = [('id','字段 ID'),('page_id','页面'),('name','名称'),('type','类型'),('length','长度'),('precision','精度'),('unit','单位'),('required','必填'),('default','默认值'),('enum','可选值'),('validation','校验'),('source','来源'),('sensitivity','敏感级别'),('unique','唯一性'),('read_write','读写规则'),('interface_mapping','接口映射')]
        lines.extend(_table([x[1] for x in columns],[[r.get(x[0]) for x in columns] for r in detail.get('fields',[])]))
        lines.append('')
        _section(lines, '权限与状态转换 ' + ui['artifact_id'])
        lines.extend(_table(['权限 ID','页面/字段','角色','动作/条件','效果'],
            [[r.get('id'),[r.get('page_id'),r.get('field_ids')],r.get('role'),[r.get('action'),r.get('condition')],r.get('effect')] for r in detail.get('permissions',[])]))
        lines.append('')
        lines.extend(_table(['状态规则 ID','页面','原状态 → 目标状态','动作/允许角色','失败反馈'],
            [[r.get('id'),r.get('page_id'),[r.get('from'),r.get('to')],[r.get('action'),r.get('allowed_roles')],r.get('error')] for r in detail.get('states',[])]))
        lines.append('')
    refs = _source_lines(lines, [x for x in (req, bp) if x] + _by_kind(records,'solution') + _by_kind(records,'ui_design'), customer=False)
    return "\n".join(lines) + "\n", refs


def _plan(info: dict[str, Any], records: dict[str, dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
    plan = _first(records, "plan")
    data = _data(plan)
    lines = _base("项目总计划与 WBS", info, "internal")
    _section(lines, "计划基线与状态")
    lines.extend(_table(["计划版本", "进度截止", "日历", "关键路径", "来源/状态"], [[plan and plan["revision"], data.get("progress_as_of"), data.get("calendar"), data.get("critical_path"), plan and plan.get("status")]]))
    lines.append("")
    _section(lines, "WBS 工作包")
    headers = ["工作包 ID", "阶段", "标题", "需求 IDs", "完成规则", "负责人", "依赖", "基线开始", "基线完成", "当前计划开始", "当前计划完成", "实际开始", "实际完成", "剩余工作", "预测完成", "里程碑", "证据", "修订"]
    lines.extend(_table(headers, [[x.get("id"), x.get("phase"), x.get("title"), x.get("requirement_ids"), x.get("completion_rule"), x.get("owner"), x.get("dependencies"), x.get("baseline_start"), x.get("baseline_finish"), x.get("current_plan_start"), x.get("current_plan_finish"), x.get("actual_start"), x.get("actual_finish"), x.get("remaining_work"), x.get("forecast_finish"), x.get("milestone_id"), _refs_text(x.get("completion_evidence", x.get("evidence_refs"))), x.get("revision")] for x in data.get("work_packages", [])]))
    lines.append("")
    _section(lines, "里程碑与变更影响")
    lines.extend(_table(["里程碑 ID", "名称", "工作包 IDs", "验收/完成条件"], [[x.get("id"), x.get("title"), x.get("work_package_ids"), x.get("criterion", x.get("criteria"))] for x in data.get("milestones", [])]))
    lines.append("")
    lines.extend(_table(["变更/偏差 ID", "工作包", "基线偏差天数", "原因", "影响", "行动", "决策依据"], [[x.get("id"), x.get("work_package_id"), x.get("baseline_variance_days"), x.get("cause"), x.get("impact"), x.get("action"), _refs_text(x.get("decision_refs"))] for x in data.get("variance_analysis", [])]))
    lines.append("")
    lines.append("日期、工时和关键路径只在运行时已有事实时展示；缺少日历或工期依据不推算。")
    lines.append("")
    refs = _source_lines(lines, [plan] if plan else [], customer=False)
    return "\n".join(lines) + "\n", refs


def _integration(info: dict[str, Any], records: dict[str, dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
    integration = _aggregate(records, "integration")
    data = _data(integration)
    lines = _base("接口数据迁移与联调", info, "internal")
    _section(lines, "接口与字段映射")
    lines.extend(_table(["接口 ID", "来源系统", "目标系统", "方向", "触发", "协议", "认证", "环境", "状态"], [[x.get("id"), x.get("source_system"), x.get("target_system"), x.get("direction"), x.get("trigger"), x.get("protocol"), x.get("auth_method"), x.get("environment"), x.get("status")] for x in data.get("interfaces", [])]))
    lines.append("")
    lines.extend(_table(["映射 ID", "接口", "来源字段", "目标字段", "类型", "单位", "转换", "业务键"], [[x.get("id"), x.get("interface_id"), x.get("source_field"), x.get("target_field"), x.get("type"), x.get("unit"), x.get("transform"), x.get("business_key")] for x in data.get("mappings", [])]))
    lines.append("")
    _section(lines, "幂等、重试与联调核对")
    lines.extend(_table(["接口 ID", "幂等", "超时", "重试", "补偿", "状态"], [[x.get("id"), x.get("idempotency"), x.get("timeout"), x.get("retry"), x.get("compensation"), x.get("status")] for x in data.get("interfaces", [])]))
    lines.append("")
    lines.extend(_table(["核对 ID", "接口", "环境", "状态", "数量核对", "业务键核对"], [[x.get("id"), x.get("interface_id"), x.get("environment"), x.get("state"), x.get("counts"), x.get("business_keys")] for x in data.get("reconciliations", [])]))
    lines.append("")
    for reconciliation in data.get("reconciliations", []):
        lines.append(f"### 构建证据 {reconciliation.get('id', '待确认')}")
        build = reconciliation.get("build_ref")
        labels = {"repository":"代码仓库/位置", "revision":"构建修订/提交", "sha256":"完整 SHA256"}
        lines.extend(_table(["构建规格", "内容"], [[labels.get(key, _label(key)), value] for key,value in build.items()] if isinstance(build, dict) else [["构建",build]]))
        lines.append("")
    _section(lines, "数据迁移与完成边界")
    migration = data.get("migration", {})
    lines.extend(_table(["适用", "责任人", "状态", "备份", "回滚", "冻结窗口", "演练", "核对"], [[migration.get("applicable"), migration.get("owner"), migration.get("state"), _refs_text(migration.get("backup_refs")), migration.get("rollback_plan"), migration.get("freeze_window"), _refs_text(migration.get("rehearsal_refs")), _refs_text(migration.get("reconciliation_refs"))]]))
    lines.append("")
    lines.append("连接成功、文件生成或模拟数据不等于业务一致性或生产可用；未执行项必须继续列为待确认。")
    lines.append("")
    for module in _by_kind(records,'integration'):
        if module['data'].get('migration'):
            _section(lines, '迁移范围 ' + module['artifact_id'])
            migration = module['data']['migration']
            lines.extend(_table(['迁移项','内容'], [[_label(k), v] for k,v in migration.items()]))
            lines.append('')
    refs = _source_lines(lines, _by_kind(records,'integration'), customer=False)
    return "\n".join(lines) + "\n", refs


def _acceptance(info: dict[str, Any], records: dict[str, dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
    milestones = _by_kind(records, "milestone")
    target_keys = {_ref_key(_record_ref(r)) for r in [*milestones, *_by_kind(records,'blueprint')]}
    confirmations = [r for r in _by_kind(records, "confirmation")
                     if _ref(_data(r).get('target')) and _ref_key(_data(r)['target']) in target_keys]
    lines = _base("里程碑验收包", info, "customer")
    _section(lines, "范围与证据")
    for milestone in milestones:
        data = _data(milestone, True)
        _section(lines, f"里程碑 {data.get('milestone_id', milestone['artifact_id'])}")
        lines.extend(_table(["项目", "内容"], [
            ["范围", data.get("scope")],
            ["交付物版本与指纹", _refs_text(data.get("deliverable_refs"))],
            ["遗留项", data.get("open_items")],
            ["来源", data.get("origin")],
            ["真实客户验收", data.get("real_customer_acceptance")],
        ]))
        lines.append("")
        criteria = data.get("criteria", [])
        for index, criterion in enumerate(criteria if isinstance(criteria, list) else [criteria], 1):
            lines.append(f"### 验收标准 {index}")
            if isinstance(criterion, dict):
                lines.extend(_table(["项目", "要求与证据"], [[_label(key), value] for key, value in criterion.items()]))
            else:
                lines.append(_display(criterion))
            lines.append("")
    if not milestones:
        lines.append("尚未登记里程碑，范围、标准及交付证据待确认。")
    lines.append("")
    _section(lines, "确认记录")
    lines.extend(_table(["确认记录", "目标", "状态", "确认人/角色", "日期/渠道", "来源", "范围", "遗留项"], [[x["artifact_id"], _refs_text(_data(x, True).get("target")), _data(x, True).get("state"), f"{_data(x, True).get('actor')} / {_data(x, True).get('role')}", f"{_data(x, True).get('date')} / {_data(x, True).get('channel')}", _data(x, True).get("origin"), _data(x, True).get("scope"), _data(x, True).get("open_items")] for x in confirmations]))
    lines.append("")
    lines.append("本包为讨论稿。状态 accepted/conditional/rejected 仅复述运行时登记，不能推断签署者身份、真实客户验收或付款已触发。")
    lines.append("")
    refs = _source_lines(lines, [*milestones, *confirmations], customer=True)
    return "\n".join(lines) + "\n", refs


BUILDERS = {
    "kickoff": _kickoff,
    "blueprint": _blueprint,
    "solution": _solution,
    "plan": _plan,
    "integration": _integration,
    "milestone": _acceptance,
}


def _write_docx(markdown_path: Path, output: Path) -> None:
    global export_templates
    if export_templates is None:
        spec = importlib.util.spec_from_file_location("delivery_export_templates", SCRIPT_DIR / "export_templates.py")
        if spec is None or spec.loader is None:
            raise CoreExportError("unable to load DOCX template helpers")
        export_templates = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(export_templates)
    from docx import Document
    metadata, blocks = export_templates._parse_markdown(markdown_path)
    doc = Document()
    export_templates._style_document(doc)
    title_block = next((value for kind, value in blocks if kind == "heading" and value[0] == 1), None)
    title = title_block[1] if title_block else metadata.get("title", "交付文档")
    doc.add_paragraph(title, "Title")
    doc.add_paragraph(f"项目 {metadata.get('project_id', '待确认')}；状态 讨论稿；可见范围 {metadata.get('visibility', 'internal')}")
    for kind, value in blocks:
        if kind == "heading":
            level, text = value
            if level == 1:
                continue
            doc.add_paragraph(text, f"Heading {level}")
        elif kind == "paragraph":
            doc.add_paragraph(value)
        elif kind == "list":
            ordered, items = value
            for item in items:
                doc.add_paragraph(item, "List Number" if ordered else "List Bullet")
        elif kind == "code":
            _, code = value
            doc.add_paragraph(code, "Code")
        elif kind == "table":
            export_templates._add_table(doc, value)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        footer.add_run('第 ')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'),'PAGE')
        footer._p.append(field)
        footer.add_run(' 页')
    doc.core_properties.title = title
    doc.core_properties.subject = "Runtime delivery discussion draft"
    doc.core_properties.comments = "Generated from exact runtime project records"
    doc.save(output)


def export_core_documents(project: str | Path, out_dir: str | Path, mode: str = "discussion") -> dict[str, Any]:
    if mode != "discussion":
        raise CoreExportError("core document export only supports discussion drafts")
    root, snapshot, records = _load_project(project)
    out = Path(out_dir).expanduser().resolve()
    if out == root or out.is_relative_to(root / "00-control"):
        raise CoreExportError("output must be outside the runtime project control directory")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.mkdir(parents=True, exist_ok=True)
    names = [name + suffix for name, _title, _visibility, _kind in DOCS for suffix in (".md", ".docx")]
    names.append("core-documents-manifest.json")
    conflicts = [out / name for name in names if (out / name).exists() or (out / name).is_symlink()]
    if conflicts:
        raise CoreExportError("refusing to overwrite existing output: " + ", ".join(str(path) for path in conflicts))
    with tempfile.TemporaryDirectory(prefix=".core-documents-", dir=out) as temporary:
        staging = Path(temporary)
        manifest_docs: list[dict[str, Any]] = []
        for name, title, visibility, kind in DOCS:
            builder = BUILDERS[kind]
            markdown, refs = builder(dict(snapshot, **{key: snapshot.get(key) for key in ("project_id", "purpose")}), records)
            md = staging / f"{name}.md"
            md.write_text(markdown, encoding="utf-8")
            # Keep metadata available to the native authoring path without
            # leaking placeholders into the reader-facing DOCX.
            docx = staging / f"{name}.docx"
            _write_docx(md, docx)
            manifest_docs.append({"name": name, "title": title, "visibility": visibility,
                                  "status": "discussion-draft", "source_refs": refs,
                                  "markdown_sha256": _sha(md), "docx_sha256": _sha(docx)})
        manifest = {"manifest_version": "1.0.0", "project": {"project_id": snapshot["project_id"], "purpose": snapshot["purpose"]},
                    "mode": mode, "documents": manifest_docs,
                    "software_build_refs": [{"source_artifact": _record_ref(record), "build": build}
                                            for record in records.values() for build in _software_builds(record.get("data", {}))],
                    "runtime_effective": {key: snapshot["effective"][key] for key in sorted(snapshot.get("effective", {}))},
                    "alerts": sorted(snapshot.get("alerts", []))}
        manifest_path = staging / "core-documents-manifest.json"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        installed: list[Path] = []
        try:
            for source in [*staging.glob("*.md"), *staging.glob("*.docx"), manifest_path]:
                target = out / source.name
                if target.exists() or target.is_symlink():
                    raise CoreExportError(f"refusing to overwrite existing output: {target}")
                os.link(source, target)
                installed.append(target)
        except Exception:
            for target in installed:
                target.unlink(missing_ok=True)
            raise
    return {"out": str(out), "mode": mode, "project_id": snapshot["project_id"],
            "documents": [dict(item, markdown_path=str(out / (item["name"] + ".md")), docx_path=str(out / (item["name"] + ".docx"))) for item in manifest_docs],
            "manifest": str(out / "core-documents-manifest.json")}


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", required=True, help="runtime project directory containing 00-control")
    parser.add_argument("--out", required=True, help="new output directory")
    parser.add_argument("--mode", default="discussion", choices=["discussion"], help="document mode")
    args = parser.parse_args(argv)
    try:
        print(json.dumps(export_core_documents(args.project, args.out, args.mode), ensure_ascii=False, indent=2))
    except (CoreExportError, ValueError, OSError, KeyError, TypeError, ImportError) as exc:
        print(f"export_core_documents: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
