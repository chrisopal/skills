#!/usr/bin/env python3
"""Export the reusable delivery templates to editable DOCX and XLSX files.

The template index and source files are repository assets.  This command only
reads them and writes a fresh output directory; an existing output file is
never replaced.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
import tempfile
from pathlib import Path
from typing import Any, Iterable


SCRIPT_DIR = Path(__file__).resolve().parent
ORCHESTRATOR_DIR = SCRIPT_DIR.parent
TEMPLATES_DIR = ORCHESTRATOR_DIR / "assets" / "templates"
DEFAULT_INDEX = TEMPLATES_DIR / "index.json"

BLUE = "0966D9"
LIGHT_BLUE = "EDF4FF"
LIGHT_BORDER = "D9DEE7"
BODY_FONT = "Arial Unicode MS"
EAST_ASIA_FONT = "Arial Unicode MS"


class ExportError(ValueError):
    """Raised when an index or output cannot be safely exported."""


class TemplateEntry:
    __slots__ = ("template_id", "title", "visibility", "fmt", "path", "glossary")

    def __init__(self, template_id: str, title: str, visibility: str, fmt: str, path: Path, glossary: Path | None = None):
        self.template_id = template_id
        self.title = title
        self.visibility = visibility
        self.fmt = fmt
        self.path = path
        self.glossary = glossary


def _safe_id(value: Any, label: str) -> str:
    if not isinstance(value, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,99}", value):
        raise ExportError(f"{label} must be a safe template identifier")
    return value


def _rooted_template_path(raw: Any, root: Path) -> Path:
    if not isinstance(raw, str) or not raw.strip():
        raise ExportError("template path is required")
    candidate = Path(raw)
    # Accept the two useful representations while keeping index paths portable.
    parts = list(candidate.parts)
    for prefix in (("assets", "templates"), ("templates",)):
        if tuple(parts[: len(prefix)]) == prefix:
            parts = parts[len(prefix) :]
            candidate = Path(*parts)
            break
    if candidate.is_absolute() or ".." in candidate.parts:
        raise ExportError(f"template path escapes assets/templates: {raw}")
    resolved = (root / candidate).resolve()
    if not resolved.is_relative_to(root.resolve()):
        raise ExportError(f"template path escapes assets/templates: {raw}")
    if not resolved.is_file():
        raise ExportError(f"template file does not exist: {raw}")
    return resolved


def _load_entries(index_path: Path | None = None) -> list[TemplateEntry]:
    index_path = (index_path or DEFAULT_INDEX).resolve()
    if not index_path.is_file():
        raise ExportError(f"template index does not exist: {index_path}")
    try:
        value = json.loads(index_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ExportError(f"cannot read template index: {exc}") from exc
    raw_entries: Any
    if isinstance(value, list):
        raw_entries = value
    elif isinstance(value, dict):
        raw_entries = value.get("templates", value.get("entries"))
    else:
        raw_entries = None
    if not isinstance(raw_entries, list):
        raise ExportError("template index must contain a templates array")

    root = index_path.parent
    if root.name == "templates":
        template_root = root
    elif (root / "assets" / "templates").is_dir():
        template_root = root / "assets" / "templates"
    else:
        # A caller-supplied index may live beside its fixture assets.  The
        # default repository index still resolves from assets/templates.
        template_root = root
    entries: list[TemplateEntry] = []
    seen: set[str] = set()
    for raw in raw_entries:
        if not isinstance(raw, dict):
            raise ExportError("template index entries must be objects")
        template_id = _safe_id(raw.get("id"), "template id")
        if template_id in seen:
            raise ExportError(f"duplicate template id: {template_id}")
        seen.add(template_id)
        title = raw.get("title", template_id)
        if not isinstance(title, str) or not title.strip():
            raise ExportError(f"template title missing: {template_id}")
        visibility = raw.get("visibility", "internal")
        if visibility not in {"internal", "customer"}:
            raise ExportError(f"invalid visibility for {template_id}: {visibility}")
        fmt = raw.get("format", raw.get("fmt"))
        if fmt not in {"markdown", "csv"}:
            raise ExportError(f"invalid format for {template_id}: {fmt}")
        path = _rooted_template_path(raw.get("path"), template_root)
        if path.suffix.lower() != (".md" if fmt == "markdown" else ".csv"):
            raise ExportError(f"format does not match file extension for {template_id}")
        glossary_raw = raw.get("field_glossary", raw.get("glossary"))
        glossary = _rooted_template_path(glossary_raw, template_root) if glossary_raw else None
        entries.append(TemplateEntry(template_id, title.strip(), visibility, fmt, path, glossary))
    return entries


def _set_run_font(run: Any, size: int | None = None, bold: bool | None = None) -> None:
    from docx.oxml.ns import qn

    run.font.name = BODY_FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EAST_ASIA_FONT)
    if size is not None:
        run.font.size = __import__("docx").shared.Pt(size)
    if bold is not None:
        run.bold = bold


def _style_document(document: Any) -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt
    from docx.oxml.ns import qn

    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
        try:
            style = document.styles[name]
        except KeyError:
            continue
        style.font.name = BODY_FONT
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.font.color.rgb = __import__("docx").shared.RGBColor.from_string("0B1220")
    title = document.styles["Title"]
    title.font.size = Pt(22)
    document.styles['Normal'].font.size = Pt(10)
    document.styles['Normal'].paragraph_format.space_after = Pt(6)
    document.styles['Normal'].paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13)):
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        # Office document guidance keeps titles and section headings black;
        # blue is reserved for native table headers and the review UI.
        style.font.color.rgb = __import__("docx").shared.RGBColor.from_string("000000")
    if "Code" not in [style.name for style in document.styles]:
        code_style = document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "JetBrains Mono"
        code_style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "等宽")
        code_style.font.size = Pt(9)


def _shade_cell(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_borders(table: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            borders = tc_pr.first_child_found_in("w:tcBorders")
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                tc_pr.append(borders)
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                tag = "w:" + edge
                element = borders.find(qn(tag))
                if element is None:
                    element = OxmlElement(tag)
                    borders.append(element)
                element.set(qn("w:val"), "single")
                element.set(qn("w:sz"), "4")
                element.set(qn("w:color"), LIGHT_BORDER)


def _split_table_row(line: str) -> list[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    return [part.strip() for part in text.split("|")]


def _is_table_separator(line: str) -> bool:
    cells = _split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def _parse_markdown(path: Path) -> tuple[dict[str, str], list[tuple[str, Any]]]:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    metadata: dict[str, str] = {}
    if lines and lines[0].strip() == "---":
        end = next((i for i, line in enumerate(lines[1:], 1) if line.strip() == "---"), None)
        if end is not None:
            for line in lines[1:end]:
                if ":" in line:
                    key, value = line.split(":", 1)
                    metadata[key.strip()] = value.strip().strip('"\'')
            lines = lines[end + 1 :]
    blocks: list[tuple[str, Any]] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        fence = re.match(r"^\s*(```+|~~~+)(.*)$", line)
        if fence:
            marker = fence.group(1)[0]
            language = fence.group(2).strip()
            index += 1
            code: list[str] = []
            while index < len(lines) and not re.match(rf"^\s*{re.escape(marker)}{{3,}}", lines[index]):
                code.append(lines[index])
                index += 1
            if index < len(lines):
                index += 1
            blocks.append(("code", (language, "\n".join(code))))
            continue
        heading = re.match(r"^\s*(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            blocks.append(("heading", (min(len(heading.group(1)), 3), heading.group(2))))
            index += 1
            continue
        if index + 1 < len(lines) and "|" in line and _is_table_separator(lines[index + 1]):
            headers = _split_table_row(line)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip() and "|" in lines[index]:
                rows.append(_split_table_row(lines[index]))
                index += 1
            blocks.append(("table", [headers, *rows]))
            continue
        list_match = re.match(r"^\s*([-*+] |\d+[.)] )(.+)$", line)
        if list_match:
            ordered = list_match.group(1)[0].isdigit()
            items: list[str] = []
            while index < len(lines):
                current = re.match(r"^\s*([-*+] |\d+[.)] )(.+)$", lines[index])
                if not current or current.group(1)[0].isdigit() != ordered:
                    break
                items.append(current.group(2).strip())
                index += 1
            blocks.append(("list", (ordered, items)))
            continue
        paragraph: list[str] = [line.strip()]
        index += 1
        while index < len(lines) and lines[index].strip():
            current = lines[index]
            if re.match(r"^\s*(#{1,6})\s+", current) or re.match(r"^\s*([-*+] |\d+[.)] )", current):
                break
            if index + 1 < len(lines) and "|" in current and _is_table_separator(lines[index + 1]):
                break
            paragraph.append(current.strip())
            index += 1
        blocks.append(("paragraph", " ".join(paragraph)))
    return metadata, blocks


def _add_table(document: Any, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        from docx.oxml import OxmlElement
        properties = table.rows[row_index]._tr.get_or_add_trPr()
        properties.append(OxmlElement('w:cantSplit'))
        if row_index == 0:
            properties.append(OxmlElement('w:tblHeader'))
        for col_index in range(width):
            cell = table.cell(row_index, col_index)
            cell.text = row[col_index] if col_index < len(row) else ""
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = row_index == 0
                for run in paragraph.runs:
                    _set_run_font(run, 9 if row_index else 9, bold=row_index == 0)
            if row_index == 0:
                _shade_cell(cell, BLUE)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = __import__("docx").shared.RGBColor.from_string("FFFFFF")
            elif row_index % 2 == 0:
                _shade_cell(cell, LIGHT_BLUE)
    _set_cell_borders(table)
    document.add_paragraph()


def _write_docx(entry: TemplateEntry, output: Path) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_BREAK
        from docx.shared import RGBColor
    except ImportError as exc:
        raise ExportError("python-docx is required; run with the bundled workspace Python") from exc
    metadata, blocks = _parse_markdown(entry.path)
    document = Document()
    _style_document(document)
    document.core_properties.title = entry.title
    document.core_properties.subject = f"Delivery template {entry.template_id}"
    document.core_properties.comments = f"Template ID: {entry.template_id}; visibility: {entry.visibility}"
    title = document.add_paragraph(style="Title")
    title.add_run(entry.title)
    meta_rows = [
        ["模板 ID", entry.template_id],
        ["模板版本", metadata.get('template_version', '0.1.0')],
        ["可见范围", entry.visibility],
        ["项目 ID", "{{project_id}}"],
        ["状态", metadata.get('status', 'draft')],
        ["来源模板", entry.path.name],
    ]
    _add_table(document, meta_rows)
    for index, (kind, value) in enumerate(blocks):
        if kind == "heading":
            level, text = value
            if index == 0 and text == entry.title:
                continue
            paragraph = document.add_paragraph(style=f"Heading {level}")
            paragraph.add_run(text)
        elif kind == "paragraph":
            paragraph = document.add_paragraph(value)
            for run in paragraph.runs:
                _set_run_font(run, 10)
        elif kind == "list":
            ordered, items = value
            for item in items:
                paragraph = document.add_paragraph(style="List Number" if ordered else "List Bullet")
                paragraph.add_run(item)
        elif kind == "code":
            language, code = value
            paragraph = document.add_paragraph(style="Code")
            if language:
                lead = paragraph.add_run(f"[{language}]\n")
                lead.font.color.rgb = RGBColor.from_string(BLUE)
            paragraph.add_run(code)
            paragraph.paragraph_format.keep_together = True
        elif kind == "table":
            _add_table(document, value)
    document.save(output)


def _sheet_name(entry: TemplateEntry, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", entry.title).strip() or entry.template_id
    base = base[:31]
    candidate = base
    counter = 2
    while candidate in used:
        suffix = f"_{counter}"
        candidate = (base[: 31 - len(suffix)] + suffix)[:31]
        counter += 1
    used.add(candidate)
    return candidate


def _read_csv(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return [list(row) for row in csv.reader(stream)]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_xlsx(entries: list[TemplateEntry], output: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ExportError("openpyxl is required; run with the bundled workspace Python") from exc
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor=BLUE)
    header_font = Font(name=BODY_FONT, bold=True, color="FFFFFF")
    body_font = Font(name=BODY_FONT, color="0B1220")
    border = Border(*(Side(style="thin", color=LIGHT_BORDER) for _ in range(4)))
    used: set[str] = set()
    for entry in entries:
        sheet = workbook.create_sheet(_sheet_name(entry, used))
        rows = _read_csv(entry.path)
        if not rows:
            rows = [[entry.title]]
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2" if len(rows) > 1 else "A1"
        sheet.sheet_view.showGridLines = False
        max_column = max((len(row) for row in rows), default=1)
        if rows and max_column:
            sheet.auto_filter.ref = f"A1:{get_column_letter(max_column)}{len(rows)}"
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
        for column_index in range(1, max_column + 1):
            values = [str(row[column_index - 1]) if column_index - 1 < len(row) else "" for row in rows]
            width = min(60, max(12, max(len(value) for value in values) + 2))
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        sheet.row_dimensions[1].height = 28
        if entry.glossary:
            glossary_rows = _read_csv(entry.glossary)
            glossary = workbook.create_sheet(_sheet_name(TemplateEntry(entry.template_id + "_glossary", "字段说明", entry.visibility, "csv", entry.glossary), used))
            for row in glossary_rows:
                glossary.append(row)
            glossary.freeze_panes = "A2"
            glossary.sheet_view.showGridLines = False
            if glossary_rows:
                glossary.auto_filter.ref = f"A1:{get_column_letter(max(len(row) for row in glossary_rows))}{len(glossary_rows)}"
                for cell in glossary[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = border
                for row in glossary.iter_rows(min_row=2):
                    for cell in row:
                        cell.font = body_font
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                        cell.border = border
    workbook.properties.title = "Delivery templates"
    workbook.properties.subject = "Editable delivery template workbook"
    workbook.save(output)


def export_templates(out_dir: str | Path, index_path: str | Path | None = None) -> dict[str, Any]:
    entries = _load_entries(Path(index_path) if index_path else None)
    out = Path(out_dir).expanduser().resolve()
    out.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    markdown = [entry for entry in entries if entry.fmt == "markdown"]
    csv_entries = [entry for entry in entries if entry.fmt == "csv"]
    for entry in markdown:
        outputs.append(out / f"{entry.template_id}.docx")
    if csv_entries:
        outputs.append(out / "delivery-templates.xlsx")
    conflicts = [path for path in outputs if path.exists() or path.is_symlink()]
    if conflicts:
        names = ", ".join(str(path) for path in conflicts)
        raise ExportError(f"refusing to overwrite existing output: {names}")
    # Finish generation before publishing anything. Exclusive links reject files
    # and dangling symlinks created after preflight, without following them.
    with tempfile.TemporaryDirectory(prefix=".delivery-export-", dir=out) as temporary:
        staging = Path(temporary)
        for entry in markdown:
            _write_docx(entry, staging / f"{entry.template_id}.docx")
        if csv_entries:
            _write_xlsx(csv_entries, staging / "delivery-templates.xlsx")
        installed = []
        try:
            for path in outputs:
                os.link(staging / path.name, path)
                installed.append(path)
        except OSError as exc:
            for path in installed:
                path.unlink()
            raise ExportError(f"unable to publish without overwriting: {exc}") from exc
    output_manifest = []
    for entry in markdown:
        output = out / f"{entry.template_id}.docx"
        output_manifest.append(
            {
                "template_id": entry.template_id,
                "format": entry.fmt,
                "source_path": entry.path.relative_to(TEMPLATES_DIR).as_posix()
                if entry.path.is_relative_to(TEMPLATES_DIR)
                else entry.path.name,
                "output_path": str(output),
                "sha256": _sha256(output),
            }
        )
    if csv_entries:
        output = out / "delivery-templates.xlsx"
        output_hash = _sha256(output)
        for entry in csv_entries:
            output_manifest.append(
                {
                    "template_id": entry.template_id,
                    "format": entry.fmt,
                    "source_path": entry.path.relative_to(TEMPLATES_DIR).as_posix()
                    if entry.path.is_relative_to(TEMPLATES_DIR)
                    else entry.path.name,
                    "output_path": str(output),
                    "sha256": output_hash,
                }
            )
    return {
        "out": str(out),
        "docx": [str(out / f"{entry.template_id}.docx") for entry in markdown],
        "xlsx": str(out / "delivery-templates.xlsx") if csv_entries else None,
        "template_count": len(entries),
        "outputs": output_manifest,
    }


def main(argv: Iterable[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", required=True, help="new or existing directory for exported files")
    parser.add_argument("--index", help="override the portable assets/templates/index.json path")
    args = parser.parse_args(argv)
    try:
        result = export_templates(args.out, args.index)
    except (ExportError, OSError) as exc:
        print(f"export_templates: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
