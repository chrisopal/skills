import importlib.util
import json
import tempfile
import unittest
from pathlib import Path


SUITE = Path(__file__).resolve().parents[1]
RUNTIME_PATH = SUITE / "skills/delivery-orchestrator/scripts/runtime.py"
EXPORTER_PATH = SUITE / "skills/delivery-orchestrator/scripts/export_core_documents.py"


def load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class CoreDocumentExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name)
        self.project = self.base / "real-project"
        self.runtime = load(RUNTIME_PATH, "core_document_runtime")
        self.exporter = load(EXPORTER_PATH, "core_document_exporter")
        self.runtime.init_project(self.project, "REAL-CORE-001", purpose="internal")
        source = self.base / "requirements.txt"
        source.write_text("Customer requirement source R-01\n", encoding="utf-8")
        self.source_ref = self.runtime.source(self.project, source, "SRC-REQ")
        self.kickoff_ref = self.runtime.register(self.project, {
            "project_id": "REAL-CORE-001", "artifact_id": "KICKOFF", "revision": 1,
            "kind": "kickoff", "visibility": "customer", "status": "draft",
            "inputs": [self.source_ref],
            "data": {
                "goals": ["Confirm the delivery boundary"],
                "roles": [{"id": "PM", "role": "Project manager", "actor": "customer-pm",
                           "authority": "recorded role", "state": "proposed"}],
                "milestones": ["M-01"],
                "customer_actions": ["Provide sample data"],
                "communication": "Weekly review",
                "open_items": ["Confirm production environment"],
                "origin": "actual",
                "source_refs": [self.source_ref],
                "internal_rate": "900.00",
                "secret": "must never appear in customer view",
            },
        })

    def tearDown(self):
        self.temp.cleanup()

    def test_real_project_exports_six_native_documents_and_missing_facts(self):
        out = self.base / "core-documents"
        result = self.exporter.export_core_documents(self.project, out)
        self.assertEqual(result["project_id"], "REAL-CORE-001")
        self.assertEqual(len(result["documents"]), 6)
        manifest = json.loads((out / "core-documents-manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["project"]["purpose"], "internal")
        kickoff = next(item for item in manifest["documents"] if item["name"] == "01-initiation-kickoff")
        refs = {(item["artifact_id"], item["revision"], item["sha256"]) for item in kickoff["source_refs"]}
        self.assertIn(("KICKOFF", 1, self.kickoff_ref["sha256"]), refs)
        self.assertIn(("SRC-REQ", 1, self.source_ref["sha256"]), refs)

        customer_text = (out / "01-initiation-kickoff.md").read_text(encoding="utf-8")
        self.assertIn("discussion-draft", customer_text)
        self.assertIn("待确认", (out / "02-business-blueprint.md").read_text(encoding="utf-8"))
        self.assertNotIn("internal_rate", customer_text)
        self.assertNotIn("must never appear", customer_text)
        self.assertNotIn("{{", customer_text)

        from docx import Document
        doc = Document(out / "01-initiation-kickoff.docx")
        self.assertIn("项目启动会", doc.paragraphs[0].text)
        self.assertIn("REAL-CORE-001", "\n".join(p.text for p in doc.paragraphs))
        self.assertGreaterEqual(len(doc.tables), 3)
        self.assertTrue(any("版本与来源" in p.text for p in doc.paragraphs))

    def test_business_coverage_and_acceptance_keep_all_fields(self):
        def record(aid, kind, data):
            return dict(artifact_id=aid, kind=kind, data=data, revision=1, sha256='a'*64, inputs=[], status='draft', health='current', is_effective=False)
        records = {'BP': record('BP','blueprint', {'coverage':[{'requirement_id':'R-01','design_ids':['FLOW-1'],'permission_ids':['PERM-1'],'data_ids':['DATA-1'],'gap':'第三方认证待定'}]}),
                   'M': record('M','milestone', {'milestone_id':'M1','scope':['R-01'],'deliverable_refs':[], 'criteria':['审批通过'], 'open_items':['阻断缺陷BUG-9未关闭']})}
        bp, _ = self.exporter._blueprint({'project_id':'P','purpose':'simulation'},records)
        self.assertIn('DATA-1', bp)
        self.assertIn('PERM-1', bp)
        acceptance, _ = self.exporter._acceptance({'project_id':'P','purpose':'simulation'}, records)
        self.assertIn('阻断缺陷BUG-9未关闭', acceptance)
        md = self.base / 'wide.md'
        md.write_text(bp, encoding='utf-8')
        self.exporter._write_docx(md,self.base/'wide.docx')
        from docx import Document
        self.assertTrue(all(len(table.columns)<=5 for table in Document(self.base/'wide.docx').tables))

    def test_detail_export_keeps_multiple_modules_and_ui_constraints(self):
        def record(aid, kind, data):
            return dict(artifact_id=aid, kind=kind, data=data, revision=1, sha256='b'*64, inputs=[], status='draft', health='current', is_effective=False)
        records = {'S1':record('S1','solution',{'components':[{'id':'C1','purpose':'工单幂等处理'}]}),
                   'S2':record('S2','solution',{'components':[{'id':'C2','purpose':'独立ERP模块'}]}),
                   'UI':record('UI','ui_design', {'fields':[{'id':'F1','length':128,'validation':'不得重复','enum':['待审批','已审批']}],
                                               'states':[{'id':'ST1','from':'待审批','to':'已审批','error':'越权审批禁止'}]})}
        text, refs = self.exporter._solution({'project_id':'P','purpose':'internal'},records)
        for value in ['工单幂等处理','独立ERP模块','不得重复','128','越权审批禁止']:
            self.assertIn(value,text)
        self.assertEqual({r['artifact_id'] for r in refs},{'S1','S2','UI'})

    def test_plan_preserves_registered_completion_evidence(self):
        rec = dict(artifact_id='PLAN', kind='plan', revision=1, sha256='a'*64,
                   inputs=[], status='draft', health='current', is_effective=False,
                   data={'work_packages':[{'id':'WP1','evidence_refs':[self.source_ref]}]})
        text, _ = self.exporter._plan({'project_id':'P','purpose':'simulation'}, {'PLAN':rec})
        self.assertIn('SRC-REQ r1', text)

    def test_blueprint_scenarios_and_native_build_triple_are_not_lost(self):
        def rec(aid, kind, data):
            return dict(artifact_id=aid,kind=kind,data=data,revision=1,sha256='a'*64,
                        inputs=[],status='draft',health='current',is_effective=False)
        bp = rec('BP','blueprint',{'acceptance_scenarios':[{'id':'AS-1','steps':['提交并审批'],
                     'expected':'工单唯一且状态已批准','unknown_detail':'人工确认范围'}]})
        text, _ = self.exporter._blueprint({'project_id':'P','purpose':'internal'},{'BP':bp})
        for value in ['AS-1','提交并审批','工单唯一且状态已批准','人工确认范围']:
            self.assertIn(value,text)
        build={'repository':'https://example.com/build-evidence','revision':'commit-unique','sha256':'b'*64}
        integration=rec('INT','integration',{'reconciliations':[{'id':'RC-1','build_ref':build}]})
        text, _ = self.exporter._integration({'project_id':'P','purpose':'internal'},{'INT':integration})
        for value in build.values():self.assertIn(value,text)
        self.assertEqual(list(self.exporter._software_builds(integration["data"])),[build])

    def test_corrupt_or_stale_records_do_not_silently_export(self):
        path = self.runtime.artifact_path(self.project,self.kickoff_ref)
        payload=path.read_bytes()
        path.unlink()
        with self.assertRaises((ValueError,OSError)):
            self.exporter.export_core_documents(self.project,self.base/'corrupt')
        path.write_bytes(payload)
        replacement=self.base/'replacement.txt'
        replacement.write_text('改变后的来源',encoding='utf-8')
        self.runtime.source(self.project,replacement,'SRC-REQ')
        with self.assertRaises(ValueError):
            self.exporter.export_core_documents(self.project,self.base/'stale')

    def test_export_refuses_overwrite_after_first_success(self):
        out = self.base / "core-documents"
        self.exporter.export_core_documents(self.project, out)
        with self.assertRaisesRegex(self.exporter.CoreExportError, "refusing to overwrite"):
            self.exporter.export_core_documents(self.project, out)


if __name__ == "__main__":
    unittest.main()
