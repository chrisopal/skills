"""Contract tests for the reusable template exporters."""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills/delivery-orchestrator/scripts/export_templates.py"
OFFICE_AVAILABLE = bool(importlib.util.find_spec('docx') and importlib.util.find_spec('openpyxl'))


def load_module():
    spec = importlib.util.spec_from_file_location("delivery_export_templates", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class ExportTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exporter = load_module()

    @unittest.skipUnless(OFFICE_AVAILABLE, 'Office integration requires existing python-docx/openpyxl; use bundled workspace Python')
    def test_repository_index_exports_native_documents_and_workbook(self):
        from docx import Document
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            result = self.exporter.export_templates(directory)
            self.assertEqual(result["template_count"], 29)
            self.assertEqual(len(result["docx"]), 17)
            self.assertEqual(len(result["outputs"]), 29)
            self.assertTrue(all(item["template_id"] and item["output_path"] and len(item["sha256"]) == 64 for item in result["outputs"]))
            self.assertTrue(Path(result["xlsx"]).is_file())

            document = Document(result["docx"][0])
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("模板 ID", text)
            self.assertIn("{{project_id}}", text)
            self.assertTrue(any(paragraph.style.name.startswith("Heading") for paragraph in document.paragraphs))

            workbook = load_workbook(result["xlsx"])
            self.assertEqual(len(workbook.sheetnames), 12)
            sheet = workbook[workbook.sheetnames[0]]
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertTrue(sheet.auto_filter.ref)
            self.assertIsNotNone(sheet["A1"].value)
            self.assertTrue(sheet["A1"].fill.fgColor.rgb.endswith("0966D9"))

    def test_export_refuses_to_overwrite_existing_output(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "delivery-contract-handover-v1.docx"
            output.write_bytes(b"keep this file")
            with self.assertRaises(self.exporter.ExportError):
                self.exporter.export_templates(directory)
            self.assertEqual(output.read_bytes(), b"keep this file")

    def test_dangling_output_symlink_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            out = Path(directory) / 'kit'
            out.mkdir()
            outside = Path(directory) / 'outside.docx'
            (out / 'delivery-contract-handover-v1.docx').symlink_to(outside)
            with self.assertRaises(self.exporter.ExportError):
                self.exporter.export_templates(out)
            self.assertFalse(outside.exists())

    @unittest.skipUnless(OFFICE_AVAILABLE, 'Office integration requires existing python-docx/openpyxl; use bundled workspace Python')
    def test_index_paths_are_portable_and_traversal_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "template.md"
            source.write_text("# Heading\n\nA paragraph.", encoding="utf-8")
            index = Path(directory) / "index.json"
            index.write_text(
                json.dumps({"templates": [{"id": "portable", "title": "Portable", "format": "markdown", "path": "template.md", "visibility": "customer"}]}),
                encoding="utf-8",
            )
            with tempfile.TemporaryDirectory() as output:
                result = self.exporter.export_templates(output, index)
                self.assertTrue(Path(result["docx"][0]).is_file())

            bad = Path(directory) / "bad.json"
            bad.write_text(
                json.dumps({"templates": [{"id": "bad", "title": "Bad", "format": "markdown", "path": "../template.md", "visibility": "internal"}]}),
                encoding="utf-8",
            )
            with self.assertRaises(self.exporter.ExportError):
                self.exporter._load_entries(bad)


if __name__ == "__main__":
    unittest.main()
