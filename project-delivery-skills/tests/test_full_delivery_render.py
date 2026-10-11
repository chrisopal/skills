"""Review rendering preserves pinned evidence and creates usable business outputs."""
from __future__ import annotations

import importlib.util
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from html.parser import HTMLParser

SUITE = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location('delivery_full_render', SUITE / 'scripts/render_full_delivery.py')
renderer = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(renderer)


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self.ids = set()

    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        if 'id' in values:
            self.ids.add(values['id'])
        if tag == 'a':
            self.links.append(values.get('href', ''))


def fixture(path):
    path.mkdir()
    project = path / 'project'
    renderer.runtime.init_project(project, 'RENDER-SIM-001', purpose='simulation')
    source_path = path / 'source.md'
    source_path.write_text('模拟来源 <script>alert(1)</script>\n', encoding='utf-8')
    source = renderer.runtime.source(project, source_path, 'SOURCE')
    artifact = {
        'project_id': 'RENDER-SIM-001', 'artifact_id': 'REQ', 'revision': 1,
        'kind': 'requirements', 'visibility': 'internal', 'status': 'ready', 'inputs': [source],
        'data': {'items': [{'id': 'R1', 'statement': '<script>alert("x")</script> =HYPERLINK("evil")',
                            'source_refs': [source], 'priority': 'must', 'acceptance': '工单状态回读一致'}], 'open_items': []},
    }
    req = renderer.runtime.register(project, artifact)
    report = {'skill_runs': [{'skill': 'delivery-requirements-baseline', 'goal': '提取工单需求',
                                'input_refs': [source], 'output_refs': [req], 'evidence_refs': [source],
                                'unresolved': ['真实客户尚未确认']}]}
    (path / 'report.json').write_text(json.dumps(report, ensure_ascii=False), encoding='utf-8')
    return source, req


class FullRenderTests(unittest.TestCase):
    def test_text_does_not_dump_json_and_preserves_reference_identity(self):
        self.assertEqual(renderer.text({'artifact_id': 'REQ', 'revision': 2, 'sha256': 'a' * 64}), 'REQ r2')
        self.assertEqual(renderer.text({'owner': '张工', 'hours': 8}), '负责人：张工；工时：8')
        self.assertEqual(renderer.esc('<script>"x"</script>'), '&lt;script&gt;&quot;x&quot;&lt;/script&gt;')
        self.assertEqual(renderer.trace_rows([]), [])
        self.assertEqual(renderer.text({'billing_mode': 'fixed', 'mode': 'fixed', 'state': 'fixed'}), '计费模式：固定价；mode：固定价；业务状态：已修复待复测')

    def test_runtime_readback_populates_safe_linked_sections_and_exact_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            source, req = fixture(base / 'case')
            out = base / 'review'
            def office_stub(path, *args):
                path.write_bytes(b'Office authoring delegated to marked integration run')
            with patch.object(renderer, 'write_docx', side_effect=office_stub), patch.object(renderer, 'write_xlsx', side_effect=office_stub):
                manifest = renderer.render(base / 'case', out)
            self.assertEqual(len(manifest['documents']), 13)
            self.assertEqual(len(manifest['task_cards']), 25)
            self.assertEqual(manifest['template_index_count'], 58)
            output = (out / 'index.html').read_text(encoding='utf-8')
            self.assertIn('&lt;script&gt;alert', output)
            self.assertNotIn('<script>alert(', output)
            self.assertIn('真实客户验收', output)
            self.assertIn('没有匹配记录', output)
            self.assertIn("localStorage.setItem", output)
            parser = Links()
            parser.feed(output)
            for target in parser.links:
                if target.startswith('#'):
                    self.assertIn(target[1:], parser.ids)
                else:
                    self.assertFalse('://' in target or target.startswith(('/', 'javascript:', 'data:')))
                    self.assertTrue((out / target).is_file(), target)
            for _, title, _, _ in renderer.GROUPS:
                self.assertIn(title, output)
            reqdoc = next(d for d in manifest['documents'] if d['title'] == '需求基线与业务蓝图')
            self.assertEqual(reqdoc['artifact_refs'], [req])
            self.assertIn('工单状态回读一致', (out / reqdoc['markdown']).read_text(encoding='utf-8'))
            raw = out / 'raw/REQ-r1.json'
            self.assertEqual(renderer.sha(raw), req['sha256'])
            sourceinfo = next(a for a in manifest['artifacts'] if a['artifact_id'] == 'SOURCE')
            self.assertEqual(renderer.sha(out / sourceinfo['source_file']),
                             json.loads((out / sourceinfo['path']).read_text())['data']['sha256'])
            card = (out / 'tasks/delivery-requirements-baseline.md').read_text(encoding='utf-8')
            self.assertIn('提取工单需求', card)
            self.assertIn('真实客户尚未确认', card)
            self.assertIn('delivery-requirements-baseline-v1', card)
            self.assertTrue(all((out / d['path']).stat().st_size for d in manifest['documents']))
            with self.assertRaisesRegex(ValueError, 'fresh directory'):
                renderer.render(base / 'case', out)

    def test_workbook_spec_contains_typed_summary_and_distinct_ledgers(self):
        records = [{'artifact_id': 'HOURS', 'kind': 'timesheets', 'revision': 2,
                    'status': 'ready', 'health': 'current', 'is_effective': True, 'sha256': 'a' * 64,
                    'data': {'items': [{'id': 'H1', 'hours': 8, 'rate': 100, 'person': 'analyst-sim'}]}}]
        spec = renderer.workbook_spec(records, [], {'project_id': 'TEST', 'cost_summary': {'budget': '8500.00', 'forecast': '6600.00', 'complete': True}})
        self.assertEqual(spec['sheets'][0]['name'], '项目概览')
        overview = {row[0]: row for row in spec['sheets'][0]['rows']}
        self.assertEqual(overview['批准预算'][1], 8500.0)
        self.assertIsInstance(overview['完工预测'][1], float)
        hours = next(sheet for sheet in spec['sheets'] if sheet['name'] == '工时明细')
        self.assertEqual(hours['rows'][0][hours['keys'].index('hours')], 8)
        self.assertEqual(hours['rows'][0][hours['keys'].index('rate')], 100)
        self.assertFalse(any(sheet['name'] == '资源工时成本与差旅台账' for sheet in spec['sheets']))

    def test_xlsx_authoring_uses_bundled_artifact_builder_and_temporary_spec(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            fixture(base / 'case')
            _, snapshot, records, _ = renderer.load_case(base / 'case')
            with patch.object(renderer.subprocess, 'run') as run:
                renderer.write_xlsx(base / 'tracking.xlsx', records, [], snapshot)
            args = run.call_args.args[0]
            self.assertIn('/.cache/codex-runtimes/', args[0])
            self.assertTrue(args[1].endswith('build_full_delivery_workbook.mjs'))
            self.assertIn('--spec', args)
            self.assertFalse(Path(args[args.index('--spec') + 1]).exists())
            self.assertTrue(run.call_args.kwargs['check'])

    def test_corrupt_runtime_artifact_is_rejected_without_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            _, req = fixture(base / 'case')
            renderer.runtime.artifact_path(base / 'case/project', req).write_text('{}', encoding='utf-8')
            with self.assertRaises((ValueError, KeyError)):
                renderer.render(base / 'case', base / 'review')
            self.assertFalse((base / 'review/index.html').exists())

    def test_existing_output_symlink_is_not_followed(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            fixture(base / 'case')
            (base / 'real').mkdir()
            (base / 'review').symlink_to(base / 'real', target_is_directory=True)
            with self.assertRaisesRegex(ValueError, 'already exists'):
                renderer.render(base / 'case', base / 'review')
            self.assertEqual(list((base / 'real').iterdir()), [])

    @unittest.skipUnless(os.environ.get('DELIVERY_OFFICE_TESTS') == '1', 'Caller must mark authoring before native Office integration')
    def test_native_office_outputs_have_filled_business_content(self):
        from docx import Document
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            fixture(base / 'case')
            manifest = renderer.render(base / 'case', base / 'review')
            for entry in manifest['documents']:
                doc = Document(base / 'review' / entry['path'])
                self.assertEqual(doc.paragraphs[0].style.name, 'Title')
                self.assertIn('模拟演练', '\n'.join(p.text for p in doc.paragraphs))
            reqdoc = next(d for d in manifest['documents'] if d['title'] == '需求基线与业务蓝图')
            self.assertIn('工单状态回读一致', '\n'.join(p.text for p in Document(base / 'review' / reqdoc['path']).paragraphs))
            book = load_workbook(base / 'review/delivery-tracking.xlsx')
            self.assertEqual(book.sheetnames[0], '项目概览')
            self.assertIn('需求清单', book.sheetnames)
            self.assertGreater(book['需求设计测试追溯'].max_row, 1)
            self.assertEqual(book['版本清单'].freeze_panes, 'A2')
            self.assertEqual(book['需求设计测试追溯']['A2'].value, 'R1')
            self.assertTrue(book['需求设计测试追溯'].tables)


if __name__ == '__main__':
    unittest.main()
