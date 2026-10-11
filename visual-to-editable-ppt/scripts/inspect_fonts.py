#!/usr/bin/env python3
"""Read stored PPTX font values; inherited theme/master values remain unknown."""
import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def text_frames(shapes, prefix=''):
    for index, shape in enumerate(shapes, 1):
        location = f'{prefix}{index}:{shape.name}'
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from text_frames(shape.shapes, location + '/')
        elif shape.has_text_frame:
            yield location, shape.text_frame
        elif shape.has_table:
            for row_i, row in enumerate(shape.table.rows, 1):
                for col_i, cell in enumerate(row.cells, 1):
                    yield f'{location}/cell[{row_i},{col_i}]', cell.text_frame


def inspect_fonts(path: Path) -> dict:
    presentation = Presentation(path)
    records = []
    for slide_i, slide in enumerate(presentation.slides, 1):
        for location, frame in text_frames(slide.shapes):
            for para_i, paragraph in enumerate(frame.paragraphs, 1):
                for run_i, run in enumerate(paragraph.runs, 1):
                    if not run.text.strip():
                        continue
                    explicit = run.font.size
                    default = paragraph.font.size
                    size = explicit if explicit is not None else default
                    rpr = run._r.rPr
                    ea = rpr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea') if rpr is not None else None
                    records.append({
                        'slide': slide_i, 'shape': location, 'paragraph': para_i, 'run': run_i,
                        'text': run.text, 'size_pt': round(size.pt, 2) if size is not None else None,
                        'size_source': 'run' if explicit is not None else 'paragraph' if default is not None else 'unknown_inherited',
                        'latin_font': run.font.name or paragraph.font.name,
                        'east_asian_font': ea.get('typeface') if ea is not None else None,
                    })
    return {'pptx': str(path.resolve()), 'slides': len(presentation.slides),
            'slide_inches': [round(presentation.slide_width.inches, 6), round(presentation.slide_height.inches, 6)],
            'text_runs': records,
            'unknown_size_runs': sum(record['size_pt'] is None for record in records),
            'limits': ['Stored values only; does not resolve theme/master inheritance or application font substitution.',
                       'Match text to approved roles before judging size; cover titles and body have different ranges.',
                       'PowerPoint/LibreOffice rendering and visual inspection are still required.']}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('pptx', type=Path)
    args = parser.parse_args()
    print(json.dumps(inspect_fonts(args.pptx), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
