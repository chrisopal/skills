"""Synthetic contracts only: these are not human reviews or image quality tests."""
import copy
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('visual_bridge', ROOT / 'scripts/visual_bridge.py')
b = importlib.util.module_from_spec(spec)
spec.loader.exec_module(b)

class BridgeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.p = b.load_pipeline(ROOT.parent / 'consulting-ppt-image')
        self.pack = self.p.load_json(ROOT.parent / 'consulting-ppt-image/examples/minimal/slide-content-pack.json')
        self.style = self.p.load_json(ROOT / 'assets/themes/consulting-purple.json')
        self.pack['theme_profile'] = self.style['id']
        self.pack['approval_ref'] = 'SYNTHETIC TEST APPROVAL'
        self.pack['slides'][0]['chart_spec'] = {'unit': '万元', 'values': [12, 24], 'labels': ['A', 'B']}
        for s in self.pack['slides']:
            s['review_status'] = 'approved'
        self.project = self.base / 'project with spaces'
        self.pack_path = self.base / 'pack.json'
        self.p.save_json(self.pack_path, self.pack)
        self.style_path = self.base / 'style.json'
        self.p.save_json(self.style_path, self.style)
        self.p.init_project(self.pack_path, self.project, self.style_path)

    def ready(self, sid, version=None):
        path = self.base / f'{sid}-{version}.png'
        Image.new('RGB', (1600, 900), (int(sid[1:])*30, (version or 1)*40, 90)).save(path)
        v = self.p.register_image(self.project, sid, path, version, 'synthetic fixture')
        review = {'slide_id': sid, 'version': v['version'], 'artifact_sha256': v['sha256'],
                  'reviewer': 'SYNTHETIC UNIT TEST', 'checks': {k: True for k in self.p.CHECKS}, 'notes': 'not visual acceptance'}
        rp = self.base / 'review.json'
        self.p.save_json(rp, review)
        self.p.record_review(self.project, rp)
        self.p.select_image(self.project, sid, v['version'], 'SYNTHETIC TEST SELECTION')
        return v

    def test_theme_and_exact_content_survive_plan(self):
        q = b.plan(self.p, self.project, '1', False)
        prompt = Path(q['items'][0]['prompt']).read_text()
        self.assertIn(self.style['palette']['primary'], prompt)
        self.assertIn('万元', prompt)
        self.assertIn(self.pack['slides'][0]['title'], prompt)
        self.assertNotIn('墨绿色标题', prompt)
        self.assertEqual(q['items'][0]['slide_id'], 'S001')

    def test_draft_cannot_generate(self):
        self.pack['version'] = '0.2.0'
        self.pack['slides'][0]['review_status'] = 'draft'
        self.p.save_json(self.pack_path, self.pack)
        self.p.sync_pack(self.project, self.pack_path, 'SYNTHETIC change')
        with self.assertRaises(self.p.PipelineError): b.plan(self.p, self.project, '1', False)
        self.assertEqual(b.plan(self.p, self.project, '1', True)['items'][0]['status'], 'draft_only')

    def test_partial_handoff_preserves_mapping_data_and_hash(self):
        self.ready('S001'); self.ready('S003')
        out = self.base / 'handoff'
        h = b.handoff(self.p, self.project, out, '1,3')
        self.assertEqual(h['scope'], 'explicit_partial:1,3')
        self.assertEqual([(p['page_id'], p['slide_id'], p['order']) for p in h['pages']],
                         [('page_001', 'S001', 1), ('page_002', 'S003', 3)])
        text = self.p.load_json(out / h['pages'][0]['content_file'])
        self.assertEqual(text['slide']['chart_spec'], self.pack['slides'][0]['chart_spec'])
        self.assertEqual(text['slide'], self.pack['slides'][0])
        self.assertEqual(self.p.file_digest(out / h['pages'][0]['image']), h['pages'][0]['sha256'])
        self.assertEqual(b.verify_handoff(self.p, self.project, out)['pages'], 2)

    def test_unselected_newer_version_is_not_used(self):
        self.ready('S001', 1)
        path = self.base / 'new.png'
        Image.new('RGB', (1600, 900), 'red').save(path)
        self.p.register_image(self.project, 'S001', path, 2, 'test')
        h = b.handoff(self.p, self.project, self.base / 'handoff', '1')
        self.assertEqual(h['pages'][0]['selected_version'], 1)

    def test_missing_pages_block_full_export(self):
        self.ready('S001')
        with self.assertRaises(self.p.PipelineError): b.handoff(self.p, self.project, self.base / 'handoff')
        self.assertFalse((self.base / 'handoff').exists())

    def test_existing_handoff_is_never_overwritten(self):
        self.ready('S001'); out = self.base / 'handoff'
        b.handoff(self.p, self.project, out, '1')
        before = (out / 'handoff.json').read_bytes()
        with self.assertRaises(self.p.PipelineError): b.handoff(self.p, self.project, out, '1')
        self.assertEqual(before, (out / 'handoff.json').read_bytes())

    def test_changed_selected_image_blocks_handoff(self):
        v = self.ready('S001')
        (self.project / v['path']).write_bytes(b'changed')
        with self.assertRaises(self.p.PipelineError): b.handoff(self.p, self.project, self.base / 'handoff', '1')

    def test_revision_invalidates_existing_handoff(self):
        self.ready('S001'); out = self.base / 'handoff'
        b.handoff(self.p, self.project, out, '1')
        self.pack['version'] = '0.2.0'; self.pack['slides'][0]['title'] = '新的已确认标题'
        self.p.save_json(self.pack_path, self.pack)
        self.p.sync_pack(self.project, self.pack_path, 'SYNTHETIC change')
        with self.assertRaises(self.p.PipelineError): b.verify_handoff(self.p, self.project, out)

    def test_tampered_handoff_image_is_rejected(self):
        self.ready('S001'); out = self.base / 'handoff'
        h = b.handoff(self.p, self.project, out, '1')
        (out / h['pages'][0]['image']).write_bytes(b'changed')
        with self.assertRaises(self.p.PipelineError): b.verify_handoff(self.p, self.project, out)

    def test_tampered_content_and_mapping_are_rejected(self):
        self.ready('S001'); out = self.base / 'handoff'
        h = b.handoff(self.p, self.project, out, '1')
        h['pages'][0]['order'] = 2
        self.p.save_json(out / 'handoff.json', h)
        with self.assertRaises(self.p.PipelineError): b.verify_handoff(self.p, self.project, out)

    def test_all_seven_themes_work_with_existing_state(self):
        themes = list((ROOT / 'assets/themes').glob('*.json'))
        self.assertEqual(len(themes), 7)
        for path in themes:
            style = self.p.load_json(path)
            pack = copy.deepcopy(self.pack); pack['theme_profile'] = style['id']
            self.p.save_json(self.pack_path, pack)
            project = self.base / style['id']
            self.p.init_project(self.pack_path, project, path)
            q = b.plan(self.p, project, '2', False)
            self.assertIn(style['palette']['primary'], Path(q['items'][0]['prompt']).read_text())

    def test_bad_dependency_path_is_clear(self):
        with self.assertRaisesRegex(ValueError, 'consulting-ppt-image'):
            b.load_pipeline(self.base / 'missing')

    def make_pptx(self, out, flattened=False, wrong=False):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation(); prs.slide_width = Inches(13.3333); prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if flattened:
            h = self.p.load_json(out / 'handoff.json')
            slide.shapes.add_picture(str(out / h['pages'][0]['image']), 0, 0, prs.slide_width, prs.slide_height)
        else:
            values = b.required_text(self.pack['slides'][0])
            if wrong: values[0] = 'WRONG TITLE'
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(5))
            box.text_frame.text = '\n'.join(values)
        result = self.base / 'deck.pptx'; prs.save(result)
        return result

    def test_final_pptx_has_real_text_and_no_source_raster(self):
        self.ready('S001'); out = self.base / 'handoff'; b.handoff(self.p, self.project, out, '1')
        deck = self.make_pptx(out)
        report = b.audit_pptx(self.p, self.project, out, deck)
        self.assertTrue(report['passed']); self.assertFalse(report['visual_acceptance'])

    def test_final_pptx_wrong_text_fails(self):
        self.ready('S001'); out = self.base / 'handoff'; b.handoff(self.p, self.project, out, '1')
        report = b.audit_pptx(self.p, self.project, out, self.make_pptx(out, wrong=True))
        self.assertFalse(report['passed'])
        self.assertTrue(any('missing authoritative text' in e for e in report['errors']))

    def test_flattened_pptx_is_rejected(self):
        self.ready('S001'); out = self.base / 'handoff'; b.handoff(self.p, self.project, out, '1')
        report = b.audit_pptx(self.p, self.project, out, self.make_pptx(out, flattened=True))
        self.assertFalse(report['passed'])
        self.assertTrue(any('source raster' in e for e in report['errors']))

    def test_handoff_content_file_tampering_is_rejected(self):
        self.ready('S001'); out = self.base / 'handoff'; h = b.handoff(self.p, self.project, out, '1')
        cp = out / h['pages'][0]['content_file']; data = self.p.load_json(cp)
        data['slide']['chart_spec']['values'] = [99, 100]; self.p.save_json(cp, data)
        with self.assertRaises(self.p.PipelineError): b.verify_handoff(self.p, self.project, out)

    def test_handoff_carrier_preserves_notes_and_original_images(self):
        from pptx import Presentation
        self.pack['slides'][0]['speaker_notes'] = '演讲备注：成本仅为示例。'
        self.pack['version'] = '0.2.0'; self.p.save_json(self.pack_path, self.pack)
        self.p.sync_pack(self.project, self.pack_path, 'SYNTHETIC notes')
        self.ready('S001'); out = self.base / 'handoff'
        h = b.handoff(self.p, self.project, out, '1')
        prs = Presentation(out / h['input_pptx'])
        self.assertEqual(prs.slides[0].notes_slide.notes_text_frame.text, self.pack['slides'][0]['speaker_notes'])
        self.assertEqual(len(prs.slides[0].shapes), 1)

    def test_changed_carrier_notes_fail_readback(self):
        from pptx import Presentation
        self.ready('S001'); out = self.base / 'handoff'; h = b.handoff(self.p, self.project, out, '1')
        carrier = out / h['input_pptx']; prs = Presentation(carrier)
        prs.slides[0].notes_slide.notes_text_frame.text = 'tampered'; prs.save(carrier)
        with self.assertRaises(self.p.PipelineError): b.verify_handoff(self.p, self.project, out)

    def test_changed_carrier_geometry_is_rejected(self):
        from pptx import Presentation
        self.ready('S001'); out = self.base / 'handoff'; h = b.handoff(self.p, self.project, out, '1')
        carrier = out / h['input_pptx']; prs = Presentation(carrier)
        prs.slides[0].shapes[0].width //= 2; prs.save(carrier)
        with self.assertRaises(self.p.PipelineError): b.verify_handoff(self.p, self.project, out)

    def test_final_pptx_missing_notes_is_rejected(self):
        self.pack['slides'][0]['speaker_notes'] = '原稿备注'
        self.pack['version'] = '0.2.0'; self.p.save_json(self.pack_path, self.pack)
        self.p.sync_pack(self.project, self.pack_path, 'SYNTHETIC notes')
        self.ready('S001'); out = self.base / 'handoff'; b.handoff(self.p, self.project, out, '1')
        report = b.audit_pptx(self.p, self.project, out, self.make_pptx(out))
        self.assertFalse(report['passed'])
        self.assertTrue(any('notes' in e for e in report['errors']))

    def test_carrier_is_accepted_by_real_converter_normalization(self):
        runtime = ROOT.parent / 'image-to-editable-ppt/skills/image-to-editable-ppt/cli/editppt/runtime/_input_normalization.py'
        spec = importlib.util.spec_from_file_location('normalization', runtime)
        normalizer = importlib.util.module_from_spec(spec); spec.loader.exec_module(normalizer)
        self.pack['slides'][0]['speaker_notes'] = '核对成本假设。'
        self.pack['version'] = '0.2.0'; self.p.save_json(self.pack_path, self.pack)
        self.p.sync_pack(self.project, self.pack_path, 'SYNTHETIC notes')
        self.ready('S001'); self.ready('S003'); out = self.base / 'handoff'
        h = b.handoff(self.p, self.project, out, '1,3')
        deck_path = normalizer.normalize_inputs([out / h['input_pptx']], job_dir=self.base / 'convert')
        deck = self.p.load_json(deck_path)
        self.assertEqual(deck['page_count'], 2)
        notes = self.p.load_json(deck_path.parent / deck['notes_manifest'])
        self.assertEqual(notes['notes'][0]['text'], '核对成本假设。')
        for page, original in zip(deck['pages'], h['pages']):
            with Image.open(deck_path.parent / page['source_image']) as got, Image.open(out / original['image']) as expected:
                self.assertEqual(got.tobytes(), expected.tobytes())

    def test_near_widescreen_source_uses_actual_ratio_in_carrier(self):
        from pptx import Presentation
        source = self.base / 'near.png'; Image.new('RGB', (1601, 900), 'purple').save(source)
        v = self.p.register_image(self.project, 'S001', source, None, 'synthetic')
        rp = self.base / 'review.json'
        self.p.save_json(rp, {'slide_id':'S001','version':1,'artifact_sha256':v['sha256'],
                             'reviewer':'TEST','checks':{k:True for k in self.p.CHECKS},'notes':'TEST'})
        self.p.record_review(self.project, rp); self.p.select_image(self.project, 'S001', 1, 'TEST')
        out = self.base / 'handoff'; h = b.handoff(self.p, self.project, out, '1')
        prs = Presentation(out / h['input_pptx'])
        self.assertAlmostEqual(prs.slide_width/prs.slide_height, 1601/900, places=6)
        self.assertEqual(prs.slides[0].shapes[0].left, 0)

    def test_mixed_source_ratios_block_before_creating_handoff(self):
        self.ready('S001')
        source = self.base / 'near.png'; Image.new('RGB', (1601, 900), 'purple').save(source)
        v = self.p.register_image(self.project, 'S003', source, None, 'synthetic')
        rp = self.base / 'review.json'
        self.p.save_json(rp, {'slide_id':'S003','version':1,'artifact_sha256':v['sha256'],
                             'reviewer':'TEST','checks':{k:True for k in self.p.CHECKS},'notes':'TEST'})
        self.p.record_review(self.project, rp); self.p.select_image(self.project, 'S003', 1, 'TEST')
        out = self.base / 'handoff'
        with self.assertRaisesRegex(self.p.PipelineError, 'aspect ratio'):
            b.handoff(self.p, self.project, out, '1,3')
        self.assertFalse(out.exists())

if __name__ == '__main__': unittest.main()
