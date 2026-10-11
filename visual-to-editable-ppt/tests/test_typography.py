"""Final-file font evidence, rather than requested manifest values."""
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('inspect_fonts', ROOT / 'scripts/inspect_fonts.py')
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)


class TypographyTests(unittest.TestCase):
    def test_all_themes_use_points_and_keep_body_below_titles(self):
        themes = list((ROOT / 'assets/themes').glob('*.json'))
        self.assertEqual(len(themes), 7)
        for path in themes:
            with self.subTest(theme=path.stem):
                typography = json.loads(path.read_text())['typography']
                self.assertEqual(typography['unit'], 'pt')
                self.assertEqual(typography['defaults_pt']['body'], 20)
                self.assertEqual(typography['roles_pt']['body'], [18, 22])
                self.assertLess(typography['roles_pt']['body'][1], typography['roles_pt']['content_title'][0])
                for role, target in typography['defaults_pt'].items():
                    low, high = typography['roles_pt'][role]
                    self.assertLessEqual(low, target)
                    self.assertLessEqual(target, high)
                self.assertFalse(any(key.endswith('_px') for key in typography))

    def test_read_actual_mixed_runs_paragraph_default_and_unknown(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'mixed.pptx'
            p = Presentation(); slide = p.slides.add_slide(p.slide_layouts[6])
            tf = slide.shapes.add_textbox(0, 0, Inches(5), Inches(2)).text_frame
            paragraph = tf.paragraphs[0]; paragraph.font.size = Pt(20)
            a = paragraph.add_run(); a.text = '正文'; a.font.name = 'Arial'
            b = paragraph.add_run(); b.text = '局部强调'; b.font.size = Pt(24)
            paragraph = tf.add_paragraph(); paragraph.add_run().text = '继承未知'
            table = slide.shapes.add_table(1, 1, 0, Inches(3), Inches(5), Inches(1)).table
            run = table.cell(0, 0).text_frame.paragraphs[0].add_run(); run.text = '表格'; run.font.size = Pt(18)
            group = slide.shapes.add_group_shape()
            run = group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text_frame.paragraphs[0].add_run()
            run.text = '分组标签'; run.font.size = Pt(22)
            p.save(path)
            report = m.inspect_fonts(path); by_text = {r['text']: r for r in report['text_runs']}
            self.assertEqual(by_text['正文']['size_pt'], 20)
            self.assertEqual(by_text['正文']['size_source'], 'paragraph')
            self.assertEqual(by_text['局部强调']['size_pt'], 24)
            self.assertIsNone(by_text['继承未知']['size_pt'])
            self.assertEqual(report['unknown_size_runs'], 1)
            self.assertEqual(by_text['表格']['size_pt'], 18)
            self.assertEqual(by_text['分组标签']['size_pt'], 22)
            self.assertNotIn('passed', report)


if __name__ == '__main__':
    unittest.main()
